"""Gate `public-sample` — the fictional company the landing page links.

The public sample is a claim made in public: "this is what the engine does
to a trial balance". Three things would make it a false one, and each has
happened to a neighbouring artefact in this repository:

  · the published files drift from what the engine produces today (the
    proof numbers on the landing page were typed by hand on 2026-09-08 and
    never re-measured);
  · the book does not hold together (the dashboard's own example carried
    an account 121 that disagreed with classes 6/7: the engine raised a
    CRITICAL finding on the product's own example and graded it AAA);
  · a client's label reaches a public file (the exported report carried a
    client label inside its stylesheet comments until this gate found it).

WHAT THIS GATE REDS ON, once the product is correct (TC-11):

  S1  the generator stops producing a consistent ledger — a column pair
      that does not sum, a class-6/7 account left open, an opening balance
      that is not last year's closing, a fiscal code that passes the
      checksum (it could then belong to a real company);
  S2  a published file is not byte-identical to a rebuild — the workbooks,
      the served documents, the mapping, the labels, and the page's data
      (so a figure on the /sample page that is not the served document's
      is red here);
  S3  two rebuilds differ (the sample is not reproducible);
  S4  the sample book fails a check the landing page's proof block lists:
      the balance sheet does not close to 0.00, net profit is not account
      121 to the cent, turnover is not the ledger's own (a fictional
      company has no Ministry of Finance filing; its referee is the ledger
      it was written from), an EBITDA variant differs from the methodology
      layer by more than 1 RON — or the engine raises a critical finding;
  S5  a client label, a fiscal code other than the fictional one, or a
      trade-register number other than the fictional one appears ANYWHERE in
      a file a visitor can download as it is — every file under
      public/sample, public/examples and public/templates, recursively: its
      raw bytes, each member of a workbook (cells, shared strings, document
      properties), the PDF's text layer, info dictionary and XMP packet
      (scripts/public_bytes.py) — or a workbook's author / subject /
      keywords, or the PDF's /Author, holds anything but the product's name;
      fewer than 16 files, 60 workbook members or 18 PDF parts were read;
  S6  the mapping stops covering every account exactly once, leaves a mapped
      balance-sheet account without its served row, or names a P&L line
      whose accounts do not sum to the served line; the labels stop being
      the engine's own sentences, or OMIT a caveat the served document
      states (every disclosure / basis / note / refusal field is walked;
      what is not a label must be on a reasoned exclusion list); the cash
      movement the page prints beside the engine's estimate is not the two
      served balance sheets';
  S7  the /sample route loses its place outside the auth wall, or nginx
      stops handing the route to the app (the directory of the same name
      would answer 403);
  S8  the dashboard's two example workbooks are not a rebuild, say
      "anonimizat", or read differently in their two layouts;
  S9  the report is not the same bytes under another system locale and
      timezone (it printed "RON 164.171" under ro_RO and "2 October" east
      of UTC+12 until 2026-10-02);
  S10 the report (HTML and PDF), the mapping CSV or the labels file does
      not say the company is fictional; or the report carries a sentence
      that is false for a public, model-free sample ("AI-assisted",
      "Confidential", "as filed" / "filed close", "implying an opening
      balance");
  S11 a known issue's "ledger" figure is not what the PUBLISHED workbook
      gives (every sum is repeated here from the workbook's own cells), an
      "engine" figure is not the served document's, the page's list is not
      the published file's, or the estimated cash flow is published with
      no issue listed; the stock movement the page prints beside the
      account-121 reading is not the workbook's, or the derived 711 line
      no longer equals it.

WHAT S5 CANNOT SEE: a client's name that is in no real book's file name
(the labels are derived from those names); a word split across two XML runs
of a workbook; pixels; the compressed bodies of the PDF other than its text
layer (fonts, images, a form field's stream); a published file OUTSIDE the
three directories (public/og, public/logos, public/geo — images and a map).

WHAT IT CANNOT SEE: the report's HTML and PDF. Those are built by the
product's TypeScript — `frontend/pages/cfo/__tests__/publicSample.test.tsx`
holds the HTML to a byte-identical rebuild and the PDF to its text, and
`node scripts/build_public_sample_report.mjs --check` re-prints the PDF.

NO FAKE ASSEMBLER, NO MIRROR STORE: `build_public_sample.serve` drives the
production write seam and `create_app()` (tests/engine/
_real_app_comparatives.py). Plants: docs/engine_book/gates.md,
"public-sample".
"""
from __future__ import annotations

import io
import json
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path
from typing import Any, Dict, Iterator, List, Tuple

import pytest

REPO = Path(__file__).resolve().parents[2]
for _p in (REPO / "scripts", REPO / "src"):
    if str(_p) not in sys.path:
        sys.path.insert(0, str(_p))

import build_public_examples as EX  # noqa: E402
import build_public_sample as B  # noqa: E402
import build_public_sample_tb as TB  # noqa: E402

PUBLIC = B.PUBLIC_DIR
CENT = 0.005


# ── fixtures: one rebuild, read by every law ──────────────────────────


@pytest.fixture(scope="module")
def books():
    return TB.build_books()


@pytest.fixture(scope="module")
def rebuilt() -> Dict[str, bytes]:
    """Steps 1-3 of the build, run now, from the engine."""
    return dict(B.engine_files())


@pytest.fixture(scope="module")
def served() -> Dict[str, Any]:
    """The COMMITTED served documents (what the page and the report read)."""
    def load(name: str) -> Any:
        return json.loads((PUBLIC / name).read_text(encoding="utf-8"))

    return {
        B.CURRENT_YEAR: load(B.SERVED_PERIOD % B.CURRENT_YEAR),
        B.PRIOR_YEAR: load(B.SERVED_PERIOD % B.PRIOR_YEAR),
        "comparatives": load(B.SERVED_COMPARATIVES),
    }


@pytest.fixture(scope="module")
def page_data() -> Dict[str, Any]:
    return json.loads(B.PAGE_DATA.read_text(encoding="utf-8"))


# ── S1 — the ledger ───────────────────────────────────────────────────


def test_s1_every_column_pair_sums_and_the_book_is_closed(books):
    for year, book in books.items():
        rows = book["rows"]
        assert len(rows) >= 80, "FY%d lists %d accounts — the sample is a full chart" % (year, len(rows))
        for name, i in (("opening", 2), ("movement", 4), ("cumulative", 6), ("closing", 8)):
            debit = sum(r[i] for r in rows)
            credit = sum(r[i + 1] for r in rows)
            assert debit == credit, "FY%d %s: debit %s != credit %s" % (year, name, debit, credit)
        for r in rows:
            if r[0][0] in "67":
                assert r[6] == r[7] and r[8] == 0 and r[9] == 0, (
                    "FY%d account %s is not closed into 121 (cumulative %s / %s, closing %s / %s)"
                    % (year, r[0], r[6], r[7], r[8], r[9]))


def test_s1_account_121_is_the_years_profit_and_last_years_is_cleared(books):
    for year, book in books.items():
        row = next(r for r in book["rows"] if r[0] == "121")
        closing = row[9] - row[8]
        assert closing == book["facts"]["net_profit"], (year, closing, book["facts"]["net_profit"])
        # the opening (last year's result) leaves 121 in the year: guard G6
        retained = next(r for r in book["rows"] if r[0] == "117")
        assert retained[5] == book["facts"]["prior_result_transferred"] == row[3] - row[2], (
            "FY%d: 117 received %s, 121 opened with %s" % (year, retained[5], row[3] - row[2]))


def test_s1_the_current_year_opens_with_the_prior_years_closing(books):
    prior = {r[0]: r[8] - r[9] for r in books[B.PRIOR_YEAR]["rows"]}
    current = {r[0]: r[2] - r[3] for r in books[B.CURRENT_YEAR]["rows"]}
    codes = sorted(set(prior) | set(current))
    differing = [(c, prior.get(c, 0), current.get(c, 0)) for c in codes
                 if prior.get(c, 0) != current.get(c, 0)]
    assert not differing, "opening FY%d != closing FY%d on: %s" % (
        B.CURRENT_YEAR, B.PRIOR_YEAR, differing[:8])


def test_s1_the_fiscal_code_cannot_belong_to_a_real_company():
    TB.assert_fiscal_code_is_invalid()
    # the checker itself is not vacuous: it refuses a code that passes.
    digits = re.sub(r"\D", "", TB.FISCAL_CODE)
    body = digits[:-1]
    weights = "753217532"[-len(body):]
    control = (sum(int(a) * int(b) for a, b in zip(body, weights)) * 10) % 11 % 10
    with pytest.raises(AssertionError):
        TB.assert_fiscal_code_is_invalid("RO %s%d" % (body, control))
    assert "EXEMPLU" in TB.COMPANY_NAME and "FICTIV" in TB.COMPANY_NAME


# ── S2 / S3 — the published files are the engine's, reproducibly ──────


def test_s2_every_published_engine_file_is_a_byte_identical_rebuild(rebuilt):
    assert len(rebuilt) == 8, sorted(rebuilt)
    stale = [name for name, data in rebuilt.items()
             if not (PUBLIC / name).is_file() or (PUBLIC / name).read_bytes() != data]
    assert not stale, (
        "public/sample is not what the engine produces today: %s differ from a rebuild. "
        "Run scripts/build_public_sample.py and commit the result — after reading the diff."
        % stale)


def test_s2_the_pages_data_is_the_served_document(page_data, served):
    """frontend/data/publicSample.json is `page_data()` over the published
    files — and every figure in it resolves, through its own pointer, to
    the value the served document carries."""
    expected = B.canonical_json(B.page_data(PUBLIC))
    assert B.PAGE_DATA.read_text(encoding="utf-8") == expected, (
        "frontend/data/publicSample.json is stale: the /sample page would print figures "
        "that are not the published served document's. Run scripts/build_public_sample.py.")
    checked = 0
    for which, year in (("current", B.CURRENT_YEAR), ("prior", B.PRIOR_YEAR)):
        for figure in page_data["periods"][which]["figures"]:
            assert B.pointer(served[year], figure["pointer"]) == figure["value"], figure
            checked += 1
    for ratio in page_data["ratios"]:
        row = B.pointer(served["comparatives"], ratio["pointer"])
        for side in ("current", "prior"):
            assert row[side]["value_q"] == ratio[side]["value_q"], (ratio["key"], side)
            assert row[side]["band"] == ratio[side]["band"], (ratio["key"], side)
            checked += 1
    for variant in page_data["verdicts"]["ebitda"]["variants"]:
        doc = served[B.CURRENT_YEAR]
        assert B.pointer(doc, variant["pointer_engine"]) == variant["engine"]
        assert B.pointer(doc, variant["pointer_methodology"]) == variant["methodology"]
        checked += 2
    assert checked >= 50, "only %d figures were held to the served document" % checked


def test_s3_a_second_rebuild_is_byte_identical(rebuilt):
    again = B.engine_files()
    assert list(again) == list(rebuilt)
    differing = [name for name in rebuilt if again[name] != rebuilt[name]]
    assert not differing, "two rebuilds of the public sample differ on: %s" % differing


# ── S4 — the sample satisfies the checks the proof block lists ────────


@pytest.mark.parametrize("year", [B.CURRENT_YEAR, B.PRIOR_YEAR])
def test_s4_the_balance_sheet_closes(served, year):
    cbs = served[year]["statements"]["canonical_bs"]
    assert cbs["status"] == "BALANCED", cbs["status"]
    assert abs(cbs["difference"]) < CENT, cbs["difference"]
    totals = cbs["totals"]
    assert abs(totals["assets"] - totals["equity_plus_liabilities"]) < CENT, totals
    assert cbs["source_anchor"]["anchor_status"] == "MATCHED"
    assert not cbs["unmapped"], cbs["unmapped"]
    assert cbs["extraction"]["method"] == "deterministic"


@pytest.mark.parametrize("year", [B.CURRENT_YEAR, B.PRIOR_YEAR])
def test_s4_net_profit_is_account_121_and_turnover_is_the_ledgers(served, books, year):
    pl = served[year]["statements"]["assembled_pl"]
    facts = books[year]["facts"]
    assert pl["net_income_anchor_status"] == "anchored", pl["net_income_anchor_status"]
    assert abs(pl["net_income_statutory"] - float(facts["net_profit"])) < CENT, (
        "FY%d: served net profit %s, account 121 in the ledger %s"
        % (year, pl["net_income_statutory"], facts["net_profit"]))
    assert abs(pl["net_income_statutory_anchor"] - float(facts["net_profit"])) < CENT
    # The fictional company files nothing with the Ministry of Finance; the
    # ledger the trial balance was read off is its referee.
    assert abs(pl["turnover"] - float(facts["net_turnover"])) < CENT, (
        "FY%d: served turnover %s, the ledger's 70x - 709 is %s"
        % (year, pl["turnover"], facts["net_turnover"]))


@pytest.mark.parametrize("year", [B.CURRENT_YEAR, B.PRIOR_YEAR])
def test_s4_the_stock_variation_is_recovered_and_ebitda_variants_agree(served, books, year):
    pl = served[year]["statements"]["assembled_pl"]
    variation = pl["inventory_variation"]
    assert variation["refusal"] is None and pl["ebitda_refusal"] is None
    assert variation["provenance"] == "account_121_bridge", variation["provenance"]
    # account 711 is closed on the book (debit == credit): the net the
    # ledger carried before its closing entry is what the bridge must find.
    assert abs(variation["value"] - float(books[year]["facts"]["stock_variation_711"])) < CENT, (
        year, variation["value"], books[year]["facts"]["stock_variation_711"])
    verdict = B.ebitda_verdict(served[year])
    assert [v["key"] for v in verdict["variants"]] == ["reported", "strict", "cash"]
    for v in verdict["variants"]:
        assert abs(v["methodology"] - v["engine"]) <= 1.0, (
            "FY%d EBITDA %s: methodology %s vs engine %s — more than 1 RON apart"
            % (year, v["key"], v["methodology"], v["engine"]))


@pytest.mark.parametrize("year", [B.CURRENT_YEAR, B.PRIOR_YEAR])
def test_s4_the_engine_raises_no_critical_finding_on_its_own_sample(served, year):
    insights = served[year]["statements"]["insights"]["insights"]
    assert insights, "the engine raised no finding at all — the sample shows nothing"
    critical = [i["id"] for i in insights if i["severity"]["level"] == "critical"]
    assert not critical, "FY%d: critical findings on the public sample: %s" % (year, critical)
    inv = served[year]["statements"]["inventory_days"]
    assert inv["status"] == "served" and inv["basis"] == "average_two_year_ends", (
        inv["status"], inv["basis"])
    assert inv["reconciliation"]["status"] == "reconciled"


# ── S5 — nothing of a client's reaches a public file ──────────────────


def _published_texts() -> Iterator[Tuple[str, str]]:
    """(file, every character a reader of it can see) for each file under
    public/sample, the page's data and the two example workbooks."""
    from openpyxl import load_workbook

    paths = sorted(p for p in PUBLIC.iterdir() if p.is_file())
    paths += [B.PAGE_DATA] + [EX.EXAMPLES_DIR / name for name in (EX.FOUR_PAIR, EX.COMPACT)]
    for path in paths:
        if path.suffix == ".xlsx":
            wb = load_workbook(io.BytesIO(path.read_bytes()), read_only=True)
            cells = [str(c) for ws in wb.worksheets for row in ws.iter_rows(values_only=True)
                     for c in row if c is not None]
            props = wb.properties
            cells += [str(getattr(props, k) or "") for k in ("creator", "lastModifiedBy", "title")]
            yield path.name, "\n".join(cells)
        elif path.suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(path.read_bytes()))
            yield path.name, "\n".join((page.extract_text() or "") for page in reader.pages)
        else:
            yield path.name, path.read_text(encoding="utf-8")


def _client_labels() -> List[str]:
    """The labels of the REAL books this repository holds, read off their
    file and directory names at run time — never typed (scripts/
    client_labels.py, shared with the `bundle-labels` gate)."""
    import client_labels as CL

    return CL.client_labels()


#: Every directory whose files a visitor can download as they are.
PUBLISHED_ROOTS = (REPO / "public" / "sample", REPO / "public" / "examples", REPO / "public" / "templates")


def _published_parts() -> List[Tuple[str, str]]:
    """EVERY PLACE A WORD CAN SIT in a published file (scripts/public_bytes.py):
    the raw bytes of every file under public/sample, public/examples and
    public/templates, recursively; each member of a workbook, decompressed —
    the sheets, the shared strings, docProps/core.xml and app.xml (title,
    subject, keywords, creator, company …); the PDF's text layer, info
    dictionary, XMP packet and uncompressed objects; and the page's data.

    The first version read cell values and three properties of the files
    directly under public/sample. A label in a workbook's `subject` /
    `keywords`, in the PDF's /Author, or in public/sample/extra/notes.txt
    passed it; so did a label and a fiscal code typed into the sales example
    (2026-10-02)."""
    import public_bytes as PB

    out: List[Tuple[str, str]] = []
    for path in PB.files_under(*PUBLISHED_ROOTS):
        out += PB.parts(path, path.relative_to(REPO).as_posix())
    out += PB.parts(B.PAGE_DATA, B.PAGE_DATA.relative_to(REPO).as_posix())
    return out


def _scan_floor(parts: List[Tuple[str, str]]) -> Tuple[int, int, int]:
    """(files, workbook members, PDF parts) — a scan that stops reading a
    container must not pass."""
    files = {name.split("!")[0] for name, _ in parts}
    members = sum(1 for name, _ in parts if ".xlsx!" in name and not name.endswith("!<names>"))
    pdf = sum(1 for name, _ in parts if ".pdf!" in name)
    return len(files), members, pdf


def test_s5_no_client_label_appears_in_a_published_file():
    labels = _client_labels()
    assert len(labels) >= 4, "only %d client labels were derived — the scan is vacuous" % len(labels)
    pattern = re.compile(r"(?<![A-Za-z])(%s)(?![A-Za-z])" % "|".join(map(re.escape, labels)), re.I)
    parts = _published_parts()
    for name, text in parts:
        hit = pattern.search(text)
        # The label itself is not printed: a failure log is not the place
        # for a client's name either. Its position is enough to find it.
        assert hit is None, (
            "%s carries the label of a real book this repository holds, at character %d "
            "(%d characters long). Find where it comes from and remove it at the source."
            % (name, hit.start(), len(hit.group(0))))
    files, members, pdf = _scan_floor(parts)
    print("GATE-WORK public-sample-scan files=%d workbook_members=%d pdf_parts=%d labels=%d"
          % (files, members, pdf, len(labels)))
    # measured 2026-10-02: 16 files (10 + 3 + 2 and the page data), 66
    # workbook members over 7 workbooks, 32 PDF parts
    assert files >= 16, "only %d published files were scanned" % files
    assert members >= 60, "only %d workbook members were read" % members
    assert pdf >= 18, "only %d parts of the PDF were read" % pdf
    # the document properties and the PDF's metadata ARE among the parts
    names = [name for name, _ in parts]
    for must in ("!docProps/core.xml", "!docProps/app.xml", ".pdf!info", ".pdf!xmp", ".pdf!objects"):
        assert any(name.endswith(must) for name in names), "no part %s was read" % must
    for root in PUBLISHED_ROOTS:
        assert any(name.startswith(root.relative_to(REPO).as_posix() + "/") for name in names), root


def test_s5_document_properties_carry_the_product_name_only():
    """A workbook's properties and the PDF's info dictionary are read by
    every file manager and search index. They hold the product's name, the
    document's own title, the writing library's name — nobody's name, and
    nothing a generator copied from a source workbook."""
    from openpyxl import load_workbook
    from pypdf import PdfReader
    import public_bytes as PB

    seen = 0
    for path in PB.files_under(*PUBLISHED_ROOTS):
        rel = path.relative_to(REPO).as_posix()
        if path.suffix == ".xlsx":
            props = load_workbook(io.BytesIO(path.read_bytes()), read_only=True).properties
            for field in ("creator", "lastModifiedBy"):
                # the product's name (alone or with the sample's own
                # "fictional example" mark) or the writing library's default
                # — never a person, never a client
                assert (getattr(props, field) or "") in ("", "CFO AI", "CFO AI - exemplu fictiv", "openpyxl"), (
                    "%s: %s is %r — a workbook's author is the product, the writing library, or nobody"
                    % (rel, field, getattr(props, field)))
            for field in ("subject", "keywords", "category", "description", "contentStatus", "identifier"):
                assert not (getattr(props, field) or ""), (
                    "%s: the %s property is set (%d characters) — nothing of ours writes one"
                    % (rel, field, len(str(getattr(props, field)))))
            seen += 1
        elif path.suffix == ".pdf":
            info = dict(PdfReader(io.BytesIO(path.read_bytes())).metadata or {})
            for key in ("/Author", "/Subject", "/Keywords"):
                assert not info.get(key), "%s: the PDF's %s is set — the report writes none" % (rel, key)
            seen += 1
    assert seen >= 8, "only %d documents with properties were examined" % seen


_FISCAL_CODE = re.compile(r"(?:\bRO\s?|\bCUI\W{0,3}(?:RO\s?)?|\bCIF\W{0,3}(?:RO\s?)?|\bc\.\s?f\.\W{0,3}(?:RO\s?)?)(\d{2,10})\b",
                          re.I)
_TRADE_REGISTER = re.compile(r"\bJ\s?\d{1,2}\s?/\s?\d{1,7}\s?/\s?\d{4}\b")


def test_s5_the_only_fiscal_code_and_register_number_are_the_fictional_ones():
    fictional = re.sub(r"\D", "", TB.FISCAL_CODE)
    seen_code = seen_register = 0
    for name, text in _published_parts():
        for match in _FISCAL_CODE.finditer(text):
            seen_code += 1
            assert match.group(1) == fictional, (
                "%s carries a fiscal code that is not the sample's fictional one, at "
                "character %d" % (name, match.start()))
        for match in _TRADE_REGISTER.finditer(text):
            seen_register += 1
            assert re.sub(r"\s", "", match.group(0)) == TB.TRADE_REGISTER, (
                "%s carries a trade-register number that is not the sample's fictional one, "
                "at character %d" % (name, match.start()))
    assert seen_code >= 4 and seen_register >= 4, (
        "the fictional identifiers were found %d / %d times — the scan read nothing"
        % (seen_code, seen_register))


# ── S6 — the mapping and the labels ───────────────────────────────────


def test_s6_the_mapping_lists_every_account_exactly_once(served, books, page_data):
    mapping = page_data["mapping"]
    accounts = [r[0] for r in books[B.CURRENT_YEAR]["rows"]]
    assert [m["account"] for m in mapping] == accounts
    served_items = {str(i["ro_account_code"]): i for i in served[B.CURRENT_YEAR]["line_items"]}
    for row in mapping:
        if row["status"] == "mapped":
            item = served_items[row["account"]]
            assert (row["statement"], row["engine_bucket"], row["amount_ron"]) == (
                item["statement"], item["bucket"], item["amount"]), row
        else:
            assert row["account"] not in served_items, row
            assert row["status"].startswith("excluded:") or row["status"] == "no_closing_balance", row
    csv_lines = (PUBLIC / B.MAPPING_CSV).read_text(encoding="utf-8-sig").splitlines()
    assert csv_lines[0] == B.MAPPING_NOTICE
    assert csv_lines[1].split(",") == list(B.MAPPING_COLUMNS)
    assert len(csv_lines) == len(accounts) + 2


def test_s6_every_mapped_balance_sheet_account_names_its_served_row(served, page_data):
    """An account the engine carries to the balance sheet says WHICH row —
    five analytic accounts (the engine lists their rows by prefix) had none
    until 2026-10-02 — and the rows' accounts add up to the served rows."""
    cbs = served[B.CURRENT_YEAR]["statements"]["canonical_bs"]
    rows = {row["id"]: row for row in cbs["rows"]}
    by_row: Dict[str, float] = {}
    checked = 0
    for m in page_data["mapping"]:
        if m["status"] == "mapped" and m["statement"] == "BS":
            assert m["balance_sheet_row"] in rows and m["balance_sheet_section"], (
                "account %s is mapped to the balance sheet with no served row" % m["account"])
            assert m["balance_sheet_section"] == rows[m["balance_sheet_row"]]["section"], m
            by_row[m["balance_sheet_row"]] = by_row.get(m["balance_sheet_row"], 0.0) + m["amount_ron"]
            checked += 1
        if m["statement"] == "PL":
            assert m["balance_sheet_row"] is None and m["balance_sheet_section"] is None, m
    assert checked >= 35, "only %d balance-sheet accounts were checked" % checked
    for row_id, total in by_row.items():
        assert abs(abs(total) - abs(rows[row_id]["amount"])) < CENT, (
            "row %s: its accounts carry %.2f, the served row %.2f" % (row_id, total, rows[row_id]["amount"]))
    assert set(by_row) == set(rows) - {"current_year_profit"}, (
        "served rows with no account in the mapping: %s" % sorted(set(rows) - set(by_row)))


def test_s6_each_pl_account_names_the_served_line_it_sums_into(served, page_data):
    """The mapping's `pl_line` is the served P&L line, and the accounts it
    names add up to that line to the cent. 7812 / 7814 were listed as "other
    operating income" although the engine nets them, with 6812 / 6814, into
    the provisions line outside EBITDA: summing the file by bucket gave an
    "other income" the report does not print."""
    pl = served[B.CURRENT_YEAR]["statements"]["assembled_pl"]
    sums: Dict[str, float] = {}
    for m in page_data["mapping"]:
        if m["statement"] != "PL":
            assert m["pl_line"] is None, m
            continue
        assert m["pl_line"], "P&L account %s names no served line" % m["account"]
        sums.setdefault(m["pl_line"], 0.0)
        if m["pl_line"] == "net_provisions":
            sign = 1.0 if m["account"] in pl["net_provisions"]["charges"]["by_account"] else -1.0
            sums[m["pl_line"]] += sign * m["amount_ron"]
        elif m["pl_line"] == "inventory_variation":
            # the account's amount is its GROSS turnover; the line is the note
            assert m["note"] and m["note"]["key"] == "served_net_derived_from_121", m
            sums[m["pl_line"]] += m["note"]["value"]
        else:
            sums[m["pl_line"]] += m["amount_ron"]
    assert len(sums) >= 9, sorted(sums)
    for line, total in sorted(sums.items()):
        node = pl[line]
        want = node["value"] if isinstance(node, dict) else node
        assert abs(total - want) < CENT, (
            "P&L line %s: the mapping's accounts sum to %.2f, the engine serves %.2f"
            % (line, total, want))


def test_s6_every_label_is_the_engines_own_sentence(served, page_data):
    labels = page_data["labels"]
    assert len(labels) >= 8, "the sample shows %d uncertainty labels" % len(labels)
    kinds = {label["kind"] for label in labels}
    assert {"approximation", "derived", "basis", "refusal"} <= kinds, kinds
    for label in labels:
        if label["reason"] is not None:        # a refusal: served code + inputs, no sentence
            assert label["engine_en"] is None and label["engine_ro"] is None
            row = B.pointer(served[B.CURRENT_YEAR], label["source"].rsplit("/", 1)[0])
            assert row["reason"]["code"] == label["reason"]["code"]
            continue
        doc = served["comparatives"] if label["source"].startswith("comparatives:") else served[B.CURRENT_YEAR]
        node = B.pointer(doc, label["source"].split(":", 1)[-1])
        carried = json.dumps(node, ensure_ascii=False)
        for text in (label["engine_en"], label["engine_ro"]):
            if text is not None:
                assert json.dumps(text, ensure_ascii=False).strip('"') in carried or text == str(node), (
                    "label %s quotes a sentence the served document does not carry at %s"
                    % (label["key"], label["source"]))
    assert served[B.CURRENT_YEAR]["statements"]["assembled_cf"]["is_approximated"] == any(
        label["key"].startswith("cash_flow_approximated") for label in labels)


def test_s6_no_caveat_the_engine_stated_is_left_off_the_label_list(served, page_data):
    """The page says "each uncertainty label". Every disclosure, basis,
    note and refusal the served documents state is a label, or sits on the
    exclusion list with its reason — and no exclusion is dead."""
    result = B.label_completeness(served[B.CURRENT_YEAR], served["comparatives"], page_data["labels"])
    assert result["examined"] >= 100, "only %d disclosure fields were walked" % result["examined"]
    assert not result["missing"], (
        "the served document states %d caveat(s) the label list omits; first: %s — %r. Add it "
        "to uncertainty_labels(), or to LABEL_EXCLUSIONS with the reason it is not one."
        % (len(result["missing"]), result["missing"][0][0], result["missing"][0][1][:120]))
    assert not result["dead_exclusions"], (
        "LABEL_EXCLUSIONS entries that match nothing: %s" % result["dead_exclusions"])
    keys = {label["key"] for label in page_data["labels"]}
    # the three that were missing until 2026-10-02
    assert {"bands_general_sme", "ccc_inventory_term_basis", "methodology_note_dscr_approx"} <= keys, keys
    published = json.loads((PUBLIC / B.LABELS_JSON).read_text(encoding="utf-8"))
    assert published["labels"] == page_data["labels"]


def test_s6_the_cash_movement_beside_the_estimate_is_the_two_balance_sheets(served, page_data):
    """The engine's cash flow is an estimate that does not read the prior
    year; the page prints the real movement beside it. That movement is the
    difference of two SERVED cash balances, re-read here."""
    cash_flow = page_data["verdicts"]["cash_flow"]
    closing = B.pointer(served[B.CURRENT_YEAR], cash_flow["pointer_cash"])
    opening = B.pointer(served[B.PRIOR_YEAR], cash_flow["pointer_cash"])
    assert (cash_flow["closing_cash"], cash_flow["prior_closing_cash"]) == (closing, opening)
    assert abs(cash_flow["balance_sheet_movement"] - (closing - opening)) < CENT
    assert cash_flow["net_change_estimated"] == B.pointer(
        served[B.CURRENT_YEAR], cash_flow["pointer_net_change"])
    assert cash_flow["is_approximated"] is True
    inv = page_data["verdicts"]["inventory_days"]
    total = served[B.CURRENT_YEAR]["statements"]["inventory_days"]["total"]
    assert inv["ccc_term_value_q"] == total["closing_value_q"] != total["value_q"], (
        "the cycle's inventory term is the period-end one, printed on its own")


# ── S11 — the known issues are the ledger's arithmetic ────────────────
#
# The sample exposed three defects of the engine (an estimated cash flow
# that contradicts the ledgers, a finding's quick ratio on another
# definition, an asset-age finding whose bases include land, construction
# in progress and intangibles). They are published FLAGGED: /sample and the
# report open with a "Known issues in this report" box, and each line
# carries the ledger's figure beside the engine's. This law is what makes
# "the ledger's figure" true: it reads the PUBLISHED workbook cell by cell
# (openpyxl — not the generator's rows) and repeats every sum on its own.


def _published_ledger(year: int) -> Dict[str, Tuple[float, ...]]:
    """{account: (opening D, opening C, movement D, movement C, total D,
    total C, closing D, closing C)} read from the published workbook."""
    import openpyxl

    sheet = openpyxl.load_workbook(PUBLIC / TB.workbook_filename(year), read_only=True)[TB.SHEET_NAME]
    out: Dict[str, Tuple[float, ...]] = {}
    for row in sheet.iter_rows(values_only=True):
        code = row[0]
        if isinstance(code, str) and code.isdigit():
            assert code not in out, "account %s is listed twice" % code
            out[code] = tuple(float(v or 0) for v in row[2:10])
    assert len(out) >= 80, "only %d accounts read from the published workbook" % len(out)
    return out


def _over(ledger: Dict[str, Tuple[float, ...]], prefixes: Tuple[str, ...], column: int) -> float:
    return sum(values[column] for code, values in ledger.items() if code.startswith(prefixes))


_OPEN_D, _OPEN_C, _MOVE_D, _MOVE_C, _CLOSE_D, _CLOSE_C = 0, 1, 2, 3, 6, 7


def test_s11_the_known_issues_are_listed_and_the_page_carries_the_published_list(page_data, served):
    published = json.loads((PUBLIC / B.KNOWN_ISSUES_JSON).read_text(encoding="utf-8"))
    assert published["schema"] == B.KNOWN_ISSUES_SCHEMA
    assert published["company"]["fictional"] is True
    assert page_data["known_issues"] == published["issues"], (
        "the /sample page would print a known-issues list that is not the published file's")
    ids = [issue["id"] for issue in published["issues"]]
    # While the engine estimates cash flow without the prior period, the
    # issue MUST be listed: an approximated cash flow published with no
    # flag is what this law exists to stop.
    cf = served[B.CURRENT_YEAR]["statements"]["assembled_cf"]
    cash = page_data["verdicts"]["cash_flow"]
    if cf["is_approximated"] or abs(cf["net_change_in_cash"] - cash["balance_sheet_movement"]) > CENT:
        assert "cash_flow_estimated" in ids, "the estimated cash flow is published without its known issue"
    assert ids == sorted(set(ids), key=ids.index), "an issue is listed twice"


def test_s11_every_ledger_figure_of_a_known_issue_is_repeated_from_the_published_workbook(served):
    issues = {issue["id"]: issue
              for issue in json.loads((PUBLIC / B.KNOWN_ISSUES_JSON).read_text(encoding="utf-8"))["issues"]}
    ledger = _published_ledger(B.CURRENT_YEAR)
    prior = _published_ledger(B.PRIOR_YEAR)
    checked = 0

    def same(stated: float, recomputed: float, what: str) -> None:
        nonlocal checked
        assert abs(stated - recomputed) < CENT, "%s: stated %r, the published workbook gives %r" % (
            what, stated, recomputed)
        checked += 1

    cash_codes = ("5121", "5124", "5311")
    if "cash_flow_estimated" in issues:
        issue = issues["cash_flow_estimated"]
        led, eng = issue["ledger"], issue["engine"]
        opening = _over(ledger, cash_codes, _OPEN_D) - _over(ledger, cash_codes, _OPEN_C)
        closing = _over(ledger, cash_codes, _CLOSE_D) - _over(ledger, cash_codes, _CLOSE_C)
        same(led["cash_opening"], opening, "cash at the start of the year")
        # …which is the PRIOR workbook's closing cash, not just this one's opening column
        same(led["cash_opening"], _over(prior, cash_codes, _CLOSE_D) - _over(prior, cash_codes, _CLOSE_C),
             "cash at the end of the prior year")
        same(led["cash_closing"], closing, "cash at the end of the year")
        same(led["cash_movement"], closing - opening, "the year's cash movement")
        same(led["fixed_asset_additions_ex_vat"],
             _over(ledger, ("20", "21", "23"), _MOVE_D) - _over(ledger, ("231",), _MOVE_C),
             "fixed-asset additions")
        same(led["paid_to_fixed_asset_suppliers"], _over(ledger, ("404",), _MOVE_D), "paid to fixed-asset suppliers")
        drawn = _over(ledger, ("1621",), _MOVE_C)
        repaid = _over(ledger, ("1621",), _MOVE_D)
        lease = _over(ledger, ("167",), _MOVE_D)
        line = ((_over(ledger, ("5191",), _CLOSE_C) - _over(ledger, ("5191",), _CLOSE_D))
                - (_over(ledger, ("5191",), _OPEN_C) - _over(ledger, ("5191",), _OPEN_D)))
        dividends = _over(ledger, ("457",), _MOVE_D)
        interest = _over(ledger, ("666",), _MOVE_D)
        for key, value in (("loans_drawn", drawn), ("loans_repaid", repaid), ("lease_repaid", lease),
                           ("credit_line_change", line), ("dividends", dividends), ("interest", interest)):
            same(led[key], value, key)
        same(led["financing"], drawn - repaid - lease + line - dividends - interest, "financing on the ledger")
        # the engine's side is the served document's, through its own pointers
        doc = served[B.CURRENT_YEAR]
        for key in ("net_change_in_cash", "cash_from_investing", "cash_from_financing", "cash_from_operating"):
            assert eng[key] == B.pointer(doc, eng["pointers"][key]), key
            checked += 1
        same(eng["implied_opening_cash"],
             B.pointer(doc, eng["pointers"]["closing_cash_actual"]) - eng["net_change_in_cash"],
             "the opening cash the engine's net change implies")
        # and the issue is REAL: the engine's net change is not the ledger's
        assert abs(eng["net_change_in_cash"] - led["cash_movement"]) > CENT or eng["is_approximated"]

    if "quick_ratio_two_definitions" in issues:
        issue = issues["quick_ratio_two_definitions"]
        table, finding = issue["ratio_table"], issue["finding"]
        side = B.pointer(served["comparatives"], table["pointer"])
        assert side["value"] == table["value"] and side["value_q"] == table["value_q"]
        assert B.pointer(served[B.CURRENT_YEAR], finding["pointer"]) == finding["value"]
        assert abs((table["cash"] + table["trade_receivables"]) / table["current_liabilities"]
                   - table["value"]) < 1e-9, "the ratio table's quick ratio is not cash + trade receivables"
        assert abs((finding["current_assets"] - finding["inventory"]) / finding["current_liabilities"]
                   - finding["value"]) < 1e-9, "the finding's quick ratio is not current assets less inventory"
        same(table["cash"], _over(ledger, cash_codes, _CLOSE_D) - _over(ledger, cash_codes, _CLOSE_C),
             "cash in the quick ratio")
        assert abs(table["value"] - finding["value"]) >= 0.0005, "the two quick ratios agree — the issue is stale"
        checked += 4

    if "asset_age_bases" in issues:
        issue = issues["asset_age_bases"]
        led, eng = issue["ledger"], issue["engine"]
        gross = _over(ledger, ("212", "213", "214"), _CLOSE_D)
        accumulated = _over(ledger, ("281",), _CLOSE_C) - _over(ledger, ("281",), _CLOSE_D)
        charge = _over(ledger, ("281",), _MOVE_C)
        same(led["depreciable_tangible_gross"], gross, "depreciable tangible assets, gross")
        same(led["accumulated_depreciation"], accumulated, "accumulated depreciation")
        same(led["annual_depreciation_charge"], charge, "the year's depreciation charge")
        same(led["remaining_net_book_value"], gross - accumulated, "remaining net book value")
        same(led["land"], _over(ledger, ("211",), _CLOSE_D), "land")
        same(led["construction_in_progress"], _over(ledger, ("231",), _CLOSE_D), "construction in progress")
        same(led["intangibles_net"],
             _over(ledger, ("20",), _CLOSE_D) - (_over(ledger, ("280",), _CLOSE_C) - _over(ledger, ("280",), _CLOSE_D)),
             "intangible assets, net")
        assert abs(led["depreciated_share"] - accumulated / gross) < 1e-6
        assert abs(led["remaining_life_years"] - (gross - accumulated) / charge) < 1e-6
        measures = {m["key"]: m["value"] for m in B.pointer(served[B.CURRENT_YEAR], eng["pointer"])}
        assert measures["depreciated_share"] == eng["depreciated_share"]
        assert measures["remaining_book_life"] == eng["remaining_book_life_years"]
        assert measures["net_book_value"] == eng["net_book_value"]
        # the engine's bases DO include what the box says they include
        same(eng["gross_ppe"], gross + led["land"] + led["construction_in_progress"],
             "the finding's gross base = depreciable tangible + land + construction in progress")
        checked += 5
    assert checked >= 12 * min(1, len(issues)), "only %d figures were repeated from the workbook" % checked
    print("GATE-WORK public-sample known-issue figures=%d issues=%d" % (checked, len(issues)))


def test_s11_the_stock_movement_beside_the_121_reading_is_the_published_ledgers(page_data, served, books):
    """On a closed book "net income equals account 121" holds by
    construction; the page prints the real cross-check beside it — the
    derived 711 line against the movement of work in progress and finished
    goods. That movement is re-read here from the published workbook."""
    stated = page_data["verdicts"]["stock_variation"]["ledger_stock_movement"]
    ledger = _published_ledger(B.CURRENT_YEAR)
    codes = ("331", "341", "345", "348")
    opening = _over(ledger, codes, _OPEN_D) - _over(ledger, codes, _OPEN_C)
    closing = _over(ledger, codes, _CLOSE_D) - _over(ledger, codes, _CLOSE_C)
    assert abs(stated["opening"] - opening) < CENT and abs(stated["closing"] - closing) < CENT
    assert abs(stated["value"] - (closing - opening)) < CENT
    assert stated["accounts"] == sorted(code for code in ledger if code.startswith(codes))
    # the cross-check itself: the engine's DERIVED line equals the ledger's movement
    derived = served[B.CURRENT_YEAR]["statements"]["assembled_pl"]["inventory_variation"]["value"]
    assert abs(derived - stated["value"]) < CENT, (
        "the derived 711 line (%r) is not the ledger's stock movement (%r)" % (derived, stated["value"]))
    assert abs(float(books[B.CURRENT_YEAR]["facts"]["stock_variation_711"]) - stated["value"]) < CENT


# ── S7 — the page is reachable, signed out, by URL ────────────────────


def test_s7_the_route_is_public_and_nginx_hands_it_to_the_app():
    app = (REPO / "frontend" / "App.tsx").read_text(encoding="utf-8")
    route = re.search(r'<Route\s+path="/sample"\s+element=\{([^}]*)\}', app)
    assert route, "App.tsx registers no /sample route"
    assert "AuthGuard" not in route.group(1) and "FeatureRoute" not in route.group(1), (
        "the /sample route is behind a wall: %s" % route.group(1))
    layout_at = app.index("<Route element={<AppLayout />}>")
    assert app.index('path="/sample"') < layout_at, (
        "the /sample route sits inside the signed-in layout")
    # public/sample is a DIRECTORY in the built image: without its own
    # location the SPA fallback's `$uri/` answers /sample with a 403.
    nginx = (REPO / "nginx.conf").read_text(encoding="utf-8")
    bare = "\n".join(line.split("#", 1)[0] for line in nginx.splitlines())
    block = re.search(r"location\s+~\s+\^/sample/\?\$\s*\{([^}]*)\}", bare)
    assert block, "nginx.conf has no location for the /sample route"
    assert re.search(r"rewrite\s+\^\s+/index\.html\s+last;", block.group(1)), block.group(1)
    assert bare.index("location ~ ^/sample/?$") < bare.index("location / {"), (
        "the /sample location must be declared before the SPA fallback it pre-empts")


# ── S8 — the dashboard's two example workbooks ────────────────────────


def test_s8_the_examples_are_a_rebuild_and_say_fictional():
    assert EX.stale_files() == [], (
        "public/examples is stale — run scripts/build_public_examples.py")
    for name, text in _published_texts():
        if name in (EX.FOUR_PAIR, EX.COMPACT):
            assert "date fictive" in text, name
            assert "anonimiz" not in text.lower(), (
                "%s says the book is anonymised; it is fictional, which is a different claim" % name)


def test_s8_both_example_layouts_read_as_one_consistent_book():
    reads = {}
    for name, data in EX.example_files().items():
        body = B.serve_workbooks(
            [(EX.EXAMPLE_YEAR, name.replace(".xlsx", "_31.12.%d.xlsx" % EX.EXAMPLE_YEAR), data)]
        )["periods"][EX.EXAMPLE_YEAR]
        st = body["statements"]
        cbs, pl = st["canonical_bs"], st["assembled_pl"]
        assert cbs["status"] == "BALANCED" and abs(cbs["difference"]) < CENT, (name, cbs["status"])
        check = cbs["invariants"]["p121_cross_check"]
        assert check["ok"] and abs(check["p121"] - check["cls7_minus_cls6"]) < CENT, (
            "%s: account 121 (%s) disagrees with classes 6/7 (%s) — the defect the old "
            "example shipped with" % (name, check["p121"], check["cls7_minus_cls6"]))
        assert pl["net_income_anchor_status"] == "anchored" and pl["ebitda_refusal"] is None
        levels = [i["severity"]["level"] for i in st["insights"]["insights"]]
        assert "critical" not in levels, (name, levels)
        reads[name] = (cbs["extraction"]["source_format"], pl["turnover"], pl["net_income_statutory"],
                       cbs["totals"]["assets"], pl["ebitda"])
    four, compact = reads[EX.FOUR_PAIR], reads[EX.COMPACT]
    assert four[0] != compact[0], "the two examples are the same layout: %s" % four[0]
    assert four[1:] == compact[1:], "the two layouts of one book read differently: %s vs %s" % (
        four, compact)


# ── S9 — the same bytes under another locale and timezone ─────────────


def test_s9_the_report_is_the_same_bytes_under_another_locale_and_timezone():
    """`build_public_sample_report.mjs --check --no-pdf` under a Romanian
    system locale and a timezone thirteen hours east. Until 2026-10-02 both
    changed the report: a recommendation printed its amounts with the
    machine's own grouping ("RON 164.171"), and the cover printed the day in
    local time ("2 October")."""
    if shutil.which("node") is None or not (REPO / "node_modules").exists():
        pytest.skip("node / node_modules are not on this host: the report cannot be rebuilt here")
    env = dict(os.environ, LANG="ro_RO.UTF-8", LC_ALL="ro_RO.UTF-8", TZ="Pacific/Auckland")
    proc = subprocess.run(
        ["node", str(B.REPORT_SCRIPT), "--check", "--no-pdf"],
        cwd=str(REPO), env=env, capture_output=True, text=True, timeout=180)
    out = proc.stdout + proc.stderr
    assert proc.returncode == 0 and "byte-identical rebuild" in out, (
        "the sample report is not the same bytes under LANG=ro_RO.UTF-8 / TZ=Pacific/Auckland:\n%s"
        % out[-900:])
    # the environment was honoured: this Node formats a number the Romanian way
    probe = subprocess.run(
        ["node", "-e", "process.stdout.write((1234567.5).toLocaleString() + '|' + "
                       "new Date('2026-10-01T12:00:00Z').getDate())"],
        env=env, capture_output=True, text=True, timeout=60)
    grouped, _, day = probe.stdout.partition("|")
    assert day == "2", "TZ was not applied to the rebuild (day %r)" % day
    assert grouped.startswith("1.234.567"), (
        "this Node build has no Romanian locale data (%r): the locale half of the law is "
        "vacuous on this host" % grouped)


# ── S10 — what the published files say about themselves ───────────────

#: Sentences the product's export printed on every report until 2026-10-02
#: that are false for a public sample built without a model.
_FALSE_ON_A_PUBLIC_SAMPLE = (
    r"AI-assisted", r"Confidential", r"\bas filed\b", r"filed close", r"company filed",
    r"implying an opening balance", r"at any level of effort",
    # 2026-10-02, the third review: the source of this report is a trial
    # balance, nobody's "own books" and no "filing"; nothing in it was
    # adjusted; a generated document names no bank.
    r"\bthe filing\b", r"your own books", r"\badjusted Debt/EBITDA", r"Banca Transilvania|ING Romania|\bBCR\b",
)


def _report_texts() -> List[Tuple[str, str]]:
    import html as _html

    out = []
    for name, text in _published_texts():
        if name == B.REPORT_HTML:
            body = text[text.index("</style>"):]
            body = re.sub(r"<script[\s\S]*?</script>", " ", body)
            out.append((name, re.sub(r"\s+", " ", _html.unescape(re.sub(r"<[^>]+>", " ", body)))))
        elif name == B.REPORT_PDF:
            out.append((name, re.sub(r"\s+", " ", text)))
    return out


def test_s10_the_report_says_the_company_is_fictional_and_nothing_false_about_itself():
    texts = _report_texts()
    assert [name for name, _ in texts] == [B.REPORT_HTML, B.REPORT_PDF]
    for name, text in texts:
        for word in ("Fictional company", "not a real entity", "Companie fictivă"):
            assert word in text, "%s does not say the company is fictional (%r)" % (name, word)
        for pattern in _FALSE_ON_A_PUBLIC_SAMPLE:
            hit = re.search(pattern, text)
            assert hit is None, "%s prints %r: …%s…" % (
                name, hit.group(0), text[max(0, hit.start() - 80):hit.end() + 80])
        assert "no AI model read or produced these figures" in text, (
            "%s does not say who read the document" % name)
    # every printed page of the PDF but the cover carries the short notice
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO((PUBLIC / B.REPORT_PDF).read_bytes()))
    pages = [(page.extract_text() or "") for page in reader.pages]
    without = [i + 1 for i, text in enumerate(pages) if "Fictional company" not in text]
    assert not without, "PDF pages with no fictional-company notice: %s" % without


def test_s10_the_mapping_and_the_labels_say_the_company_is_fictional():
    first = (PUBLIC / B.MAPPING_CSV).read_text(encoding="utf-8-sig").splitlines()[0]
    assert first.startswith("#") and "fictional company" in first and "companie fictivă" in first, first
    labels = json.loads((PUBLIC / B.LABELS_JSON).read_text(encoding="utf-8"))
    assert labels["company"]["fictional"] is True
    assert "Fictional company" in labels["company"]["notice_en"]
    assert "Companie fictivă" in labels["company"]["notice_ro"]
