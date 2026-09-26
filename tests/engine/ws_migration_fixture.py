"""Synthetic world for the one-company-per-workspace migration tests.

Every company, CUI, figure and file below is INVENTED; the files are built
here, in memory. The SHAPES are the owner's production ones (2026-09-21):

  * a company workspace ("alfa food") whose served Dec 2024 / Dec 2025
    periods must survive, but which also holds other companies' files — a
    real-estate sister company's book attached to the served Dec 2025
    period, a stack of failed uploads of another company's PDF, a failed
    PDF of a third company, a user-deleted book that is the source of an
    old period, and an empty placeholder month;
  * a "Q&A" workspace with no company of its own: one period per foreign
    company, same-month periods colliding with the company workspace's
    served ones, a 2025 book filed under 2017, duplicates with and without
    a content hash, failed copies, a delegation itinerary PDF (not a
    balance), trash, two empty months with the SAME date, and the owner's
    conversation (with its messages) grounded in one of its periods;
  * a second user with a clean one-company workspace;
  * a two-member team workspace (never migrated).

``FakeSupabase`` answers the PostgREST / Storage requests the scripts make,
over ``httpx.MockTransport``. It pages (``max_rows`` cap, like
``db-max-rows``), refuses unknown columns and tables, records every write,
and REFUSES every DELETE — the migration and the restore must never issue
one.
"""
from __future__ import annotations

import copy
import io
import json
import re
from typing import Any, Dict, List, Optional, Tuple

import httpx
import openpyxl

from engine.workspaces.company_identity import cui_control_digit

NOW0 = "2026-09-21T10:00:00+00:00"


def valid_cui(body: str) -> str:
    return body + str(cui_control_digit(body))


ALFA = valid_cui("2000001")      # the served food company
BETA = valid_cui("3000002")      # real estate (the failed-PDF company)
GAMMA = valid_cui("4000003")     # agro (the failed PDF + an analysed xlsx)
DELTA = valid_cui("5000004")     # the sister development company (no period)
SOLO = valid_cui("6000005")      # the second user's company
TEAM = valid_cui("7000006")
SIGMA = valid_cui("8000007")     # a failed PDF whose layout the reader refuses

OWNER = "user-owner"
SOLO_USER = "user-solo"
TEAM_A, TEAM_B = "user-team-a", "user-team-b"


# ── files ──────────────────────────────────────────────────────────────

def balance_xlsx(header: List[str], *, sheet: str = "Sheet1", seed: int = 0, n: int = 14) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for line in header:
        ws.append([line])
    ws.append(["Cont", "Denumire", "SI D", "SI C", "RL D", "RL C", "SF D", "SF C"])
    for i in range(n):
        v = 1000 + seed * 10 + i
        ws.append([str(1011 + i * 7), "Cont %d" % i, 0, v, 0, 0, 0, v])
    # A customer analytic account naming ANOTHER company: never an identity.
    ws.append(["4111.00001", "Clienti OMEGA TRADING SRL", 500, 0, 0, 0, 500, 0])
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def text_pdf(lines: List[str]) -> bytes:
    """A one-page PDF with a real text layer (pdfplumber reads it)."""
    def esc(s: str) -> str:
        return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    body = "BT /F1 9 Tf 12 TL 40 800 Td " + " ".join("(%s) '" % esc(l) for l in lines) + " ET"
    objs = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources "
        "<< /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
        "<< /Length %d >>\nstream\n%s\nendstream" % (len(body.encode("latin-1")), body),
    ]
    out = b"%PDF-1.4\n"
    offs = []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += ("%d 0 obj\n%s\nendobj\n" % (i, o)).encode("latin-1")
    xref = len(out)
    out += ("xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)).encode()
    for o in offs:
        out += ("%010d 00000 n \n" % o).encode()
    out += ("trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)).encode()
    return out


def balance_pdf(header: List[str], *, seed: int = 0, n: int = 14) -> bytes:
    lines = list(header) + ["Solduri initiale an Rulaje perioada Sume totale Solduri finale",
                            "Cont Denumirea contului"]
    for i in range(n):
        v = 2000 + seed + i
        lines.append("%d Cont %d 0.00 %d.00 0.00 0.00 0.00 %d.00 0.00 %d.00" % (1011 + i * 7, i, v, v, v))
    return text_pdf(lines)


def refused_layout_pdf(header: List[str], *, seed: int = 0, n: int = 14) -> bytes:
    """A balanță whose rows the account-row reader cannot parse: every row
    starts with a label, not the account code, so no line matches the
    reader's row shape — readable text, zero account rows (the 2026-09-26
    production shape: the deterministic readers refused the layout and the
    Claude fallback had no credit, so the upload failed)."""
    lines = list(header) + ["Simbol cont | Denumire cont | Sold initial | Rulaj debitor | Rulaj creditor | Sold final"]
    for i in range(n):
        v = 3000 + seed + i
        lines.append("Cont %d | Cont %d | %d,00 | 0,00 | 0,00 | %d,00" % (1011 + i * 7, i, v, v))
    return text_pdf(lines)


def itinerary_pdf() -> bytes:
    return text_pdf([
        "OFFICIAL PROGRAMME",
        "Partner Delegation - Romania Visit",
        "9 - 15 September 2026",
        "Hosted and coordinated by Alfa Food. A programme of institutional meetings,",
        "business discussions, industrial site visits and hospitality.",
        "ALFA FOOD S.R.L",
    ])


# ── the world ──────────────────────────────────────────────────────────

def _ts(day: str, hh: int = 10) -> str:
    return "%sT%02d:00:00+00:00" % (day, hh)


def build_world() -> Tuple[Dict[str, List[Dict[str, Any]]], Dict[str, bytes], List[Dict[str, Any]]]:
    """(tables, storage objects by path, operator-verified identity rules)."""
    import hashlib

    blobs = {
        "ALFA24": balance_xlsx(["Alfa Food SRL", "Sibiu, Str. Podului 1   Balanta de Verificare - Decembrie 2024",
                                "Cod fiscal: %s" % ALFA], sheet="Document_CH14", seed=1),
        "ALFA25A": balance_xlsx([], seed=2),                       # no header: the xls-like export
        "ALFA25B": balance_xlsx([], seed=3),                       # the LV export (served)
        "ALFA25C": balance_xlsx([], seed=4),                       # "alfa trial balance 2025"
        "DELTA": balance_xlsx([], sheet="Document_CH14", seed=5),
        "BETA": balance_pdf(["BETA IMOBILIARE SRL c.f. %s r.c. J40/1/2004 Capital social 45200" % BETA,
                             "BUCURESTI sect. 1", "Balanta de verificare", "01.12.2025 -- 31.12.2025"], seed=6),
        "GAMMA_PDF": balance_pdf(["Balanta analitica", "Societate: GAMMA AGRO SRL",
                                  "Adresa: Str. Biruintei 25 C Decembrie 2025", "C.U.I: RO%s" % GAMMA], seed=7),
        "GAMMA_XLSX": balance_xlsx([], sheet="Gamma Agro", seed=8),
        "CARNEX": balance_xlsx([], sheet="Carnex", seed=9),
        "OMEGA": balance_xlsx([], sheet="Omega Retail", seed=10),
        "ITIN": itinerary_pdf(),
        "SOLO": balance_xlsx(["Solo Services SRL", "Cod fiscal: %s" % SOLO,
                              "Balanta de verificare la data de 31.12.2025"], seed=11),
        "TEAM": balance_xlsx(["Team Co SRL", "Cod fiscal: %s" % TEAM,
                              "Balanta de verificare la data de 31.12.2025"], seed=12),
        "SIGMA_PDF": refused_layout_pdf(["Balanta de verificare Decembrie 2020", "SIGMA MOTORS SRL",
                                         "CUI: RO%s" % SIGMA, "Sibiu, str. Uzinei 4"], seed=13),
    }
    sha = {k: hashlib.sha256(v).hexdigest() for k, v in blobs.items()}

    orgs = [
        {"id": "org-sf", "name": "alfa food", "industry_key": "fmcg",
         "industry_display_name": "FMCG", "default_currency": "RON", "caen_code": None,
         "caen_code_source": None, "caen_code_confirmed_at": None, "archived_at": None,
         "purge_after": None, "created_at": _ts("2026-09-09"), "updated_at": _ts("2026-09-09")},
        {"id": "org-qa", "name": "Q&A", "industry_key": "fmcg", "industry_display_name": "FMCG",
         "default_currency": "RON", "caen_code": None, "caen_code_source": None,
         "caen_code_confirmed_at": None, "archived_at": None, "purge_after": None,
         "created_at": _ts("2026-05-30"), "updated_at": _ts("2026-09-09")},
        {"id": "org-solo", "name": "solo", "industry_key": None, "industry_display_name": None,
         "default_currency": "RON", "caen_code": None, "caen_code_source": None,
         "caen_code_confirmed_at": None, "archived_at": None, "purge_after": None,
         "created_at": _ts("2026-08-01"), "updated_at": _ts("2026-08-01")},
        {"id": "org-team", "name": "team", "industry_key": None, "industry_display_name": None,
         "default_currency": "RON", "caen_code": None, "caen_code_source": None,
         "caen_code_confirmed_at": None, "archived_at": None, "purge_after": None,
         "created_at": _ts("2026-08-02"), "updated_at": _ts("2026-08-02")},
    ]
    memberships = [
        {"org_id": "org-sf", "user_id": OWNER, "role": "owner", "created_at": _ts("2026-09-09")},
        {"org_id": "org-qa", "user_id": OWNER, "role": "owner", "created_at": _ts("2026-05-30")},
        {"org_id": "org-solo", "user_id": SOLO_USER, "role": "owner", "created_at": _ts("2026-08-01")},
        {"org_id": "org-team", "user_id": TEAM_A, "role": "owner", "created_at": _ts("2026-08-02")},
        {"org_id": "org-team", "user_id": TEAM_B, "role": "member", "created_at": _ts("2026-08-02")},
    ]

    docs: List[Dict[str, Any]] = []
    storage: Dict[str, bytes] = {}

    def doc(did, org, filename, blob, *, status="analyzed", period=None, created="2026-09-10",
            deleted=None, hashed=True, hint=None, stored=True, scope="financial", hour=10, error=None):
        ext = filename.rsplit(".", 1)[-1].lower()
        path = "%s/uploads/%s.%s" % (org, did, ext)
        docs.append({
            "id": did, "org_id": org, "original_filename": filename, "display_name": None,
            "status": status, "period_id": period, "content_hash": sha[blob] if hashed else None,
            "size_bytes": len(blobs[blob]), "mime_type": "application/pdf" if ext == "pdf" else
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "created_at": _ts(created, hour), "deleted_at": deleted, "is_active": True,
            "metered_extra": False, "period_end_hint": hint,
            "uploaded_by": OWNER if org in ("org-sf", "org-qa") else
            (SOLO_USER if org == "org-solo" else TEAM_A),
            "scope": scope, "error": error, "storage_path": path, "updated_at": _ts(created, hour)})
        if stored:
            storage[path] = blobs[blob]

    # alfa food — the served company workspace
    doc("d-sf24-src", "org-sf", "Balanta decembrie 2024_extern .xlsx", "ALFA24", period="per-sf24", created="2026-09-20")
    doc("d-sf24-copy", "org-sf", "Balanta Alfa Food_FY2024.xlsx", "ALFA24", period="per-sf24", created="2026-09-09")
    doc("d-sf25-xls", "org-sf", "Balanta Alfa Food_FY2025.xlsx", "ALFA25A", period="per-sf25", created="2026-09-20")
    doc("d-sf25-lv", "org-sf", "Balanta Alfa Food_31.12.2025 LV.xlsx", "ALFA25B", period="per-sf25",
        created="2026-09-21", hour=13)
    doc("d-delta-1", "org-sf", "Trial_Balance_Alfa_Dev_31.12.2025.xlsx", "DELTA", period="per-sf25",
        created="2026-09-21", hour=7)
    doc("d-beta-fail-1", "org-sf", "balanta verificare BETA dec 2025.pdf", "BETA", status="failed",
        created="2026-09-20", error="HTTPException: 502: extraction failed")
    doc("d-beta-fail-2", "org-sf", "balanta verificare BETA dec 2025.pdf", "BETA", status="failed",
        created="2026-09-21", error="HTTPException: 502: extraction failed", stored=False)
    # a failed PDF of GAMMA whose bytes are NOT the analysed xlsx's: left for a retry
    doc("d-gamma-pdf", "org-sf", "Balanta GAMMA_FY2025.pdf", "GAMMA_PDF", status="failed",
        created="2026-09-20", error="HTTPException: 502: extraction failed")
    # a failed balanță whose row layout the reader refuses; its header
    # prints the company and its CUI (the 2026-09-26 production shape)
    doc("d-sigma-pdf", "org-sf", "Balanta Decembrie 2020 - Sigma Motors.pdf", "SIGMA_PDF", status="failed",
        created="2026-09-21", hour=15, error="HTTPException: 502: Claude extraction failed")
    doc("d-omega-trash", "org-sf", "Omega Retail Trial Balance.xlsx", "OMEGA", period="per-sf21",
        created="2026-09-20", deleted=_ts("2026-09-20", 18))

    # Q&A — no company of its own
    doc("q-carnex-src", "org-qa", "Carnex Trial Balance_FY2025.xlsx", "CARNEX", period="per-carnex",
        created="2026-09-20", hint="2017-12-31")
    doc("q-carnex-old", "org-qa", "Carnex Trial Balance 2025.xlsx", "CARNEX", created="2026-05-30", hashed=False)
    # a failed copy of the SAME bytes as the analysed source (hashed at
    # upload), whose object never landed: a duplicate, archived
    doc("q-carnex-fail", "org-qa", "Carnex Trial Balance 2025.xlsx", "CARNEX", status="failed",
        created="2026-05-30", stored=False, error="Pipeline never started")
    doc("q-sf25-src", "org-qa", "alfa trial balance 2025.xlsx", "ALFA25C", period="per-q25", created="2026-09-08")
    doc("q-delta-2", "org-qa", "Trial_Balance_Alfa_Dev_31.12.2025.xlsx", "DELTA", period="per-q25", created="2026-09-06")
    doc("q-gamma", "org-qa", "Gamma Trial Balance 2025.xlsx", "GAMMA_XLSX", period="per-q25", created="2026-09-07")
    doc("q-beta-src", "org-qa", "balanta verificare BETA dec 2025.pdf", "BETA", period="per-beta",
        created="2026-09-21", hour=14)
    doc("q-beta-dup", "org-qa", "balanta verificare BETA dec 2025.pdf", "BETA", period="per-beta",
        created="2026-09-21", hour=12)
    doc("q-sf24-src", "org-qa", "Balanta Alfa Food_FY2024.xlsx", "ALFA24", period="per-q24", created="2026-09-09")
    doc("q-itin", "org-qa", "Alfa_Delegation_Itinerary.pdf", "ITIN", status="failed", created="2026-09-11",
        error="HTTPException: 502: extraction failed")
    doc("q-gamma-trash", "org-qa", "Gamma Trial Balance 2025.xlsx", "GAMMA_XLSX", created="2026-08-02",
        deleted=_ts("2026-08-04"))
    doc("q-sku-trash", "org-qa", "Trial_Balance_Alfa_Dev_31.12.2025.xlsx", "DELTA", created="2026-08-02",
        deleted=_ts("2026-08-04"), scope="sku")

    # second user, one company
    doc("s-src", "org-solo", "balanta solo dec 2025.xlsx", "SOLO", period="per-solo", created="2026-08-01")
    # team workspace
    doc("t-src", "org-team", "balanta team 2025.xlsx", "TEAM", period="per-team", created="2026-08-02")

    def period(pid, org, end, source, *, start=None, created="2026-09-10", canonical=None, envelope=None):
        return {"id": pid, "org_id": org, "period_start": start or end, "period_end": end,
                "source_document_id": source, "currency": "RON", "assembled_canonical_v1": canonical,
                "detection_envelope": envelope, "extraction_confidence": None,
                "methodology_version": envelope.get("methodology_version") if envelope else None,
                "created_at": _ts(created), "updated_at": _ts(created)}

    # The §7 detection envelope stage_persist wrote for the 2017-filed
    # period (production's shape, 2026-09-21: its dates are the day of the
    # last re-analysis, not the row's — inconsistent before any migration)
    carnex_envelope = {
        "detection_envelope_version": "1.0.0",
        "country": {"iso2": "", "confidence": 0.0, "evidence": []},
        "standard": {"code": "", "confidence": 0.0, "evidence": []},
        "doc_type": {"code": "", "confidence": 0.0, "evidence": []},
        "industry": {"key": "fmcg", "caen_code": None, "confidence": 1.0, "evidence": ["operator_assigned"]},
        "currency": "RON", "fiscal_year_end": "2026-09-20", "period_start": "2026-09-20",
        "period_end": "2026-09-20", "methodology_version": "ro_ras_2025_v1", "routing_decision": None,
        "source_data_quality": {"raw_imbalance_pct": 0.0, "raw_imbalance_abs": 0.0, "warn": False},
    }

    periods = [
        period("per-sf24", "org-sf", "2024-12-31", "d-sf24-src", canonical={"v": 1}),
        period("per-sf25", "org-sf", "2025-12-31", "d-sf25-lv", canonical={"v": 1}),
        period("per-sf21", "org-sf", "2021-12-31", "d-omega-trash"),
        period("per-sf-empty", "org-sf", "2026-09-30", None, created="2026-09-21"),
        # filed under 2017 by a user-confirmed hint; the engine's own
        # detection record says so (stage_persist's shape, verbatim keys)
        period("per-carnex", "org-qa", "2017-12-31", "q-carnex-src", envelope=carnex_envelope, canonical={
            "canonical_bs": {"v": 1},
            "period_detection": {
                "hint": "2017-12-31", "mismatch": True, "confidence": 1.0,
                "signal_used": "user_confirmed",
                "evidence_snippet": "user-confirmed period end: 2017-12-31",
                "resolved_period_end": "2017-12-31",
                "detected": {"candidates": [{"signal": "filename", "period_end": "2025-12-31",
                                             "evidence_snippet": "Carnex Trial Balance_FY2025.xlsx"}],
                             "confidence": 0.6, "signal_used": "filename",
                             "evidence_snippet": "Carnex Trial Balance_FY2025.xlsx",
                             "proposed_period_end": "2025-12-31"}}}),
        period("per-q25", "org-qa", "2025-12-31", "q-sf25-src"),
        period("per-beta", "org-qa", "2025-12-31", "q-beta-src"),
        period("per-q24", "org-qa", "2024-12-31", "q-sf24-src"),
        period("per-qa-empty-a", "org-qa", "2026-09-30", None, created="2026-09-09"),
        period("per-qa-empty-b", "org-qa", "2026-09-30", None, created="2026-09-21"),
        period("per-solo", "org-solo", "2025-12-31", "s-src"),
        period("per-team", "org-team", "2025-12-31", "t-src"),
    ]

    def metric(i, pid, org):
        return {"id": "cm-%d" % i, "org_id": org, "period_id": pid, "name": "ebitda_margin", "value": 0.1,
                "unit": "ratio", "direction": None, "industry_p25": None, "industry_p50": None,
                "industry_p75": None, "percentile_band": None, "severity_vs_industry": None,
                "computed_at": _ts("2026-09-10")}

    metrics = [metric(i, pid, org) for i, (pid, org) in enumerate([
        ("per-sf24", "org-sf"), ("per-sf25", "org-sf"), ("per-q25", "org-qa"), ("per-beta", "org-qa"),
        ("per-q24", "org-qa"), ("per-solo", "org-solo"), ("per-team", "org-team")])]

    tables: Dict[str, List[Dict[str, Any]]] = {
        "organizations": orgs,
        "memberships": memberships,
        "org_prefs": [{"org_id": "org-sf", "prefs": {"display_currency": "RON"}, "updated_at": _ts("2026-09-09")}],
        "user_prefs": [{"user_id": OWNER, "active_org_id": "org-qa", "prefs": {"theme": "dark"},
                        "updated_at": _ts("2026-09-21")},
                       {"user_id": SOLO_USER, "active_org_id": "org-solo", "prefs": {},
                        "updated_at": _ts("2026-08-01")}],
        "financial_periods": periods,
        "documents": docs,
        "calculated_metrics": metrics,
        "statement_line_items": [{"id": "sli-1", "period_id": "per-carnex", "statement": "BS", "bucket": "cash",
                                  "amount": 10.0, "is_derived": False, "ro_account_code": "5121",
                                  "ro_account_name": "Conturi la banci"}],
        "alerts": [
            {"id": "al-1", "org_id": "org-qa", "period_id": "per-q25", "document_id": "q-sf25-src",
             "alert_key": "cash:per-q25", "title": "t", "body": "b", "category": "c", "severity": "low",
             "payload": {}, "resolved_at": None, "created_at": _ts("2026-09-10"), "updated_at": _ts("2026-09-10")},
            {"id": "al-2", "org_id": "org-qa", "period_id": "per-beta", "document_id": "q-beta-src",
             "alert_key": "cash:per-beta", "title": "t", "body": "b", "category": "c", "severity": "low",
             "payload": {}, "resolved_at": None, "created_at": _ts("2026-09-10"), "updated_at": _ts("2026-09-10")},
        ],
        "alert_states": [{"alert_id": "al-2", "org_id": "org-qa", "status": "new", "owner": None, "notes": None,
                          "acknowledged_at": None, "resolved_at": None, "updated_at": _ts("2026-09-10"),
                          "updated_by": None}],
        "briefings": [{"id": "br-1", "org_id": "org-qa", "period_id": "per-beta", "body": "x", "language": "ro",
                       "model": "m", "created_at": _ts("2026-09-10")}],
        "company_industry_assignments": [{"id": "cia-1", "organization_id": "org-qa", "period_id": "per-beta",
                                          "caen_code": None, "company_name": None, "confidence": None,
                                          "detected_industry_key": None, "locked_by_user": False,
                                          "selected_industry_key": None, "source": "auto",
                                          "created_at": _ts("2026-09-10"), "updated_at": _ts("2026-09-10")}],
        "user_valuation_assumptions": [{"user_id": OWNER, "period_id": "per-beta", "multiple_used": 8,
                                        "ebitda_used": None, "cash_used": None, "debt_used": None, "notes": None,
                                        "updated_at": _ts("2026-09-10")}],
        "recommendations": [],
        "valuations": [],
        "chat_threads": [{"id": "ct-1", "org_id": "org-qa", "user_id": OWNER, "title": "q",
                          "active_period_id": "per-q25", "active_period_label": "Dec 2025",
                          "created_at": _ts("2026-09-10"), "updated_at": _ts("2026-09-10")}],
        # no workspace column: a message belongs to its conversation
        "chat_messages": [{"id": "cm-q-1", "thread_id": "ct-1", "role": "user", "content": "q",
                           "grounded_period": "Dec 2025", "created_at": _ts("2026-09-10")},
                          {"id": "cm-q-2", "thread_id": "ct-1", "role": "assistant", "content": "a",
                           "grounded_period": "Dec 2025", "created_at": _ts("2026-09-10", 11)}],
        "billing_events": [{"id": "be-1", "org_id": "org-sf", "event_type": "x", "payload": {},
                            "stripe_event_id": None, "created_at": _ts("2026-09-10")}],
        "subscriptions": [{"id": "sub-1", "user_id": OWNER, "plan": "professional", "status": "trial",
                           "stripe_customer_id": None, "updated_at": _ts("2026-07-31")}],
        "user_usage": [{"id": "uu-1", "user_id": OWNER, "month": "2026-09", "uploads": 40,
                        "updated_at": _ts("2026-09-21")}],
    }
    rules = [
        {"user_id": OWNER, "filename_glob": "Balanta Alfa Food_*", "cui": ALFA,
         "company_name": "ALFA FOOD SRL", "evidence": "filing match (synthetic)"},
        {"user_id": OWNER, "filename_glob": "alfa trial balance 2025.xlsx", "cui": ALFA,
         "company_name": "ALFA FOOD SRL", "evidence": "filing match (synthetic)"},
        {"user_id": OWNER, "filename_glob": "Trial_Balance_Alfa_Dev_*", "cui": DELTA,
         "company_name": "DELTA DEVELOPMENT SRL", "evidence": "filing match (synthetic)"},
        {"user_id": OWNER, "filename_glob": "Gamma Trial Balance*", "cui": GAMMA,
         "company_name": "GAMMA AGRO SRL", "evidence": "filing match (synthetic)"},
        {"user_id": OWNER, "filename_glob": "Carnex Trial Balance*", "cui": None,
         "company_name": "Carnex", "evidence": "no filing matches (synthetic)"},
        # a rule for ANOTHER user never applies to the owner's files
        {"user_id": SOLO_USER, "filename_glob": "*", "cui": DELTA, "company_name": "WRONG"},
    ]
    return tables, storage, rules


def plant_second_user_move(tables, storage):
    """An ANALYSED book of BETA (no period of its own) in the SECOND user's
    one-company workspace: its bytes name another company, so the plan
    creates BETA's workspace for that user and MOVES the document (chosen
    for re-analysis) — the second user's block of the plan then carries a
    storage copy of its own. Returns the document id. (Every other test
    sees the world without it. Analysed, not failed: a failed upload that
    is no copy of a live document stays where it is, rule 5.)"""
    src = next(d for d in tables["documents"] if d["id"] == "q-beta-src")
    did = "s-beta-book"
    doc = dict(src, id=did, org_id="org-solo", status="analyzed", period_id=None, uploaded_by=SOLO_USER,
               storage_path="org-solo/uploads/%s.pdf" % did, error=None,
               created_at=_ts("2026-08-05"), updated_at=_ts("2026-08-05"))
    tables["documents"].append(doc)
    storage[doc["storage_path"]] = storage[src["storage_path"]]
    return did


def inventory_for(tables, storage):
    """What ``db_snapshot`` records: per document, whether its object
    resolved when the snapshot was taken."""
    return {str(d["id"]): {"path": d["storage_path"], "org_id": d["org_id"],
                           "exists": d["storage_path"] in storage}
            for d in tables.get("documents") or [] if d.get("storage_path")}


def facts_for(tables, storage, rules, registry=None, objects="from_storage"):
    from engine.workspaces.migration_plan import facts_from_documents

    def fetch(d):
        path = d.get("storage_path")
        if path in storage:
            return storage[path], True, None
        return None, False, "storage object missing"

    if objects == "from_storage":
        objects = inventory_for(tables, storage)
    return facts_from_documents(tables, fetch, registry=registry, rules=rules, objects=objects)


# ── a PostgREST + Storage double ───────────────────────────────────────

EXTRA_COLUMNS = {
    "recommendations": ["id", "org_id", "period_id", "title", "status", "created_at", "updated_at"],
    "valuations": ["id", "org_id", "period_id", "created_at"],
    "countries": ["code", "display_name"],
}
PKS = {"memberships": ["org_id", "user_id"], "org_prefs": ["org_id"], "user_prefs": ["user_id"],
       "alert_states": ["alert_id"], "user_valuation_assumptions": ["user_id", "period_id"],
       "countries": ["code"]}


class FakeSupabase:
    URL = "http://supabase.test"

    def __init__(self, tables: Dict[str, List[Dict[str, Any]]], storage: Dict[str, bytes], *,
                 max_rows: int = 5) -> None:
        self.tables = copy.deepcopy(tables)
        self.tables.setdefault("countries", [{"code": "RO", "display_name": "Romania"}])
        self.columns = {}
        for t, rows in self.tables.items():
            cols = set(EXTRA_COLUMNS.get(t, []))
            for r in rows:
                cols |= set(r)
            self.columns[t] = sorted(cols)
        self.objects = dict(("documents/" + k, v) for k, v in storage.items())
        self.max_rows = max_rows
        self.writes: List[Tuple[str, str, Any]] = []
        self.deletes: List[str] = []
        self.clock = "2026-09-21T12:00:00+00:00"
        #: Functions the OpenAPI document lists under /rpc/ (read, never
        #: called — the double serves no RPC): the two hold-guard markers.
        self.rpcs = {"workspace_hold_guard_version", "workspace_archive_hold_guard_version"}

    def pk(self, t):
        return PKS.get(t, ["id"])

    # ── transport ─────────────────────────────────────────────────────
    def handle(self, req: httpx.Request) -> httpx.Response:
        path = req.url.path
        if req.method == "DELETE":
            self.deletes.append(str(req.url))
            return httpx.Response(405, json={"message": "DELETE refused by the test double"})
        if path.startswith("/rest/v1"):
            table = path[len("/rest/v1"):].strip("/")
            if not table:
                return self._openapi()
            if table not in self.tables:
                return httpx.Response(404, json={"code": "PGRST205", "message": "no table %s" % table})
            return self._rest(req, table)
        if path.startswith("/storage/v1/object/sign/"):
            bucket_path = path[len("/storage/v1/object/sign/"):]
            if req.method == "POST":
                if bucket_path not in self.objects:
                    return httpx.Response(400, json={"statusCode": "404", "error": "not_found"})
                return httpx.Response(200, json={"signedURL": "/object/sign/%s?token=t" % bucket_path})
            if bucket_path not in self.objects:
                return httpx.Response(400, json={"statusCode": "404"})
            return httpx.Response(200, content=self.objects[bucket_path])
        if path.startswith("/storage/v1/object/") and req.method == "POST":
            bucket_path = path[len("/storage/v1/object/"):]
            if bucket_path in self.objects:
                return httpx.Response(400, json={"statusCode": "409", "error": "Duplicate"})
            self.objects[bucket_path] = req.content
            self.writes.append(("upload", bucket_path, len(req.content)))
            return httpx.Response(200, json={"Key": bucket_path})
        return httpx.Response(404, json={"message": "unrouted %s %s" % (req.method, path)})

    def _openapi(self) -> httpx.Response:
        defs = {}
        for t, cols in self.columns.items():
            defs[t] = {"properties": {c: {"description": "Note:\nThis is a Primary Key.<pk/>"
                                          if c in self.pk(t) else "", "type": "string"} for c in cols}}
        paths = {"/rpc/%s" % name: {"post": {}} for name in sorted(self.rpcs)}
        return httpx.Response(200, json={"definitions": defs, "paths": paths})

    def _filters(self, req: httpx.Request, table: str):
        out = []
        for k, v in req.url.params.multi_items():
            if k in ("select", "order", "limit", "offset", "on_conflict"):
                continue
            if k not in self.columns[table]:
                return None, httpx.Response(400, json={"code": "42703", "message": "column %s.%s does not exist"
                                                                                    % (table, k)})
            out.append((k, v))
        return out, None

    @staticmethod
    def _match(row, filters) -> bool:
        for k, v in filters:
            val = row.get(k)
            if v == "is.null":
                if val is not None:
                    return False
            elif v.startswith("eq."):
                if val is None or str(val) != v[3:]:
                    return False
            elif v.startswith("in.("):
                items = [x.strip().strip('"') for x in v[4:-1].split(",")]
                if str(val) not in items:
                    return False
            else:
                raise AssertionError("operator not modelled: %s=%s" % (k, v))
        return True

    def _rest(self, req: httpx.Request, table: str) -> httpx.Response:
        filters, err = self._filters(req, table)
        if err is not None:
            return err
        rows = self.tables[table]
        prefer = req.headers.get("prefer", "")
        if req.method == "GET":
            hit = [r for r in rows if self._match(r, filters)]
            order = req.url.params.get("order")
            if order:
                for part in reversed(order.split(",")):
                    col, _, direction = part.partition(".")
                    hit.sort(key=lambda r: "" if r.get(col) is None else str(r.get(col)),
                             reverse=direction == "desc")
            total = len(hit)
            off = int(req.url.params.get("offset") or 0)
            lim = int(req.url.params.get("limit") or 10 ** 9)
            page = hit[off:off + min(lim, self.max_rows)]
            headers = {}
            if "count=exact" in prefer:
                headers["content-range"] = "%d-%d/%d" % (off, off + max(len(page) - 1, 0), total)
            return httpx.Response(200, json=copy.deepcopy(page), headers=headers)
        body = json.loads(req.content or b"null")
        if req.method == "PATCH":
            bad = [c for c in body if c not in self.columns[table]]
            if bad:
                return httpx.Response(400, json={"code": "PGRST204", "message": "unknown columns %s" % bad})
            hit = [r for r in rows if self._match(r, filters)]
            for r in hit:
                r.update(copy.deepcopy(body))
                if "updated_at" in self.columns[table]:
                    r["updated_at"] = self.clock
            self.writes.append(("patch", table, {"filters": filters, "set": body}))
            return httpx.Response(200, json=copy.deepcopy(hit) if "return=representation" in prefer else [])
        if req.method == "POST":
            items = body if isinstance(body, list) else [body]
            pk = self.pk(table)
            upsert = "merge-duplicates" in prefer
            for item in items:
                bad = [c for c in item if c not in self.columns[table]]
                if bad:
                    return httpx.Response(400, json={"code": "PGRST204", "message": "unknown columns %s" % bad})
                key = [str(item.get(c)) for c in pk]
                existing = [r for r in rows if [str(r.get(c)) for c in pk] == key]
                if existing and not upsert:
                    return httpx.Response(409, json={"code": "23505", "message": "duplicate key"})
                if existing:
                    existing[0].update(copy.deepcopy(item))
                else:
                    row = {c: None for c in self.columns[table]}
                    if "created_at" in row:
                        row["created_at"] = self.clock
                    if "updated_at" in row:
                        row["updated_at"] = self.clock
                    row.update(copy.deepcopy(item))
                    rows.append(row)
                self.writes.append(("upsert" if upsert else "insert", table, item))
            return httpx.Response(201, json=[])
        return httpx.Response(405, json={"message": "method %s" % req.method})

    def client(self):
        from engine.api._supabase import SupabaseClient

        c = SupabaseClient(self.URL, "service-role-test-key")
        c._client.close()
        c._client = httpx.Client(transport=httpx.MockTransport(self.handle), headers=c._headers, timeout=5)
        return c

    def row_writes(self):
        return [w for w in self.writes if w[0] != "upload"]
