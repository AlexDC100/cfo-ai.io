"""Books for the inventory-days gates (owner spec 2026-09-26, inventory days; design
B1-B5): `inventory-split-reconciles`, `inventory-basis-label`,
`stock-claim-policy` and `benchmark-basis-separation-engine`.

Every book here is COMMITTED corpus data (corpus/*/input.*, anonymised) or a
variant of it written to the test's own tmp directory — never a client book,
never a new committed fixture. Each goes through the REAL production write
path (parse -> `_deterministic_tb_parsed`, which measures the
`inventory_stock/1` evidence -> `stage_map` -> `stage_persist`) and is read
back through the REAL router (GET /api/period over the projection-faithful
double of `test_rebuild_net_income_anchor`), exactly as a served period is.

Variants (tmp only):
  * `predates(bk)` — the same period with the stored evidence removed: the
    state of EVERY period written before the measurement (production today,
    until `scripts/reprocess_periods_definition.py` runs). The block serves
    the period-end snapshot from the line items.
  * `with_stock_moved(case, tmp, delta)` — the corpus workbook with `delta`
    added to one materials account's period-end stock and the same amount to
    one supplier account (bought on credit): the trial balance still
    balances, the movement convention still holds, the P&L does not move —
    only the stock, and so the inventory days.
"""
from __future__ import annotations

import copy
import hashlib
import sys
import types
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

HERE = Path(__file__).resolve().parent
if str(HERE) not in sys.path:
    sys.path.insert(0, str(HERE))

import test_insights_wire as W  # noqa: E402
import test_rebuild_net_income_anchor as A  # noqa: E402
from _pytest.monkeypatch import MonkeyPatch  # noqa: E402

REPO = HERE.parents[1]
CORPUS = REPO / "corpus"

#: Every corpus case whose trial balance carries class-3 stock (measured
#: 2026-09-27 through the real write path: ten books, 185 stock accounts).
#: The five without stock lines (contra_sign_flip, exact_zero,
#: imbalance_03pct, rounding_004pct, unmapped_equals_delta) serve a refused
#: block (`no_stock_lines`) and are read by `test_a_book_without_stock_serves_no_days`.
STOCK_BOOKS = ("saga_10_col", "saga_10_col_agras", "saga_10_col_carniprod",
               "saga_10_col_realestate", "saga_10_col_retail", "pdf_positional",
               "saga_compact_6_col", "csv", "generic_4_col", "dup_totals_row")
NO_STOCK_BOOKS = ("contra_sign_flip", "exact_zero", "imbalance_03pct",
                  "rounding_004pct", "unmapped_equals_delta")
#: The books whose file carries the fiscal-year opening (movement convention
#: B or C): the two-year-end average.
OPENING_BOOKS = ("saga_10_col", "saga_10_col_agras", "saga_10_col_carniprod",
                 "saga_10_col_realestate", "saga_10_col_retail", "pdf_positional",
                 "saga_compact_6_col")
#: The books whose file does NOT carry it: a 4-column layout (no `si`), a
#: file whose movement columns do not say which date the `si` is, and a
#: 4-pair file where B wins only by a tie with A and no class 6/7 row dates
#: the `si` at 1 January (the csv book has no profit-and-loss account).
SNAPSHOT_BOOKS = {"generic_4_col": "si_column_absent", "dup_totals_row": "si_convention_undecided",
                  "csv": "si_date_undetermined"}

_BOOKS: Dict[str, Any] = {}
_BODIES: Dict[str, Dict[str, Any]] = {}


def corpus_book(case: str) -> Any:
    if case not in _BOOKS:
        _BOOKS[case] = A._Book(CORPUS / case)
    return _BOOKS[case]


def book_from_path(path: Path, *, key: str, period_end: str = "2025-12-31") -> Any:
    """A workbook at `path` through the real write path (a tmp variant)."""
    import _real_app_comparatives as RA

    bk = RA.book_from_workbook(path, period_end=period_end, key=key)
    return types.SimpleNamespace(period=bk.period, line_items=bk.line_items,
                                 period_id=bk.period_id,
                                 org={"id": bk.period.get("org_id") or "org-local",
                                      "name": "Variant %s" % key})


def with_org(bk: Any, **org: Any) -> Any:
    """The same persisted period, served to a workspace whose organization
    row carries `org` (e.g. a CAEN code)."""
    return types.SimpleNamespace(period=bk.period, line_items=bk.line_items,
                                 period_id=bk.period_id, org=dict(bk.org, **org))


def predates(bk: Any) -> Any:
    """The period as it was written before the stock evidence existed."""
    period = copy.deepcopy(bk.period)
    env = period.get("assembled_canonical_v1") or {}
    env.pop("inventory_stock", None)
    env.pop("inventory_days", None)
    period["assembled_canonical_v1"] = env
    return types.SimpleNamespace(period=period, line_items=bk.line_items,
                                 period_id=bk.period_id, org=dict(bk.org))


def served(bk: Any, cache_key: Optional[str] = None) -> Dict[str, Any]:
    """GET /api/period/{id} through the real router. `cache_key` names the
    book (period ids repeat across corpus books — never a cache key)."""
    if cache_key is not None and cache_key in _BODIES:
        return _BODIES[cache_key]
    mp = MonkeyPatch()
    try:
        status, body = W._served_via_route(bk, mp)
    finally:
        mp.undo()
    assert status == 200, (bk.period_id, status)
    if cache_key is not None:
        _BODIES[cache_key] = body
    return body


def block_of(body: Dict[str, Any]) -> Dict[str, Any]:
    block = (body.get("assembled_metrics") or {}).get("inventory_days")
    assert isinstance(block, dict) and block.get("schema") == "inventory_days/1", sorted(body)
    assert block == body["statements"]["inventory_days"], "two blocks served for one period"
    return block


def parsed_rows(case: str) -> List[Dict[str, Any]]:
    """The trial balance's own rows, parsed exactly as the write path parses
    them — the gate's INDEPENDENT reading of the stock (never the stored
    evidence the block was built from)."""
    import corpus_replay

    case_dir = CORPUS / case
    pack = corpus_replay.get_pack("RO")
    meta = corpus_replay._load_meta(case_dir)
    path = corpus_replay._input_path(case_dir)
    content = path.read_bytes()
    rows = (pack.parse_trial_balance_csv(content, path.name)
            if str(meta["expected_parser"]) == "csv"
            else pack.parse_trial_balance(content, path.name))
    return [dict(r) for r in rows if isinstance(r, dict)]


def digits(code: Any) -> str:
    return "".join(c for c in str(code or "") if c.isdigit())


def class3_leaves(rows: List[Dict[str, Any]]) -> Dict[str, Tuple[float, float]]:
    """code -> (opening si_d - si_c, closing sf_d - sf_c) for every class-3
    row the assembler reads (leaves and distinct accounts; an aggregate
    row's figures are already in its leaves)."""
    from engine.country_packs.ro_romania.stock_variation import row_reading

    ordered = sorted(rows, key=lambda r: str(r.get("cont") or "").strip())
    out: Dict[str, Tuple[float, float]] = {}
    for r, kind in zip(ordered, row_reading(ordered)):
        code = str(r.get("cont") or "").strip()
        if kind == "aggregate" or not digits(code).startswith("3"):
            continue

        def f(k: str) -> float:
            try:
                return float(r.get(k) or 0.0)
            except (TypeError, ValueError):
                return 0.0
        o, c = f("si_d") - f("si_c"), f("sf_d") - f("sf_c")
        prev = out.get(code, (0.0, 0.0))
        out[code] = (prev[0] + o, prev[1] + c)
    return out


# ── the owner's literal groups (design B1), restated HERE, not read from
# the pack, so a pack edit that moves an account is a red, not a new law ──
OWNER_STOCK = {"materials": ("301", "302", "303", "308"),
               "finished_goods_wip": ("331", "341", "345", "348"),
               "merchandise": ("371", "378")}
OWNER_PROVISIONS = {"materials": ("391", "392"),
                    "finished_goods_wip": ("393", "394"),
                    "merchandise": ("397",)}


def owner_group(code: str) -> Tuple[str, str]:
    """(part key, 'stock' | 'provision') under the ruling."""
    d = digits(code)
    for key, prefixes in OWNER_STOCK.items():
        if d.startswith(prefixes):
            return key, "stock"
    for key, prefixes in OWNER_PROVISIONS.items():
        if d.startswith(prefixes):
            return key, "provision"
    return "other_stock", "provision" if d.startswith("39") else "stock"


def canonical_inventory(statements: Dict[str, Any]) -> float:
    """Σ the served canonical_bs inventory rows (the balance sheet's stock)."""
    cbs = statements.get("canonical_bs") or (statements.get("assembled_canonical_v1") or {}).get("canonical_bs")
    rows = [r for r in (cbs or {}).get("rows") or []
            if str(r.get("id") or "").startswith("inventory_") and r.get("amount") is not None]
    assert rows, "the served canonical_bs carries no inventory row"
    return round(sum(float(r["amount"]) for r in rows), 2)


# ── tmp variants of a corpus workbook ─────────────────────────────────────


def _xlsx_rows(case: str) -> Tuple[List[Any], List[List[Any]]]:
    import openpyxl

    wb = openpyxl.load_workbook(CORPUS / case / "input.xlsx", read_only=True)
    ws = wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return rows[0], rows[1:]


def with_stock_moved(case: str, tmp: Path, delta: float, *, stock_prefix: str = "301",
                     supplier_prefix: str = "401", name: str = "variant") -> Path:
    """The corpus workbook (10-column SAGA layout: SI, RL, RC, SF pairs) with
    `delta` of materials bought on credit and still in stock at the period
    end: the first `stock_prefix` account's cumulative debit and closing
    debit rise by `delta`, the first `supplier_prefix` account's cumulative
    credit and closing credit rise by the same. Written to `tmp` only."""
    import openpyxl

    header, rows = _xlsx_rows(case)
    si_d, rl_d, rc_d, sf_d = 2, 4, 6, 8
    moved = {"stock": False, "supplier": False}
    for r in rows:
        code = digits(r[0])
        if not moved["stock"] and code.startswith(stock_prefix) and (r[sf_d] or 0):
            r[rc_d] = float(r[rc_d] or 0) + delta
            r[sf_d] = float(r[sf_d] or 0) + delta
            moved["stock"] = True
        elif not moved["supplier"] and code.startswith(supplier_prefix) and (r[sf_d + 1] or 0):
            r[rc_d + 1] = float(r[rc_d + 1] or 0) + delta
            r[sf_d + 1] = float(r[sf_d + 1] or 0) + delta
            moved["supplier"] = True
    assert moved == {"stock": True, "supplier": True}, (case, moved)
    del si_d, rl_d
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(header)
    for r in rows:
        ws.append(r)
    out = tmp / ("%s_%s_%s.xlsx" % (case, name, hashlib.sha256(repr(delta).encode()).hexdigest()[:6]))
    wb.save(out)
    return out
