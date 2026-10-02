"""THE workspace_v2 GATE SUITE (G1–G4, G7, G8), on the REAL app and the REAL identifier.

Owner spec (2026-09-21): one company per workspace, keyed by CUI; one upload
component; a file lands in the company its own header names; the period is
the document's, never the file name's; the same file twice is not stored,
analysed or counted; no period without an analysed file; removals are
archives. Every gate below is plant-proven — the plant, its red output and
the revert are recorded in docs/engine_book/gates.md ("workspace-v2").

What is real here: `engine.api.create_app()` with every middleware and wall
on the request path; `engine.workspaces.company_identity.identify_document`
over WORKBOOK BYTES BUILT IN THE TEST (the anonymized real Agras book from
corpus/saga_10_col_agras, with a company header in front of it); the
duplicate definition (`engine.api._doc_dedupe`); the meter's HTTP shapes;
and — in G7 — the pipeline itself: `_run_pipeline_sync` with the REAL
`stage_extract` (the deterministic trial-balance path, reading the bytes the
commit stored), `stage_map`, `stage_persist`, `stage_compute`,
`stage_validate`, the council and the narrate stage, followed by the served
`GET /api/companies/{org}/years` and `GET /api/period/{id}`.

What is doubled: PostgREST + Storage (`test_workspace_uploads.WorkspaceDouble`
widened to every table the migrations declare, production-shaped: no
`organizations.cui`), the Anthropic SDK (a stand-in that fails every request
exactly as production does with an empty credit balance — the narrate stage
is non-fatal by design), the account meter (`_usage_gate` reserve / commit /
release, recorded), the daemon thread (`_enqueue` records; the test runs the
thread's body). The NETWORK IS CLOSED: any httpx request that is not the
double's own storage raises.

CUIs are the spec's: Scandia 16070576, Agras 46355095 (both carry a valid
control digit — the identifier enforces key 753217532).

WHAT EACH GATE REDS ON, with the product correct (TC-11):
  G1  an Agras file dropped while Scandia is on screen landing anywhere but
      Agras — identify's target, the commit's document row, its storage
      prefix, the answer's company — or ANY row or object appearing in
      Scandia.
  G2  a period taken from the file NAME: "balanta_2017.xlsx" whose period
      line says 31.12.2025 must be offered, stored as the hint, and
      persisted (G7) as 2025-12-31; a file named for December 2025 whose
      document says 31.12.2024 is 2024.
  G3  the second copy of a file (same bytes, account, company, period)
      being stored (a documents row or a storage object), analysed
      (enqueued, run) or counted (reserved, committed) — at identify, at
      commit, and after the first copy was really analysed.
  G4  identify or commit creating a period; an analysed run leaving any
      period without an analysed source document; `scripts/check_no_empty_periods.py`
      not finding the empty periods of a production-shaped snapshot, or
      passing vacuously.
  G7  identify -> commit (one tap) -> the five pipeline stages -> the
      analysed period served by the dashboard's own route, with the year
      tile's revenue equal to that served revenue.
  G8  the migration planning a DELETE, a restore that does not put back the
      pre-state exactly, or the migration / restore / snapshot scripts
      carrying a delete call.

Python 3.9 — no `match`, no `X | Y` unions.
"""
from __future__ import annotations

import copy
import glob
import hashlib
import importlib.util
import io
import json
import os
import re
import sys
import types
from pathlib import Path
from typing import Any, Dict, List, Optional

import httpx
import pytest
from fastapi.testclient import TestClient

import _real_app_comparatives as RA
import firm_postgrest_double as D
import test_workspace_uploads as WU
from engine.api import _supabase, _uploads, _usage_gate, pipeline

REPO = Path(__file__).resolve().parents[2]
AGRAS_BOOK = REPO / "corpus" / "saga_10_col_agras" / "input.xlsx"
SCANDIA_BOOK = REPO / "corpus" / "saga_10_col" / "input.xlsx"

USER = WU.USER
OUTSIDER = WU.OUTSIDER
ORG_SCANDIA = WU.ORG_SCANDIA
ORG_AGRAS = WU.ORG_AGRAS
CUI_SCANDIA = "16070576"
CUI_AGRAS = "46355095"
XLSX = WU.XLSX

STORAGE_HOST = "storage.double.invalid"


# ══════════════════════════════════════════════════════════════════════
# Workbooks built in the test
# ══════════════════════════════════════════════════════════════════════


def book_workbook(book: Path, *, name: str, cui: str,
                  period_line: Optional[str] = "Balanta de verificare la data de 31.12.2025") -> bytes:
    """A committed corpus trial balance (codes and every figure to the cent)
    with the company header an accounting export prints above the columns:
    the name, the fiscal code, and the period line."""
    import openpyxl

    wb = openpyxl.load_workbook(str(book))
    ws = wb[wb.sheetnames[0]]
    header = [name, "Cod fiscal: RO%s" % cui]
    if period_line:
        header.append(period_line)
    ws.insert_rows(1, amount=len(header) + 1)
    for i, line in enumerate(header, start=1):
        ws.cell(row=i, column=1, value=line)
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def agras_workbook(*, period_line: Optional[str] = "Balanta de verificare la data de 31.12.2025",
                   cui: str = CUI_AGRAS, name: str = "AGRAS SRL") -> bytes:
    """The anonymized real Agras FY2025 book (corpus/saga_10_col_agras)
    under Agras's header."""
    return book_workbook(AGRAS_BOOK, name=name, cui=cui, period_line=period_line)


def balance_pdf_of(name: str, cui: str) -> bytes:
    """A balance PDF under a company's header: the synthetic five-pair book of
    tests/engine/test_pdf_balanta_stage_extract.py (PyMuPDF, invented
    figures, "Decembrie 2025" in its title block) — a book the pipeline's
    strict text-line reader reads whole, with no model."""
    import test_pdf_balanta_stage_extract as PB

    lines = []
    for line in PB._synthetic_five_pair_lines():
        if line.startswith("Societate:"):
            line = "Societate: %s" % name
        elif line.startswith("C.U.I:"):
            line = "C.U.I: RO%s" % cui
        lines.append(line)
    return PB._pdf_bytes(lines)


#: The byte shapes of review round 4 (2026-10-02). Neither comes out of an
#: ERP or a PDF writer; both are PDFs to the pipeline and were "unreadable"
#: to the identifier, so the takeover wall had no CUI to hold.
#:  · pdf_behind_PK — "PK" + 100 bytes + the PDF: the sniffer demands the
#:    four-byte zip signature, so this is a PDF under EVERY name.
#:  · pdf_junk_both_ends — the header past the sniffer's 1,024-byte window
#:    and the trailer ahead of the last 2,048 bytes: a PDF to the pipeline
#:    exactly where the branch is pdf or xlsx, and the branch of a name with
#:    no extension is the declared MIME type's.
PDF_SHAPES = {
    "pdf": lambda pdf: pdf,
    "pdf_behind_PK": lambda pdf: b"PK" + b"J" * 100 + pdf,
    "pdf_junk_both_ends": lambda pdf: b"J" * 2000 + pdf + b"\n" + b"K" * 3000,
}
XLS_MIME = "application/vnd.ms-excel"
OCTET = "application/octet-stream"


# ══════════════════════════════════════════════════════════════════════
# The double: every table the migrations declare, production-shaped
# ══════════════════════════════════════════════════════════════════════


def _all_migration_columns() -> Dict[str, List[str]]:
    """Every `create table` + `alter table … add column` over supabase/*.sql
    — then shaped to PRODUCTION exactly as test_workspace_uploads shapes it
    (organizations has no cui / firm_id; documents carries the three columns
    that predate the migrations), plus the two tables get_period reads that
    no migration here creates (the same declaration _real_app_comparatives
    makes)."""
    cols = {}  # type: Dict[str, List[str]]
    for path in sorted(glob.glob(str(REPO / "supabase" / "*.sql"))):
        text = Path(path).read_text(encoding="utf-8")
        for table, names in D.parse_table_columns(text).items():
            have = cols.setdefault(table, [])
            have += [n for n in names if n not in have]
        bare = "\n".join(line.split("--", 1)[0] for line in text.splitlines())
        for m in RA._ALTER_RX.finditer(bare):
            have = cols.setdefault(m.group(1).lower(), [])
            have += [n for n in RA._ADD_RX.findall(m.group(2)) if n not in have]
    cols["organizations"] = [c for c in cols["organizations"] if c not in ("cui", "firm_id")]
    for col in ("deleted_at", "display_name", "is_active"):
        if col not in cols["documents"]:
            cols["documents"].append(col)
    cols.setdefault("valuations", ["id", "period_id", "org_id"])
    cols.setdefault("user_valuation_assumptions", ["user_id", "period_id"])
    return cols


class GateDouble(WU.WorkspaceDouble):
    """`WorkspaceDouble` (projection-faithful PostgREST + tenant-guarded
    Storage) over every table the migrations declare, with the two storage
    reads the pipeline makes: `signed_url` (tenant-asserted, like the real
    client) and the download of that URL (served by `storage_transport`)."""

    def __init__(self) -> None:
        super(GateDouble, self).__init__()
        self.columns.update(_all_migration_columns())
        self.upserts = []  # type: List[Any]

    def signed_url(self, bucket: str, path: str, *, org_id: str, expires_in: int = 300) -> str:
        _supabase.assert_tenant_path(bucket, path, org_id, op="sign")
        return "https://%s/%s/%s" % (STORAGE_HOST, bucket, path)

    def upsert(self, table: str, rows: Any, *, on_conflict: Optional[str] = None,
               returning: bool = True, **_: Any) -> List[Dict[str, Any]]:
        body = rows if isinstance(rows, list) else [rows]
        self.upserts.append((table, on_conflict))
        keys = [k.strip() for k in (on_conflict or "id").split(",") if k.strip()]
        out = []
        for r in body:
            hit = next((x for x in self.rows(table) if all(x.get(k) == r.get(k) for k in keys)), None)
            if hit is not None and all(r.get(k) is not None for k in keys):
                unknown = set(r) - set(self.columns[table])
                if unknown:
                    raise D.unknown_column(table, sorted(unknown)[0])
                hit.update(copy.deepcopy(r))
                out.append(copy.deepcopy(hit))
            else:
                out.append(copy.deepcopy(self.add(table, dict(r))))
        return out if returning else []


def storage_transport(db: GateDouble) -> httpx.MockTransport:
    """The only host this test's network reaches: the double's storage."""
    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.host != STORAGE_HOST:
            raise RuntimeError("network is closed in this test: %s" % request.url)
        key = request.url.path.lstrip("/")
        if key not in db.storage:
            return httpx.Response(404, json={"message": "not found"})
        return httpx.Response(200, content=db.storage[key])
    return httpx.MockTransport(handler)


class _NoCreditAnthropic(object):
    """The Anthropic SDK as production has it today: constructed fine,
    every request refused (the credit balance is empty). Nothing leaves the
    process."""

    calls = []  # type: List[str]

    def __init__(self, *a: Any, **kw: Any) -> None:
        self.messages = self

    def create(self, *a: Any, **kw: Any) -> Any:
        _NoCreditAnthropic.calls.append(str(kw.get("model")))
        raise RuntimeError("Your credit balance is too low to access the Anthropic API.")

    def stream(self, *a: Any, **kw: Any) -> Any:
        return self.create(*a, **kw)


class Meter(object):
    """`_usage_gate` recorded: what was reserved, committed (COUNTED) and
    released."""

    def __init__(self) -> None:
        self.reserved = []  # type: List[str]
        self.committed = []  # type: List[Any]
        self.released = []  # type: List[Any]

    def reserve(self, user_id: str) -> Any:
        self.reserved.append(user_id)
        return _usage_gate.DocReserveDecision(
            kind="allowed", plan_key="professional", used=3, reserved=1, cap=15,
            extra_doc_eur=3.0, message="", was_extra=False)


class GateWorld(object):
    def __init__(self, db: GateDouble, meter: Meter) -> None:
        self.db = db
        self.meter = meter
        self.enqueued = []  # type: List[str]

    def docs(self, **match: Any) -> List[Dict[str, Any]]:
        return [d for d in self.db.rows("documents") if all(d.get(k) == v for k, v in match.items())]

    def objects(self, org_id: str) -> List[str]:
        return sorted(k for k in self.db.storage if k.startswith("documents/%s/" % org_id))

    def state(self) -> str:
        return self.db.snapshot()


@pytest.fixture(scope="module")
def app():
    return RA.build_app()


@pytest.fixture()
def gw(app, monkeypatch):
    db = GateDouble()
    for org_id, name in ((ORG_SCANDIA, "Scandia Food SRL"), (ORG_AGRAS, "Agras SRL"),
                         (WU.ORG_OUTSIDE, "Outside SRL")):
        db.add("organizations", {"id": org_id, "name": name, "default_currency": "RON",
                                 "industry_key": "food_manufacturing", "archived_at": None,
                                 "created_at": "2026-01-01T00:00:00+00:00"})
    for i, org_id in enumerate((ORG_SCANDIA, ORG_AGRAS)):
        db.add("memberships", {"user_id": USER, "org_id": org_id, "role": "owner",
                               "created_at": "2026-01-0%dT00:00:00+00:00" % (i + 1)})
    db.add("memberships", {"user_id": OUTSIDER, "org_id": WU.ORG_OUTSIDE, "role": "owner",
                           "created_at": "2026-01-01T00:00:00+00:00"})
    for org_id, cui in ((ORG_SCANDIA, CUI_SCANDIA), (ORG_AGRAS, CUI_AGRAS), (WU.ORG_OUTSIDE, "31415926")):
        db.add("org_prefs", {"org_id": org_id, "prefs": {"cui": cui, "display_currency": "RON"}})
    db.add("industry_profiles", {"key": "food_manufacturing", "display_name": "Food manufacturing",
                                 "display_name_ro": "Industria alimentara"})
    meter = Meter()
    w = GateWorld(db, meter)

    for key in ("CFO_FEATURES_ACTIVE", "USAGE_LIMITS_ENABLED", "PUBLIC_TEST_MODE"):
        monkeypatch.delenv(key, raising=False)
    # The xlsx path of stage_extract asks for a key before it reads the
    # file (production has one; its balance is empty). The stand-in SDK
    # below is what any request meets.
    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-test-no-network")
    stub = types.ModuleType("anthropic")
    stub.Anthropic = _NoCreditAnthropic  # type: ignore[attr-defined]
    stub.APIStatusError = RuntimeError  # type: ignore[attr-defined]
    stub.APIError = RuntimeError  # type: ignore[attr-defined]
    monkeypatch.setitem(sys.modules, "anthropic", stub)
    _NoCreditAnthropic.calls = []

    # The network is closed: the double's storage is the only host.
    transport = storage_transport(db)
    real_client = httpx.Client

    def _client(*a: Any, **kw: Any) -> httpx.Client:
        kw.pop("transport", None)
        return real_client(*a, transport=transport, **kw)

    shim = dict((k, getattr(httpx, k)) for k in dir(httpx) if not k.startswith("__"))
    shim["Client"] = _client
    monkeypatch.setattr(pipeline, "httpx", types.SimpleNamespace(**shim))

    def _refuse(self: Any, request: httpx.Request) -> httpx.Response:
        raise RuntimeError("network is closed in this test: %s" % request.url)

    monkeypatch.setattr(httpx.HTTPTransport, "handle_request", _refuse)

    monkeypatch.setattr(_supabase, "per_user", lambda jwt: WU._UserClient(db, D.verified_identity(jwt).get("id")))
    monkeypatch.setattr(_supabase, "admin", lambda: db)
    monkeypatch.setattr(_uploads, "_open_registry", lambda: None)
    monkeypatch.setattr(pipeline, "_enqueue", lambda doc_id: w.enqueued.append(doc_id))
    monkeypatch.setattr(_usage_gate, "enforcement_enabled", lambda: True)
    monkeypatch.setattr(_usage_gate, "reserve_document", meter.reserve)
    # `month`: the settlement commits / releases in the month the
    # reservation was made in (fix/dedupe-quota 11f84259, the restart repair).
    monkeypatch.setattr(_usage_gate, "commit_document",
                        lambda uid, was_extra=False, month=None: meter.committed.append((uid, was_extra)))
    monkeypatch.setattr(_usage_gate, "release_document",
                        lambda uid, was_extra=False, month=None: meter.released.append((uid, was_extra)))
    from engine.api import _doc_dedupe
    monkeypatch.setattr(pipeline, "_QUOTA_RUNS", {})
    monkeypatch.setattr(_doc_dedupe, "_ARCHIVED_HERE", set())
    return w


def _http(app) -> TestClient:
    return TestClient(app, raise_server_exceptions=False)


def _headers(user: str = USER, org: Optional[str] = None) -> Dict[str, str]:
    out = {"Authorization": "Bearer %s" % D.mint_jwt(user)}
    if org:
        out["X-Org-Id"] = org
    return out


PDF_MIME = "application/pdf"


def identify(app, content: bytes, filename: str, *, org: Optional[str] = ORG_SCANDIA, mime: str = XLSX):
    return _http(app).post("/api/uploads/identify", headers=_headers(USER, org),
                           files={"file": (filename, content, mime)})


def commit(app, content: bytes, filename: str, *, mime: str = XLSX, **form: Any):
    data = dict((k, v if isinstance(v, str) else json.dumps(v)) for k, v in form.items() if v is not None)
    return _http(app).post("/api/uploads/commit", headers=_headers(USER),
                           files={"file": (filename, content, mime)}, data=data)


def one_tap(app, content: bytes, filename: str, *, on_screen: str = ORG_SCANDIA,
            mime: str = XLSX) -> Dict[str, Any]:
    """What the card's one primary button sends: identify, then commit the
    identity exactly as the card shows it."""
    r = identify(app, content, filename, org=on_screen, mime=mime)
    assert r.status_code == 200, r.text[:400]
    ident = r.json()
    target = ident["target"]
    form = {"period_end": ident["identity"]["period_end"], "output_language": "ro"}
    if target["org_id"]:
        form["target_org_id"] = target["org_id"]
    else:
        form["create_company"] = {"name": target["name"], "cui": ident["identity"]["cui"],
                                  "caen_code": ident["identity"]["caen_code"],
                                  "industry_key": ident["identity"]["industry_key"]}
    c = commit(app, content, filename, mime=mime, **form)
    assert c.status_code == 200, c.text[:400]
    return {"identify": ident, "commit": c.json()}


def run_analysis(gw: GateWorld, doc_id: str) -> Dict[str, Any]:
    """The daemon thread's body, for real."""
    assert gw.enqueued[-1] == doc_id, gw.enqueued
    pipeline._run_pipeline_sync(doc_id)
    (doc,) = gw.docs(id=doc_id)
    return doc


# ══════════════════════════════════════════════════════════════════════
# G1 — an Agras file dropped while Scandia is open lands in Agras
# ══════════════════════════════════════════════════════════════════════


def test_g1_an_agras_file_dropped_on_a_scandia_page_lands_in_agras(app, gw):
    content = agras_workbook()
    before_scandia = (gw.docs(org_id=ORG_SCANDIA), gw.objects(ORG_SCANDIA))
    out = one_tap(app, content, "balanta.xlsx", on_screen=ORG_SCANDIA)
    ident, done = out["identify"], out["commit"]
    assert ident["identity"]["cui"] == CUI_AGRAS, ident["identity"]
    assert ident["identity"]["sources"]["cui"]["signal"] == "document_header_cui", ident["identity"]["sources"]
    assert ident["target"] == {"org_id": ORG_AGRAS, "name": "Agras SRL", "is_new": False,
                               "reason": "cui_match"}, ident["target"]
    assert done["status"] == "queued" and done["org_id"] == ORG_AGRAS, done
    assert done["company_name"] == "Agras SRL" and done["created_company"] is False, done
    (doc,) = gw.docs(id=done["document_id"])
    assert doc["org_id"] == ORG_AGRAS and doc["uploaded_by"] == USER, doc
    assert doc["storage_path"].startswith(ORG_AGRAS + "/uploads/"), doc["storage_path"]
    assert gw.objects(ORG_AGRAS) == ["documents/" + doc["storage_path"]]
    assert (gw.docs(org_id=ORG_SCANDIA), gw.objects(ORG_SCANDIA)) == before_scandia, \
        "G1: something of the Agras file landed in Scandia"
    assert gw.enqueued == [doc["id"]] and gw.meter.reserved == [USER]


def test_g1_an_agras_workbook_named_pdf_is_identified_by_its_bytes_and_lands_in_agras(app, gw):
    """Review round 3 (2026-10-02), HIGH. Since the one upload policy the
    card's routes ADMIT a workbook named .pdf (the pipeline reads it). The
    identifier chose its reader by NAME, so the same Agras workbook named
    balanta.pdf came back "unreadable" — no CUI, no company, no period —
    the card offered the company ON SCREEN, and one tap stored and analysed
    Agras's book in Scandia, which then served Agras's turnover. (On
    production's tree the routes refused the file, so neither was reachable.)

    The identity is the bytes': named .pdf, the file is identified exactly
    as under .xlsx, lands in Agras, is analysed there, and Scandia holds
    nothing of it.

    Reds on: a workbook named .pdf identified by its name (no CUI, or the
    on-screen company as the target); anything of it stored or analysed in
    the company on screen."""
    content = agras_workbook()
    as_xlsx = identify(app, content, "balanta.xlsx", org=ORG_SCANDIA)
    assert as_xlsx.status_code == 200, as_xlsx.text[:300]
    before_scandia = (gw.docs(org_id=ORG_SCANDIA), gw.objects(ORG_SCANDIA), gw.db.rows("financial_periods"))
    out = one_tap(app, content, "balanta.pdf", on_screen=ORG_SCANDIA, mime=PDF_MIME)
    ident, done = out["identify"], out["commit"]
    assert ident["identity"]["cui"] == CUI_AGRAS, ident["identity"]
    assert ident["identity"]["sources"]["cui"]["signal"] == "document_header_cui", ident["identity"]["sources"]
    for field in ("cui", "company_name", "period_end", "document_kind", "caen_code", "industry_key"):
        assert ident["identity"][field] == as_xlsx.json()["identity"][field], \
            "the identity's %s differs by NAME: %r as .pdf, %r as .xlsx" % (
                field, ident["identity"][field], as_xlsx.json()["identity"][field])
    assert ident["identity"]["document_kind"] == "trial_balance", ident["identity"]
    assert ident["target"] == {"org_id": ORG_AGRAS, "name": "Agras SRL", "is_new": False,
                               "reason": "cui_match"}, ident["target"]
    assert done["status"] == "queued" and done["org_id"] == ORG_AGRAS, done
    doc = run_analysis(gw, done["document_id"])
    assert doc["status"] == "analyzed" and doc["org_id"] == ORG_AGRAS, (doc["status"], doc.get("error"))
    (period,) = gw.db.rows("financial_periods")
    assert period["org_id"] == ORG_AGRAS and period["source_document_id"] == doc["id"], period
    assert (gw.docs(org_id=ORG_SCANDIA), gw.objects(ORG_SCANDIA)) == before_scandia[:2], \
        "G1: something of the Agras file named .pdf landed in Scandia"
    assert [p for p in gw.db.rows("financial_periods") if p["org_id"] == ORG_SCANDIA] == \
        [p for p in before_scandia[2] if p["org_id"] == ORG_SCANDIA]


@pytest.mark.parametrize("shape,filename,mime", [
    ("pdf_behind_PK", "balanta.pdf", PDF_MIME),
    ("pdf_behind_PK", "balanta.xls", "application/vnd.ms-excel"),
    ("pdf_junk_both_ends", "balanta", PDF_MIME),
    ("pdf_junk_both_ends", "balanta.csv", PDF_MIME),
    ("pdf_junk_both_ends", "balanta", XLSX),
])
def test_g1_an_agras_pdf_is_identified_on_the_branch_the_pipeline_reads_it_on_and_lands_in_agras(
        app, gw, shape, filename, mime):
    """Review round 4 (2026-10-02). Two byte shapes the pipeline reads as a
    PDF and the identifier answered "unreadable" for (`PDF_SHAPES`): the
    card offered the company ON SCREEN and one tap stored and analysed
    Agras's PDF in Scandia. Through the real routes — /api/uploads/identify
    is handed the declared MIME type by the browser, and must hand it to the
    identifier: the pipeline picks its branch from the name AND that type.

    Reds on: the identifier reading the magic bytes its own way ("PK" as a
    zip); the identifier, the route or the seam between them classifying by
    the name alone; anything of the file stored or analysed in the company
    on screen."""
    clean = balance_pdf_of("AGRAS SRL", CUI_AGRAS)
    content = PDF_SHAPES[shape](clean)
    as_pdf = identify(app, clean, "balanta.pdf", org=ORG_SCANDIA, mime=PDF_MIME)
    assert as_pdf.status_code == 200 and as_pdf.json()["identity"]["cui"] == CUI_AGRAS, as_pdf.text[:300]
    before_scandia = (gw.docs(org_id=ORG_SCANDIA), gw.objects(ORG_SCANDIA))

    r = identify(app, content, filename, org=ORG_SCANDIA, mime=mime)
    assert r.status_code == 200, r.text[:300]
    ident = r.json()
    # The period is not compared: the title block's printed period is read
    # off a PDF whose header sits at offset 0 only, under every name alike.
    for field in ("cui", "company_name", "document_kind", "caen_code", "industry_key"):
        assert ident["identity"][field] == as_pdf.json()["identity"][field], \
            "%s as (%r, %r): the identity's %s is %r, the same PDF named balanta.pdf has %r" % (
                shape, filename, mime, field, ident["identity"][field], as_pdf.json()["identity"][field])
    assert ident["identity"]["sources"]["cui"]["signal"] == "document_header_cui", ident["identity"]["sources"]
    assert ident["target"] == {"org_id": ORG_AGRAS, "name": "Agras SRL", "is_new": False,
                               "reason": "cui_match"}, ident["target"]

    c = commit(app, content, filename, mime=mime, target_org_id=ident["target"]["org_id"],
               period_end="2025-12-31", output_language="ro")
    assert c.status_code == 200 and c.json()["status"] == "queued" and c.json()["org_id"] == ORG_AGRAS, c.text[:300]
    doc = run_analysis(gw, c.json()["document_id"])
    # Non-vacuity: the pipeline READS these bytes (with a model that answers
    # every request "no credit" — the extraction is the strict reader's): the
    # identity is of a file that is analysed, not of one refused anyway.
    assert doc["status"] == "analyzed" and doc["org_id"] == ORG_AGRAS, (doc["status"], doc.get("error"))
    (period,) = gw.db.rows("financial_periods")
    assert period["org_id"] == ORG_AGRAS and period["source_document_id"] == doc["id"], period
    assert (gw.docs(org_id=ORG_SCANDIA), gw.objects(ORG_SCANDIA)) == before_scandia, \
        "G1: something of the Agras PDF (%s as %r) landed in Scandia" % (shape, filename)


def test_g1_scandia_on_screen_never_captures_a_cui_it_does_not_hold(app, gw):
    """The mirror: a Scandia file dropped on the AGRAS page lands in Scandia."""
    content = agras_workbook(cui=CUI_SCANDIA, name="SCANDIA FOOD SRL")
    out = one_tap(app, content, "balanta.xlsx", on_screen=ORG_AGRAS)
    assert out["identify"]["target"]["org_id"] == ORG_SCANDIA, out["identify"]["target"]
    assert out["commit"]["org_id"] == ORG_SCANDIA
    assert gw.docs(org_id=ORG_AGRAS) == [] and gw.objects(ORG_AGRAS) == []


# ══════════════════════════════════════════════════════════════════════
# G2 — the period is the document's, never the file name's
# ══════════════════════════════════════════════════════════════════════


def test_g2_a_2017_file_name_whose_period_line_says_2025_is_2025(app, gw):
    content = agras_workbook(period_line="Balanta de verificare la data de 31.12.2025")
    out = one_tap(app, content, "balanta_2017.xlsx")
    ident = out["identify"]["identity"]
    assert ident["period_end"] == "2025-12-31", ident
    assert ident["sources"]["period_end"]["signal"] not in _uploads.FILENAME_SIGNALS, ident["sources"]
    assert "31.12.2025" in (ident["sources"]["period_end"]["evidence"] or ""), ident["sources"]
    (doc,) = gw.docs(id=out["commit"]["document_id"])
    assert doc["period_end_hint"] == "2025-12-31", doc
    assert doc["original_filename"] == "balanta_2017.xlsx"
    # The persist stage files the analysis under the confirmed period, not
    # the one the file name suggests.
    assert pipeline.resolve_period_end_for_persist(doc, {"period_end": None})[0] == "2025-12-31"


def test_g2_a_file_named_for_december_2025_whose_document_says_2024_is_2024(app, gw):
    content = agras_workbook(period_line="Balanta de verificare la data de 31.12.2024")
    ident = identify(app, content, "balanta_12_2025.xlsx").json()["identity"]
    assert ident["period_end"] == "2024-12-31", ident


def test_g2_no_period_line_means_no_period_offered_never_the_file_name(app, gw):
    content = agras_workbook(period_line=None)
    ident = identify(app, content, "balanta_31.12.2019.xlsx").json()["identity"]
    assert ident["cui"] == CUI_AGRAS
    assert ident["period_end"] != "2019-12-31", ident
    # And commit never invents one.
    r = commit(app, content, "balanta_31.12.2019.xlsx", target_org_id=ORG_AGRAS)
    assert r.status_code == 422, r.text[:200]
    assert gw.docs() == [] and gw.meter.reserved == []


# ══════════════════════════════════════════════════════════════════════
# G3 — the same file twice: not stored, not analysed, not counted
# ══════════════════════════════════════════════════════════════════════


def test_g3_the_same_file_twice_is_stored_analysed_and_counted_once(app, gw):
    content = agras_workbook()
    first = one_tap(app, content, "balanta.xlsx")
    doc = run_analysis(gw, first["commit"]["document_id"])
    assert doc["status"] == "analyzed", (doc["status"], doc.get("error"))
    assert gw.meter.committed == [(USER, False)], "the first copy is counted once"
    before = gw.state()
    counts = (len(gw.meter.reserved), len(gw.meter.committed), len(gw.meter.released), list(gw.enqueued))

    # identify already knows — the card says "Already uploaded — open it".
    again = identify(app, content, "balanta (1).xlsx")
    assert again.status_code == 200, again.text[:300]
    assert again.json()["duplicate"] == {"document_id": doc["id"], "period_id": doc["period_id"],
                                         "org_id": ORG_AGRAS}, again.json()["duplicate"]
    # A commit anyway (a stale tab, a second click) is answered, not stored.
    r = commit(app, content, "balanta (1).xlsx", target_org_id=ORG_AGRAS, period_end="2025-12-31")
    assert r.status_code == 200, r.text[:300]
    assert r.json() == {"status": "duplicate", "document_id": doc["id"], "period_id": doc["period_id"],
                        "org_id": ORG_AGRAS, "company_name": "Agras SRL"}, r.json()
    assert gw.state() == before, "G3: the second copy was stored"
    assert (len(gw.meter.reserved), len(gw.meter.committed), len(gw.meter.released),
            list(gw.enqueued)) == counts, "G3: the second copy was analysed or counted"


def test_g3_a_counted_book_re_uploaded_on_the_card_after_its_analysis_failed_is_not_counted_again(app, gw):
    """Lane A's lens S (fix/dedupe-quota a8c1c8cf, S10) through the
    redesign's own entry. A book the plan COUNTED whose later correction
    failed — `failed`, its period gone — is no live original, so the card
    stores and analyses the re-upload (the user gets the book back); but the
    book was counted at its first analysis and is never counted — nor
    reserved, nor at the cap put to the €-dialog — again. /run reads that
    fact from the quota ledger (`pipeline._book_already_counted`); the
    commit, which meters BEFORE it stores, must read the same fact."""
    from engine.api import _quota_ledger
    content = agras_workbook()
    first = one_tap(app, content, "balanta.xlsx")
    doc = run_analysis(gw, first["commit"]["document_id"])
    assert doc["status"] == "analyzed", (doc["status"], doc.get("error"))
    assert gw.meter.committed == [(USER, False)]
    assert [r["document_id"] for r in gw.db.rows(_quota_ledger.TABLE) if r.get("committed_at")] == [doc["id"]]
    # The free correction re-run failed (a PDF on an empty Anthropic balance,
    # §24): the counted copy is `failed` and its analysis is gone.
    period_id = doc["period_id"]
    gw.db.update("documents", {"status": "failed", "period_id": None, "error": "correction failed"},
                 filters={"id": "eq.%s" % doc["id"]})
    gw.db.tables["financial_periods"] = [p for p in gw.db.rows("financial_periods") if p["id"] != period_id]
    reserved = list(gw.meter.reserved)

    again = identify(app, content, "balanta (2).xlsx", org=ORG_AGRAS)
    assert again.status_code == 200 and again.json()["duplicate"] is None, again.text[:300]
    r = commit(app, content, "balanta (2).xlsx", target_org_id=ORG_AGRAS, period_end="2025-12-31")
    assert r.status_code == 200 and r.json()["status"] == "queued", r.text[:300]
    assert gw.meter.reserved == reserved, "G3: the card reserved a book the plan already counted"
    second = run_analysis(gw, r.json()["document_id"])
    assert second["status"] == "analyzed", (second["status"], second.get("error"))
    assert gw.meter.committed == [(USER, False)], "G3: the book was counted a second time"


def test_g3_a_commit_whose_hand_off_fails_gives_its_reservation_back_once_ledger_included(app, gw, monkeypatch):
    """Nothing ran, so nothing is counted — and the slot goes back EXACTLY
    once. Since fix/dedupe-quota 11f84259 (the restart repair) a registered
    reservation is also a row in the quota ledger, which an orphan sweep
    releases when its owner process is gone. The commit's failure path used
    to hand the slot back to the meter and leave that row outstanding: the
    sweep would then give the same slot back a second time."""
    from engine.api import _quota_ledger

    def _no_thread(doc_id: str) -> None:
        raise RuntimeError("the daemon thread could not start")

    monkeypatch.setattr(pipeline, "_enqueue", _no_thread)
    content = agras_workbook()
    ident = identify(app, content, "balanta.xlsx", org=ORG_AGRAS).json()
    r = commit(app, content, "balanta.xlsx", target_org_id=ORG_AGRAS,
               period_end=ident["identity"]["period_end"])
    assert r.status_code == 500, r.text[:300]
    assert gw.meter.reserved == [USER] and gw.meter.committed == []
    assert gw.meter.released == [(USER, False)], "G3: the slot was not given back exactly once"
    (doc,) = gw.docs(org_id=ORG_AGRAS)
    rows = [x for x in gw.db.rows(_quota_ledger.TABLE) if x["document_id"] == doc["id"]]
    assert rows and all(not x.get("reserved_at") for x in rows), (
        "G3: the ledger still holds the reservation — the orphan sweep would release it again", rows)


def test_g3_the_same_bytes_for_another_period_or_company_are_not_duplicates(app, gw):
    content = agras_workbook()
    first = one_tap(app, content, "balanta.xlsx")
    run_analysis(gw, first["commit"]["document_id"])
    r = commit(app, content, "balanta.xlsx", target_org_id=ORG_AGRAS, period_end="2024-12-31")
    assert r.json()["status"] == "queued", r.text[:300]
    r = commit(app, content, "balanta.xlsx", target_org_id=ORG_SCANDIA, period_end="2025-12-31")
    assert r.json()["status"] == "queued", r.text[:300]


def test_g3_a_confirmed_extra_on_the_card_is_one_reservation_its_run_holds_as_its_own(app, gw, monkeypatch):
    """The card's 402 -> Confirm -> the same commit with confirm_extra=1,
    the quota ledger present. The confirm records the extra's reservation
    in the ledger (fix/dedupe-quota, the reservation-outstanding repair);
    the commit's run registers THAT reservation as its own. Recording it a
    second time is refused as outstanding — and that refusal gave the
    confirmed slot straight back and left the run unmetered, the extra
    never counted."""
    from engine.api import _plan_state

    rpcs = []  # type: List[str]

    def _rpc(fn: str, params: Any) -> Any:
        rpcs.append(fn)
        return {"used": 15, "reserved": 1} if fn == "reserve_user_upload_extra" else {}

    monkeypatch.setattr(_usage_gate, "_rpc", _rpc)
    monkeypatch.setattr(_plan_state, "get_plan_state", lambda uid: types.SimpleNamespace(
        docs_used_this_period=15,
        plan=types.SimpleNamespace(key="professional", extra_doc_eur=3.0, included_docs=15)))
    c = commit(app, agras_workbook(), "balanta.xlsx", target_org_id=ORG_AGRAS, period_end="2025-12-31",
               output_language="ro", confirm_extra="1")
    assert c.status_code == 200 and c.json()["status"] == "queued", c.text[:300]
    doc_id = c.json()["document_id"]
    (row,) = [r for r in gw.db.rows("document_quota_ledger") if r["document_id"] == doc_id]
    run = pipeline._QUOTA_RUNS.get(doc_id)
    assert run is not None and run.was_extra, "G3: the confirmed extra's run holds no reservation"
    assert run.reservation_id == row["reservation_id"], "the run recorded a second reservation"
    assert gw.meter.released == [] and gw.meter.reserved == [], "the confirmed slot was given back"
    assert rpcs.count("reserve_user_upload_extra") == 1, rpcs
    doc = run_analysis(gw, doc_id)
    assert doc["status"] == "analyzed", (doc["status"], doc.get("error"))
    assert gw.meter.committed == [(USER, True)], "the confirmed extra is counted once, as the extra"
    assert gw.meter.released == []


# ══════════════════════════════════════════════════════════════════════
# G7 — drop -> one tap -> the five stages -> the analysed dashboard
# ══════════════════════════════════════════════════════════════════════


E2E_DIR = REPO / "e2e" / "fixtures" / "workspace_v2"
#: The figures the e2e spec reads off the page; held equal to the route.
E2E_PINNED = (
    ("period", "period", "period_end"),
    ("period", "period", "methodology_version"),
    ("period", "organization", "id"),
    ("period", "organization", "name"),
    ("period", "statements", "incomeStatement", "revenue"),
    ("period", "statements", "assembled_pl", "revenue"),
    ("period", "statements", "assembled_pl", "ebitda"),
    ("period", "statements", "assembled_pl", "net_income_statutory"),
    ("years", 0, "revenue"),
    ("years", 0, "year"),
    ("years", 0, "period_id"),
    ("identify", "identity"),
    ("identify", "target"),
    ("identify", "duplicate"),
    ("identify", "companies"),
    ("commit",),
)

#: (fixture, corpus book, name printed in its header, CUI, company on screen
#: when it is dropped, company it must land in, fixed e2e ids).
G7_BOOKS = (
    ("agras_fy2025", AGRAS_BOOK, "AGRAS SRL", CUI_AGRAS, ORG_SCANDIA, ORG_AGRAS,
     "5ea50000-0000-4000-8000-00000000a925", "d0c50000-0000-4000-8000-00000000a925"),
    ("scandia_fy2025", SCANDIA_BOOK, "SCANDIA FOOD SRL", CUI_SCANDIA, ORG_SCANDIA, ORG_SCANDIA,
     "5ea50000-0000-4000-8000-0000000051f5", "d0c50000-0000-4000-8000-0000000051f5"),
)


def _dig(obj: Any, path: Any) -> Any:
    for key in path:
        obj = obj[key]
    return obj


def _e2e_fixture(body: Dict[str, Any], tile: Dict[str, Any], *, identify: Dict[str, Any],
                 commit: Dict[str, Any], doc_id: str, period_id: str,
                 e2e_period: str, e2e_doc: str) -> Dict[str, Any]:
    """What the engine answered at each step for this book — identify (with
    Scandia on screen), commit, the year tile and the served period — with
    the double's sequential ids replaced by the fixed ones the e2e double
    routes by."""
    text = json.dumps({"identify": identify, "commit": commit, "period": body, "years": [tile]},
                      sort_keys=True, default=str)
    text = text.replace(period_id, e2e_period).replace(doc_id, e2e_doc)
    return json.loads(text)


def _served_revenue(env: Dict[str, Any]) -> float:
    from engine.serving.facts import FactsGateway
    return FactsGateway.from_envelope(copy.deepcopy(env), currency="RON").revenue().to_float()


@pytest.mark.parametrize("fixture_name,book,header_name,cui,on_screen,lands_in,e2e_period,e2e_doc",
                         G7_BOOKS, ids=[b[0] for b in G7_BOOKS])
def test_g7_drop_one_tap_five_stages_then_the_served_dashboard(app, gw, monkeypatch, fixture_name, book,
                                                                header_name, cui, on_screen, lands_in,
                                                                e2e_period, e2e_doc):
    content = book_workbook(book, name=header_name, cui=cui)
    statuses = []  # type: List[str]
    real_set = pipeline._admin_set_status

    def _watch(doc_id: str, status: str, **kw: Any) -> Any:
        statuses.append(status)
        return real_set(doc_id, status, **kw)

    monkeypatch.setattr(pipeline, "_admin_set_status", _watch)
    out = one_tap(app, content, "balanta_2017.xlsx", on_screen=on_screen)
    assert out["commit"]["org_id"] == lands_in, out["commit"]
    doc = run_analysis(gw, out["commit"]["document_id"])

    assert doc["status"] == "analyzed", (doc["status"], doc.get("error"))
    # The five live steps the card shows, in order, never walking back.
    order = ["extracting", "mapping", "computing", "narrating", "analyzed"]
    steps = [s for s in statuses if s in order]
    assert [s for s in order if s in steps] == order, statuses
    assert steps == sorted(steps, key=order.index), statuses
    # G4 + G2 at the end of the run: exactly one period, in the company the
    # header names, December 2025, backed by this analysed document.
    periods = gw.db.rows("financial_periods")
    assert len(periods) == 1, periods
    (period,) = periods
    assert period["org_id"] == lands_in and period["source_document_id"] == doc["id"], period
    assert str(period["period_end"])[:10] == "2025-12-31", period["period_end"]
    assert doc["period_id"] == period["id"]
    assert gw.meter.committed == [(USER, False)] and gw.meter.released == []
    from engine.workspaces.migration_plan import empty_live_periods
    assert empty_live_periods(gw.db.tables) == [], "G4: the analysed run left an empty period"
    # F4.3: the run persists its detection envelope and methodology pin.
    # Both are written by a "non-fatal" try that swallowed a NameError on
    # this branch (the effective industry key left behind when the
    # valuation block moved into `_compute_and_persist_valuation`), so no
    # upload carried them — the served period said methodology_version null.
    assert period.get("methodology_version") == "ro_ras_2025_v1", period.get("methodology_version")
    assert (period.get("detection_envelope") or {}).get("methodology_version") == "ro_ras_2025_v1", (
        period.get("detection_envelope"))

    # The company page's year tile and the dashboard read ONE authority.
    years = _http(app).get("/api/companies/%s/years" % lands_in, headers=_headers(USER))
    assert years.status_code == 200, years.text[:300]
    (tile,) = years.json()
    served = _served_revenue(period["assembled_canonical_v1"])
    # Turnover and its growth only (owner ruling 2026-09-26): the page's
    # `revenue` names ARE the turnover.
    assert tile == {"period_id": period["id"], "year": 2025, "period_end": "2025-12-31",
                    "turnover": served, "turnover_change_pct": None,
                    "revenue": served, "revenue_change_pct": None,
                    "basis": {"ro": "cifra de afaceri netă (70x − 709)",
                              "en": "net turnover (70x − 709)"},
                    "currency": "RON"}, tile
    assert served > 0
    dash = _http(app).get("/api/period/%s" % period["id"], headers=_headers(USER, lands_in))
    assert dash.status_code == 200, dash.text[:400]
    body = dash.json()
    assert body["period"]["id"] == period["id"] and body["organization"]["id"] == lands_in, body["period"]
    assert body["period"]["period_end"] == "2025-12-31", body["period"]
    assert body["period"]["source_document"]["status"] == "analyzed", body["period"]["source_document"]
    # One authority: the tile's revenue IS the revenue the dashboard prints.
    assert body["statements"]["incomeStatement"]["revenue"] == tile["revenue"], \
        (body["statements"]["incomeStatement"]["revenue"], tile["revenue"])
    assert body["statements"]["assembled_pl"]["revenue"] == tile["revenue"]
    assert body["line_items"] and body["metrics"], "an analysed period with nothing persisted"

    # The browser-side G7 (e2e/workspace-v2.spec.ts) serves THIS payload from
    # its hermetic double. The committed fixture must be what this route
    # serves for this book — every figure the e2e reads, at least.
    path = E2E_DIR / ("%s.json" % fixture_name)
    served_fixture = _e2e_fixture(body, tile, identify=out["identify"], commit=out["commit"],
                                  doc_id=doc["id"], period_id=period["id"],
                                  e2e_period=e2e_period, e2e_doc=e2e_doc)
    if os.environ.get("WS_V2_WRITE_FIXTURE") == "1":
        E2E_DIR.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(served_fixture, indent=1, sort_keys=True, default=str) + "\n",
                        encoding="utf-8")
    assert path.is_file(), "missing %s — regenerate with WS_V2_WRITE_FIXTURE=1" % path
    fixture = json.loads(path.read_text(encoding="utf-8"))
    for key in E2E_PINNED:
        assert _dig(fixture, key) == _dig(served_fixture, key), (
            "e2e fixture drifted from the served route at %s: %r != %r — regenerate with "
            "WS_V2_WRITE_FIXTURE=1" % (".".join(str(k) for k in key), _dig(fixture, key),
                                       _dig(served_fixture, key)))


def test_g7_the_e2e_feature_registry_is_what_the_app_serves(app, monkeypatch):
    """The e2e double serves `GET /api/features/status` from a committed
    capture; it must be byte-for-byte what this app serves (workspace_v2 as
    `preview`), so the browser gate opens the redesign exactly the way
    production would for an opted-in user — never through an invented row."""
    monkeypatch.delenv("CFO_FEATURES_ACTIVE", raising=False)
    served = _http(app).get("/api/features/status").json()
    assert served["features"]["workspace_v2"]["status"] == "preview", served["features"]["workspace_v2"]
    path = E2E_DIR / "features_status.json"
    if os.environ.get("WS_V2_WRITE_FIXTURE") == "1":
        E2E_DIR.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(served, indent=1, sort_keys=True) + "\n", encoding="utf-8")
    assert path.is_file(), "missing %s — regenerate with WS_V2_WRITE_FIXTURE=1" % path
    assert json.loads(path.read_text(encoding="utf-8")) == served, \
        "the e2e feature registry drifted from the app — regenerate with WS_V2_WRITE_FIXTURE=1"


# ══════════════════════════════════════════════════════════════════════
# G4 — no empty period: the code path, and the check for production data
# ══════════════════════════════════════════════════════════════════════


def test_g4_identify_and_commit_create_no_period(app, gw):
    content = agras_workbook()
    one_tap(app, content, "balanta.xlsx")
    assert gw.db.rows("financial_periods") == [], "G4: the upload created a period before the analysis"


def test_g4_a_run_that_fails_after_persist_leaves_no_period(app, gw, monkeypatch):
    content = agras_workbook()
    out = one_tap(app, content, "balanta.xlsx")

    def _boom(*a: Any, **kw: Any) -> Any:
        raise RuntimeError("compute failed")

    monkeypatch.setattr(pipeline, "stage_compute", _boom)
    doc = run_analysis(gw, out["commit"]["document_id"])
    assert doc["status"] == "failed", doc["status"]
    assert gw.db.rows("financial_periods") == [], "G4: a failed run left its period"
    assert gw.meter.committed == [] and gw.meter.released == [(USER, False)], "a failure is never counted"


# ── G4 — a same-month re-upload never empties the month ────────────────
#
# The production defect (2026-09, the owner's December 2025): a second file
# for a month that already had an analysed one made `stage_persist` re-point
# the month's period at the NEW document and wipe the first document's line
# items and envelope BEFORE the new run had succeeded. An ordinary failure
# then left a period whose only source had failed — the company page listed
# no year at all. The takeover is now transactional by order: the new run
# persists under its own staged row, and only its terminal success replaces
# the month; a cross-company file never replaces anything.
#
# Reds on: a failed re-upload changing the month's source document, its
# envelope, any derivative row, the year tile or the served dashboard; a file
# whose CUI is another company's replacing the month (or failing with a
# stack-trace message) — the company's CUI on file, or for a company without
# one the CUI the month's own file states; a successful re-upload leaving a
# second period, a derivative row under the staged id, or the first file live
# and unarchived; a dead-letter replay of a re-upload replacing the month or
# leaving its staged row as a second period of the month.

#: The rows a run persists under a period id. Named here, not read from the
#: engine, so a RED run fails on the defect and never on a missing name.
_PERIOD_ROWS = ("statement_line_items", "calculated_metrics", "alerts",
                "recommendations", "briefings", "valuations")


def _rows_under(gw: GateWorld, period_id: str) -> Dict[str, List[str]]:
    return dict((t, sorted(json.dumps(r, sort_keys=True, default=str)
                           for r in gw.db.rows(t) if str(r.get("period_id")) == str(period_id)))
                for t in _PERIOD_ROWS)


def _served(app, org_id: str, period_id: str) -> Dict[str, Any]:
    """The year tile and the dashboard body — what the user sees of a month."""
    years = _http(app).get("/api/companies/%s/years" % org_id, headers=_headers(USER))
    assert years.status_code == 200, years.text[:300]
    dash = _http(app).get("/api/period/%s" % period_id, headers=_headers(USER, org_id))
    assert dash.status_code == 200, dash.text[:400]
    body = dash.json()
    return {"tiles": years.json(),
            "source_document": body["period"]["source_document"]["id"],
            "revenue": body["statements"]["incomeStatement"]["revenue"],
            "line_items": len(body["line_items"]), "metrics": len(body["metrics"])}


def _analysed_month(app, gw: GateWorld, content: bytes, filename: str = "balanta.xlsx") -> Dict[str, Any]:
    out = one_tap(app, content, filename)
    doc = run_analysis(gw, out["commit"]["document_id"])
    assert doc["status"] == "analyzed", (doc["status"], doc.get("error"))
    (period,) = gw.db.rows("financial_periods")
    assert period["source_document_id"] == doc["id"]
    return {"doc": doc, "period": copy.deepcopy(period),
            "rows": _rows_under(gw, period["id"]), "served": _served(app, doc["org_id"], period["id"])}


def test_g4_a_same_month_reupload_whose_run_fails_leaves_the_month_serving_the_first_analysis(app, gw, monkeypatch):
    first = _analysed_month(app, gw, agras_workbook())
    org_id = first["doc"]["org_id"]
    # The same company's December again — a corrected export: same CUI, same
    # month, different bytes (not a duplicate), a run that fails at compute.
    second = one_tap(app, agras_workbook(name="AGRAS S.R.L."), "balanta_corectata.xlsx")
    assert second["commit"]["org_id"] == org_id and second["commit"]["status"] == "queued", second["commit"]

    def _boom(*a: Any, **kw: Any) -> Any:
        raise RuntimeError("compute failed")

    monkeypatch.setattr(pipeline, "stage_compute", _boom)
    failed = run_analysis(gw, second["commit"]["document_id"])
    assert failed["status"] == "failed", failed["status"]

    periods = gw.db.rows("financial_periods")
    assert [p["id"] for p in periods] == [first["period"]["id"]], \
        "G4: the failed re-upload left a period of its own: %r" % periods
    (period,) = periods
    assert period["source_document_id"] == first["doc"]["id"], \
        "G4: the failed re-upload took the month over — its period now names the failed document"
    assert period["assembled_canonical_v1"] == first["period"]["assembled_canonical_v1"], \
        "G4: the failed re-upload overwrote the month's envelope"
    assert _rows_under(gw, period["id"]) == first["rows"], "G4: the failed re-upload changed the month's rows"
    assert _served(app, org_id, period["id"]) == first["served"], "G4: the month serves something else after the failure"
    (kept,) = gw.docs(id=first["doc"]["id"])
    assert kept["status"] == "analyzed" and kept["deleted_at"] is None and kept["period_id"] == period["id"], kept
    assert failed["period_id"] is None, "a failed document is never pinned to the month it did not replace"
    from engine.workspaces.migration_plan import empty_live_periods
    assert empty_live_periods(gw.db.tables) == [], "G4: the failure left an empty period"
    assert gw.meter.committed == [(USER, False)] and gw.meter.released == [(USER, False)], \
        "the first run counted, the failed one released"


@pytest.mark.parametrize("filename,mime", [("balanta.xlsx", XLSX), ("balanta.pdf", PDF_MIME)])
def test_g4_a_same_month_file_of_another_company_never_replaces_the_month(app, gw, filename, mime):
    """…under ANY name the file is read under. The .pdf-named twin (review
    round 3, 2026-10-02, HIGH): the wall reads the document's CUI through
    the identifier, which chose its reader by name — Scandia's workbook
    named balanta.pdf had no CUI to the wall, was analysed, replaced
    Agras's December (the served revenue became Scandia's, the first
    document archived `superseded_by`) and the meter committed twice."""
    first = _analysed_month(app, gw, agras_workbook())
    org_id = first["doc"]["org_id"]
    # Scandia's book, committed into Agras by an explicit choice on the card
    # (the route lets a member file anything into their own company). The
    # persist layer is the belt and braces: December belongs to CUI 46355095.
    scandia = book_workbook(SCANDIA_BOOK, name="SCANDIA FOOD SRL", cui=CUI_SCANDIA)
    c = commit(app, scandia, filename, mime=mime, target_org_id=org_id, period_end="2025-12-31",
               output_language="ro")
    assert c.status_code == 200 and c.json()["status"] == "queued" and c.json()["org_id"] == org_id, c.text[:300]
    refused = run_analysis(gw, c.json()["document_id"])
    _assert_the_wall_held(app, gw, first, refused, "named %s" % filename)


def _assert_the_wall_held(app, gw: GateWorld, first: Dict[str, Any], refused: Dict[str, Any], what: str) -> None:
    """Another company's file (Scandia's CUI) met Agras's served month: the
    run FAILED on the wall's own sentence, and the month, its rows, what it
    serves, its document and the meter are exactly as the first analysis
    left them."""
    org_id = first["doc"]["org_id"]
    assert refused["status"] == "failed", \
        "G4: another company's file %s was analysed over the month: %r" % (
            what, (refused["status"], refused.get("error")))
    message = str(refused.get("error") or "")
    assert CUI_SCANDIA in message and CUI_AGRAS in message, message
    assert not re.match(r"^\w*(Error|Exception|Refused)\w*:", message), \
        "a refusal is a plain sentence, not an exception's repr: %r" % message
    assert "source" not in message.lower(), message

    (period,) = gw.db.rows("financial_periods")
    assert period["id"] == first["period"]["id"] and period["source_document_id"] == first["doc"]["id"], \
        "G4: another company's file replaced the month"
    assert period["assembled_canonical_v1"] == first["period"]["assembled_canonical_v1"]
    assert _rows_under(gw, period["id"]) == first["rows"]
    assert _served(app, org_id, period["id"]) == first["served"]
    assert refused["period_id"] is None
    (kept,) = gw.docs(id=first["doc"]["id"])
    assert kept["status"] == "analyzed" and kept["deleted_at"] is None, \
        "G4: the month's own document was archived by another company's file: %r" % kept
    assert gw.meter.committed == [(USER, False)], \
        "G4: the refused file was counted: %r" % (gw.meter.committed,)
    from engine.workspaces.migration_plan import empty_live_periods
    assert empty_live_periods(gw.db.tables) == []


#: (shape, name, declared MIME type). The first row is the CONTROL — a clean
#: PDF under its own name, which the wall held before round 4 too.
_WALL_PDFS = [
    ("pdf", "balanta.pdf", PDF_MIME),
    ("pdf_behind_PK", "balanta.pdf", PDF_MIME),
    ("pdf_behind_PK", "balanta.xls", XLS_MIME),
    ("pdf_behind_PK", "balanta", OCTET),
    ("pdf_junk_both_ends", "balanta.pdf", PDF_MIME),
    ("pdf_junk_both_ends", "balanta", PDF_MIME),
    ("pdf_junk_both_ends", "balanta.csv", PDF_MIME),
    ("pdf_junk_both_ends", "balanta", XLSX),
]


@pytest.mark.parametrize("shape,filename,mime", _WALL_PDFS,
                         ids=["%s|%s|%s" % (s_, n, m.split("/")[-1][:12]) for s_, n, m in _WALL_PDFS])
def test_g4_a_same_month_pdf_of_another_company_never_replaces_the_month(app, gw, shape, filename, mime):
    """Review round 4 (2026-10-02), a measured bypass of this wall. The wall
    reads the document's CUI through the identifier
    (`pipeline._document_company_cui`). On the two byte shapes of
    `PDF_SHAPES` the pipeline read Scandia's PDF and the identifier answered
    "unreadable": no CUI, the run was ANALYSED, Agras's December was
    replaced, its own document archived and the meter committed twice —
    under all eleven names for the "PK" shape, and for the other wherever
    the declared MIME type alone (or against the extension) picks the
    branch.

    Reds on: the identifier's dispatch differing from the pipeline's on a
    byte shape; `_document_company_cui` not handing the identifier the
    document's `mime_type`."""
    first = _analysed_month(app, gw, agras_workbook())
    org_id = first["doc"]["org_id"]
    content = PDF_SHAPES[shape](balance_pdf_of("SCANDIA FOOD SRL", CUI_SCANDIA))
    c = commit(app, content, filename, mime=mime, target_org_id=org_id, period_end="2025-12-31",
               output_language="ro")
    assert c.status_code == 200 and c.json()["status"] == "queued" and c.json()["org_id"] == org_id, c.text[:300]
    refused = run_analysis(gw, c.json()["document_id"])
    _assert_the_wall_held(app, gw, first, refused, "(%s) as (%r, %r)" % (shape, filename, mime))


@pytest.mark.parametrize("shape,filename,mime", [
    ("pdf_behind_PK", "balanta.pdf", PDF_MIME),
    ("pdf_junk_both_ends", "balanta", PDF_MIME),
])
def test_g4_a_same_month_pdf_of_the_same_company_still_replaces_the_month(app, gw, shape, filename, mime):
    """The other half, and the non-vacuity of the law above: the pipeline
    READS both byte shapes — the SAME company's PDF in that shape is
    analysed and becomes the month. So for another company's, the wall is
    the only thing in the way; and the wall refuses only a provable other
    company."""
    first = _analysed_month(app, gw, agras_workbook())
    org_id = first["doc"]["org_id"]
    content = PDF_SHAPES[shape](balance_pdf_of("AGRAS SRL", CUI_AGRAS))
    c = commit(app, content, filename, mime=mime, target_org_id=org_id, period_end="2025-12-31",
               output_language="ro")
    assert c.status_code == 200 and c.json()["status"] == "queued", c.text[:300]
    replaced = run_analysis(gw, c.json()["document_id"])
    assert replaced["status"] == "analyzed", (replaced["status"], replaced.get("error"))
    (period,) = gw.db.rows("financial_periods")
    assert period["id"] == first["period"]["id"] and period["source_document_id"] == replaced["id"], period


def test_g4_the_month_of_a_company_without_a_cui_is_its_own_pdfs_whatever_mime_named_it(app, gw):
    """`_served_document_cui`: for a company with no CUI on file, the month's
    OWN document is the evidence of whose month it is — read through the
    same identifier, so it needs that document's `mime_type` too. Agras's
    December is served from a PDF stored as `balanta` / application/pdf (the
    declared type alone picks its branch); Scandia's workbook must not
    replace it.

    Reds on: the select of the month's own document dropping `mime_type`
    (the wall then finds no CUI on either record and lets the month go)."""
    own = PDF_SHAPES["pdf_junk_both_ends"](balance_pdf_of("AGRAS SRL", CUI_AGRAS))
    c = commit(app, own, "balanta", mime=PDF_MIME, target_org_id=ORG_AGRAS, period_end="2025-12-31",
               output_language="ro")
    assert c.status_code == 200 and c.json()["status"] == "queued", c.text[:300]
    doc = run_analysis(gw, c.json()["document_id"])
    assert doc["status"] == "analyzed", (doc["status"], doc.get("error"))
    (period,) = gw.db.rows("financial_periods")
    first = {"doc": doc, "period": copy.deepcopy(period), "rows": _rows_under(gw, period["id"]),
             "served": _served(app, ORG_AGRAS, period["id"])}
    (bag,) = [r["prefs"] for r in gw.db.rows("org_prefs") if r["org_id"] == ORG_AGRAS]
    gw.db.update("org_prefs", {"prefs": dict((k, v) for k, v in bag.items() if k != "cui")},
                 filters={"org_id": "eq.%s" % ORG_AGRAS})
    scandia = book_workbook(SCANDIA_BOOK, name="SCANDIA FOOD SRL", cui=CUI_SCANDIA)
    c = commit(app, scandia, "balanta.xlsx", target_org_id=ORG_AGRAS, period_end="2025-12-31", output_language="ro")
    assert c.status_code == 200 and c.json()["status"] == "queued", c.text[:300]
    refused = run_analysis(gw, c.json()["document_id"])
    _assert_the_wall_held(app, gw, first, refused, "over a month served from a PDF named by its MIME type")


def test_g4_a_same_month_file_of_another_company_never_replaces_the_month_of_a_company_without_a_cui(app, gw):
    """A company created before companies were keyed by CUI has none on
    file (`org_prefs.prefs.cui` absent). The month's own file is then the
    evidence of whose month it is: Scandia's book never replaces the
    December that Agras's book is serving, CUI on file or not."""
    first = _analysed_month(app, gw, agras_workbook())
    org_id = first["doc"]["org_id"]
    # The company's CUI is not on file — a workspace from before the CUI key.
    (bag,) = [r["prefs"] for r in gw.db.rows("org_prefs") if r["org_id"] == org_id]
    gw.db.update("org_prefs", {"prefs": dict((k, v) for k, v in bag.items() if k != "cui")},
                 filters={"org_id": "eq.%s" % org_id})
    scandia = book_workbook(SCANDIA_BOOK, name="SCANDIA FOOD SRL", cui=CUI_SCANDIA)
    c = commit(app, scandia, "balanta.xlsx", target_org_id=org_id, period_end="2025-12-31", output_language="ro")
    assert c.status_code == 200 and c.json()["status"] == "queued" and c.json()["org_id"] == org_id, c.text[:300]
    (bag,) = [r["prefs"] for r in gw.db.rows("org_prefs") if r["org_id"] == org_id]
    assert "cui" not in bag, "the commit adopted another company's CUI: %r" % bag
    refused = run_analysis(gw, c.json()["document_id"])
    assert refused["status"] == "failed", \
        "G4: another company's file replaced the month of a company without a CUI on file: %r" % (
            (refused["status"], refused.get("error")),)
    message = str(refused.get("error") or "")
    assert CUI_SCANDIA in message and CUI_AGRAS in message, message
    assert not re.match(r"^\w*(Error|Exception|Refused)\w*:", message), \
        "a refusal is a plain sentence, not an exception's repr: %r" % message

    (period,) = gw.db.rows("financial_periods")
    assert period["id"] == first["period"]["id"] and period["source_document_id"] == first["doc"]["id"], \
        "G4: another company's file replaced the month"
    assert period["assembled_canonical_v1"] == first["period"]["assembled_canonical_v1"]
    assert _rows_under(gw, period["id"]) == first["rows"]
    assert _served(app, org_id, period["id"]) == first["served"]
    (kept,) = gw.docs(id=first["doc"]["id"])
    assert kept["deleted_at"] is None and kept["status"] == "analyzed", kept
    assert refused["period_id"] is None
    from engine.workspaces.migration_plan import empty_live_periods
    assert empty_live_periods(gw.db.tables) == []


def test_g4_a_same_month_file_of_the_same_company_replaces_the_month_of_a_company_without_a_cui(app, gw):
    """The other half of the same rule: without a CUI on file, the corrected
    export of the SAME company (its own CUI, as the month's file states it)
    still replaces the month — the fallback refuses only a provable other
    company."""
    first = _analysed_month(app, gw, agras_workbook())
    org_id = first["doc"]["org_id"]
    (bag,) = [r["prefs"] for r in gw.db.rows("org_prefs") if r["org_id"] == org_id]
    gw.db.update("org_prefs", {"prefs": dict((k, v) for k, v in bag.items() if k != "cui")},
                 filters={"org_id": "eq.%s" % org_id})
    corrected = book_workbook(SCANDIA_BOOK, name="AGRAS SRL", cui=CUI_AGRAS)
    c = commit(app, corrected, "balanta_corectata.xlsx", target_org_id=org_id, period_end="2025-12-31",
               output_language="ro")
    assert c.status_code == 200 and c.json()["status"] == "queued" and c.json()["org_id"] == org_id, c.text[:300]
    replaced = run_analysis(gw, c.json()["document_id"])
    assert replaced["status"] == "analyzed", (replaced["status"], replaced.get("error"))
    (period,) = gw.db.rows("financial_periods")
    assert period["id"] == first["period"]["id"] and period["source_document_id"] == replaced["id"]
    (superseded,) = gw.docs(id=first["doc"]["id"])
    assert superseded["deleted_at"] is not None and replaced["id"] in str(superseded.get("error") or "")


def test_g4_a_same_month_reupload_that_succeeds_replaces_the_month_and_archives_the_first_file(app, gw):
    first = _analysed_month(app, gw, agras_workbook())
    org_id = first["doc"]["org_id"]
    # The corrected December: Agras's header over different figures (the
    # other corpus book), so the replacement is visible in every number.
    corrected = book_workbook(SCANDIA_BOOK, name="AGRAS SRL", cui=CUI_AGRAS)
    second = one_tap(app, corrected, "balanta_corectata.xlsx")
    assert second["commit"]["org_id"] == org_id and second["commit"]["status"] == "queued", second["commit"]
    replaced = run_analysis(gw, second["commit"]["document_id"])
    assert replaced["status"] == "analyzed", (replaced["status"], replaced.get("error"))

    periods = gw.db.rows("financial_periods")
    assert [p["id"] for p in periods] == [first["period"]["id"]], \
        "G4: the successful re-upload left a second period for the month: %r" % [p["id"] for p in periods]
    (period,) = periods
    assert period["source_document_id"] == replaced["id"], "the month is the corrected document's now"
    assert period["assembled_canonical_v1"]["provenance"]["source_document_id"] == replaced["id"]
    assert replaced["period_id"] == period["id"]
    rows = _rows_under(gw, period["id"])
    assert rows["statement_line_items"] and rows["statement_line_items"] != first["rows"]["statement_line_items"]
    assert rows["calculated_metrics"] and rows["calculated_metrics"] != first["rows"]["calculated_metrics"]
    served = _served(app, org_id, period["id"])
    assert served["source_document"] == replaced["id"] and served["revenue"] != first["served"]["revenue"], served
    # The tile's turnover (and the page's `revenue` alias of it) follows the
    # corrected document (owner ruling 2026-09-26: the tile is turnover).
    assert served["tiles"] == [dict(first["served"]["tiles"][0], revenue=served["revenue"],
                                    turnover=served["revenue"])], served["tiles"]
    assert served["revenue"] == _served_revenue(period["assembled_canonical_v1"])
    # Nothing of the run is left under a staged id: every row named a period
    # that exists.
    live = set(str(p["id"]) for p in periods)
    strays = dict((t, [r for r in gw.db.rows(t) if str(r.get("period_id")) not in live]) for t in _PERIOD_ROWS)
    assert all(not v for v in strays.values()), "rows left under a period that no longer exists: %r" % strays
    # The superseded file is ARCHIVED (restorable for 30 days), never deleted,
    # and named as superseded by the corrected one.
    (superseded,) = gw.docs(id=first["doc"]["id"])
    assert superseded["deleted_at"] is not None, "the superseded file is still live beside its replacement"
    assert superseded["status"] == "analyzed" and replaced["id"] in str(superseded.get("error") or ""), superseded
    assert "documents/" + superseded["storage_path"] in gw.db.storage, "the superseded file's bytes must stay"
    from engine.workspaces.migration_plan import empty_live_periods
    assert empty_live_periods(gw.db.tables) == []
    assert gw.meter.committed == [(USER, False), (USER, False)] and gw.meter.released == []


def test_g4_a_same_month_reupload_replayed_from_the_dead_letter_queue_never_replaces_the_month(
        app, gw, monkeypatch, tmp_path):
    """The journal's resume re-runs persist by hand and marks nothing
    analyzed. A same-month re-upload replaces the month only when its run
    succeeds end to end, so its replay from the dead-letter queue refuses
    honestly (`cannot_resume`) — the month keeps serving the first analysis,
    and no staged row is left beside it for a second period of the month."""
    from engine.journal import ResumeRefused
    from engine.journal import hooks as journal_hooks
    from engine.journal.journal import Journal
    from engine.journal.resume import replay_dlq

    first = _analysed_month(app, gw, agras_workbook())
    org_id = first["doc"]["org_id"]
    root = tmp_path / "journal"
    monkeypatch.setenv(journal_hooks.ENV_VAR, str(root))
    journal_hooks.reset_cache()
    try:
        corrected = book_workbook(SCANDIA_BOOK, name="AGRAS SRL", cui=CUI_AGRAS)
        second = one_tap(app, corrected, "balanta_corectata.xlsx")
        assert second["commit"]["org_id"] == org_id, second["commit"]
        real_persist = pipeline.stage_persist

        def _outage(*a: Any, **kw: Any) -> Any:
            raise RuntimeError("the database is unavailable")

        monkeypatch.setattr(pipeline, "stage_persist", _outage)
        failed = run_analysis(gw, second["commit"]["document_id"])
        assert failed["status"] == "failed", failed["status"]
        monkeypatch.setattr(pipeline, "stage_persist", real_persist)
        (entry,) = Journal(root).dlq_entries()
        with pytest.raises(ResumeRefused) as refused:
            replay_dlq(Journal(root), entry["run_id"])
        assert refused.value.reason == "cannot_resume", refused.value
    finally:
        journal_hooks.set_active_run(None)
        journal_hooks.reset_cache()

    periods = gw.db.rows("financial_periods")
    assert [p["id"] for p in periods] == [first["period"]["id"]], \
        "G4: the replayed re-upload left a second period for the month: %r" % [
            (p["id"], p["source_document_id"]) for p in periods]
    (period,) = periods
    assert period["source_document_id"] == first["doc"]["id"], "G4: the replay took the month over"
    assert period["assembled_canonical_v1"] == first["period"]["assembled_canonical_v1"]
    assert _rows_under(gw, period["id"]) == first["rows"]
    assert _served(app, org_id, period["id"]) == first["served"]
    # (The staged row's own derivative rows go with it by the foreign keys'
    # ON DELETE CASCADE, which this double does not model — the month's
    # rows and the served body above are what the user sees.)
    (kept,) = gw.docs(id=first["doc"]["id"])
    assert kept["status"] == "analyzed" and kept["deleted_at"] is None, kept
    assert pipeline._pop_takeover(failed["id"]) is None, "the staged takeover outlived the refusal"
    from engine.workspaces.migration_plan import empty_live_periods
    assert empty_live_periods(gw.db.tables) == []

def _load_script(name: str) -> Any:
    spec = importlib.util.spec_from_file_location(name, REPO / "scripts" / ("%s.py" % name))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)  # type: ignore[union-attr]
    return mod


def test_g4_the_production_check_finds_the_empty_periods_of_a_snapshot(tmp_path):
    """`scripts/check_no_empty_periods.py` over the owner's production
    shapes (tests/engine/ws_migration_fixture.build_world: file-less
    current-month rows, a period whose document failed, one whose document
    was deleted) — read from a db_snapshot file AND live, GET-only, through
    the paging PostgREST double."""
    from ws_migration_fixture import FakeSupabase, build_world
    from engine.workspaces import pgrest_io
    from engine.workspaces.migration_plan import empty_live_periods

    check = _load_script("check_no_empty_periods")
    snapshot_cli = _load_script("db_snapshot")
    tables, storage, _rules = build_world()
    # Every rule of the check, not only the file-less rows the world
    # already carries: a period whose document FAILED, one whose document
    # lives in ANOTHER workspace, and one whose analysed document left
    # NOTHING persisted (no envelope, no metric, no line item).
    base_doc = next(d for d in tables["documents"] if d["id"] == "d-sf24-src")
    base_per = next(p for p in tables["financial_periods"] if p["id"] == "per-sf24")
    for did, org, status in (("g4-doc-failed", "org-sf", "failed"), ("g4-doc-elsewhere", "org-qa", "analyzed"),
                             ("g4-doc-bare", "org-sf", "analyzed")):
        tables["documents"].append(dict(base_doc, id=did, org_id=org, status=status, period_id=None,
                                        content_hash="%064x" % len(tables["documents"])))
    for pid, src, env in (("g4-per-failed", "g4-doc-failed", {"v": 1}), ("g4-per-elsewhere", "g4-doc-elsewhere", {"v": 1}),
                          ("g4-per-bare", "g4-doc-bare", None)):
        tables["financial_periods"].append(dict(base_per, id=pid, source_document_id=src, period_end="2023-12-31",
                                                period_start="2023-12-31", assembled_canonical_v1=env))
    expected = empty_live_periods(tables)
    reasons = dict(expected)
    assert reasons.get("g4-per-failed") == "source failed", expected
    assert reasons.get("g4-per-elsewhere") == "source in another workspace", expected
    assert reasons.get("g4-per-bare") == "nothing persisted", expected
    assert reasons.get("per-sf-empty") == "no source document" and reasons.get("per-sf21") == "source deleted", expected
    fake = FakeSupabase(tables, storage, max_rows=5)

    lines = []  # type: List[str]
    snap = str(tmp_path / "snap.json.gz")
    assert snapshot_cli.main(["--out", snap], client_factory=fake.client, out=lines.append) == 0
    lines = []
    assert check.main(["--snapshot", snap], out=lines.append) == 1, lines
    found = sorted((l.split("  ")[1], l.split("  ")[2]) for l in lines if l.startswith("EMPTY"))
    assert found == sorted(expected), (found, expected)

    writes_before = len(fake.row_writes())
    lines = []
    assert check.main([], client_factory=fake.client, out=lines.append) == 1, lines
    live = sorted((l.split("  ")[1], l.split("  ")[2]) for l in lines if l.startswith("EMPTY"))
    assert live == sorted(expected), (live, expected)
    assert len(fake.row_writes()) == writes_before, "the check wrote to production"
    assert any(l.startswith("CHECKED") for l in lines), lines

    # Clean data passes; nothing to check is a RED, never a pass.
    clean = copy.deepcopy(tables)
    bad = set(pid for pid, _why in expected)
    clean["financial_periods"] = [p for p in clean["financial_periods"] if str(p["id"]) not in bad]
    assert empty_live_periods(clean) == []
    lines = []
    assert check.main([], client_factory=FakeSupabase(clean, storage, max_rows=5).client,
                      out=lines.append) == 0, lines
    lines = []
    assert check.main([], client_factory=FakeSupabase(dict(clean, financial_periods=[]), storage,
                                                      max_rows=5).client, out=lines.append) == 2, lines


# ══════════════════════════════════════════════════════════════════════
# G8 — archive, never delete; the restore reverts the migration exactly
# ══════════════════════════════════════════════════════════════════════


_DELETE_CALL = re.compile(r"\.(delete|delete_object|remove)\s*\(|\"DELETE\"|'DELETE'|method\s*=\s*[\"']DELETE")


def test_g8_the_migration_archives_and_db_restore_reverts_it_exactly(tmp_path):
    """The operator path end to end, over the paging PostgREST + Storage
    double that REFUSES every DELETE (ws_migration_fixture.FakeSupabase):
    db_snapshot -> workspace_migration --execute (the reviewed plan) ->
    db_restore --plan PLAN.json --apply, the rollback (feat/ws-migration
    f6f9c7b5: it undoes exactly the plan's own operations; the table-wide
    restore is an operator tool behind --whole-tables). The plan holds no
    delete; what it removes from view it archives (deleted_at /
    archived_at); after the rollback every row of the pre-state is back,
    byte for byte, and every workspace the migration created is archived,
    held, never deleted."""
    from ws_migration_fixture import FakeSupabase, build_world
    from engine.workspaces import pgrest_io
    from engine.workspaces.rowstore import pk_for, row_key, rows_equal

    snapshot_cli = _load_script("db_snapshot")
    restore_cli = _load_script("db_restore")
    migration_cli = _load_script("workspace_migration")
    tables, storage, rules = build_world()
    pre = copy.deepcopy(tables)
    fake = FakeSupabase(tables, storage, max_rows=5)
    rules_path = tmp_path / "known.json"
    rules_path.write_text(json.dumps({"rules": rules}))
    lines = []  # type: List[str]
    run = "2026-09-21T15:00:00+00:00"
    snap = str(tmp_path / "snap.json.gz")
    assert snapshot_cli.main(["--out", snap], client_factory=fake.client, out=lines.append) == 0
    args = ["--snapshot", snap, "--known-identities", str(rules_path), "--out-dir", str(tmp_path / "out"),
            "--migration-date", "2026-09-21"]
    assert migration_cli.main(args, client_factory=fake.client, out=lines.append, now=run, registry=None) == 0
    plan = json.loads((tmp_path / "out" / "plan_2026-09-21.json").read_text())
    ops = plan["ops"]
    assert ops, "vacuous: the plan is empty"
    kinds = set(op["op"] for op in ops)
    assert kinds <= {"insert", "merge_prefs", "update", "copy_object"}, sorted(kinds)
    archives = [op for op in ops if op["op"] == "update"
                and set(op.get("set") or {}) & {"deleted_at", "archived_at"}]
    assert archives, "the fixture carries duplicates and empty periods: something must be archived"
    assert migration_cli.main(["--execute", "--expect-plan-sha", plan["ops_sha256"]] + args,
                              client_factory=fake.client, out=lines.append, now=run,
                              registry=None) == 0, "\n".join(lines[-20:])
    assert fake.deletes == [], "G8: the migration deleted"
    for table, rows in pre.items():
        pk = pk_for(table)
        have = set(row_key(r, pk) for r in fake.tables[table])
        assert all(row_key(r, pk) in have for r in rows), "G8: a %s row vanished" % table

    plan_path = str(tmp_path / "out" / "plan_2026-09-21.json")
    assert restore_cli.main([snap, "--plan", plan_path, "--apply"], client_factory=fake.client,
                            out=lines.append, now="2026-09-22T00:00:00+00:00") == 0, "\n".join(lines[-20:])
    assert fake.deletes == [], "G8: the restore deleted"
    pks = pgrest_io.snapshot_pks(pgrest_io.load_snapshot(snap))
    for table, rows in pre.items():
        if table in pgrest_io.BILLING_TABLES:
            assert fake.tables[table] == rows, "the migration touched %s" % table
            continue
        pk = pk_for(table, pks)
        cur = dict((row_key(r, pk), r) for r in fake.tables[table])
        for r in rows:
            assert rows_equal(r, cur[row_key(r, pk)]), ("G8: not restored", table, r)
    pre_orgs = set(o["id"] for o in pre["organizations"])
    created = [o for o in fake.tables["organizations"] if o["id"] not in pre_orgs]
    assert created, "vacuous: the migration created no workspace"
    assert all(o["archived_at"] and o["purge_after"] is None for o in created), created


def test_g8_the_migration_scripts_and_their_wire_carry_no_delete():
    """A static census: the migration, restore and snapshot scripts and the
    wire they write through name no DELETE — of a row or of an object."""
    files = [REPO / "scripts" / "workspace_migration.py", REPO / "scripts" / "db_restore.py",
             REPO / "scripts" / "db_snapshot.py", REPO / "scripts" / "check_no_empty_periods.py",
             REPO / "src" / "engine" / "workspaces" / "pgrest_io.py",
             REPO / "src" / "engine" / "workspaces" / "rowstore.py",
             REPO / "src" / "engine" / "workspaces" / "migration_plan.py"]
    offenders = []  # type: List[str]
    for path in files:
        assert path.is_file(), path
        for n, line in enumerate(path.read_text(encoding="utf-8").splitlines(), start=1):
            code = line.split("#", 1)[0]
            if _DELETE_CALL.search(code):
                offenders.append("%s:%d: %s" % (path.relative_to(REPO), n, line.strip()))
    assert offenders == [], "G8: a delete on the migration path:\n" + "\n".join(offenders)
