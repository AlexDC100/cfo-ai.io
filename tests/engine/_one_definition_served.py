"""The ENGINE'S served surfaces for the one-definition gates (design A8).

Owner ruling 2026-09-26: net 711 ("Variația stocurilor de produse") and net
72x inside EBITDA and the operating result, outside turnover; margins and
growth over net turnover (70x − 709); one definition across dashboard,
report, benchmark and forecast; a refused 711 refuses EBITDA and everything
built on it with the same typed reason.

The frontend halves of `one-ebitda`, `turnover-denominator` and
`refusal-carries` (frontend/lib/__tests__/oneEbitdaSurfaces / turnover
Denominator / refusalCarries) hold every BROWSER surface to the served
figure. These are the ENGINE halves: every engine surface that carries an
EBITDA, an operating result, a margin or a refusal of them, read off the
same served period, on the same books:

  corpus (real client-shaped exports committed in corpus/, through the real
  write path — stage_map -> stage_persist — and the real GET /api/period):
      agras, carniprod, realestate (the developer), retail
  CONSTRUCTED (the SYNTHETIC books of tests/engine/test_net_711_rule.py, no
  client data, through the same write seam and route):
      closed_bridge, bridge_with_722, open, closed_no_activity,
      unanchored (711 refused: account 121 absent; the sheet rebuilt to
      balance without it),
      g6_uncleared (711 refused: 121's opening not cleared),
      unanchored_unbalanced (711 refused: account 121 dropped from a real
      export — the sheet short by the missing year's result)
  DERIVED from a committed corpus book (critic round 2, 2026-09-27):
      realestate_no121 — corpus/saga_10_col_realestate/input.xlsx with its
      account-121 row (121101) DELETED FROM THE FILE, in a temporary copy
      (dropping rows in the parser is not enough: canonical_bs re-reads the
      bytes, still sees 121, and reports delta 0). The real developer with
      no anchor: 711 and the net result refused, the sheet short by the
      year's result, related-party balances present.

A surface is a small reader: (served bundle) -> a number, None, or a
Refused(code). Nothing here computes an EBITDA — the expectation is always
the served `assembled_pl` field.
"""
from __future__ import annotations

import contextlib
import copy
import types
from pathlib import Path
from typing import Any, Callable, Dict, Iterator, List, Optional, Tuple

from _pytest.monkeypatch import MonkeyPatch

import test_net_711_rule as N
import test_rebuild_net_income_anchor as ANCHOR

REPO = Path(__file__).resolve().parents[2]
AUTH = {"Authorization": "Bearer test"}

CORPUS_BOOKS = ("agras", "carniprod", "realestate", "retail")
CONSTRUCTED_BOOKS = ("closed_bridge", "bridge_with_722", "open", "closed_no_activity",
                     "unanchored", "g6_uncleared", "unanchored_unbalanced")
DERIVED_BOOKS = ("realestate_no121",)
ALL_BOOKS = CORPUS_BOOKS + CONSTRUCTED_BOOKS + DERIVED_BOOKS
REFUSED_BOOKS = ("unanchored", "g6_uncleared", "unanchored_unbalanced", "realestate_no121")
SERVED_BOOKS = tuple(b for b in ALL_BOOKS if b not in REFUSED_BOOKS)

CENT = 0.01


class Refused(object):
    """A surface that REFUSED, with the code it carries (None = refused
    without a code, e.g. an absent figure)."""

    __slots__ = ("code",)

    def __init__(self, code: Optional[str]) -> None:
        self.code = code

    def __repr__(self) -> str:
        return "Refused(%r)" % (self.code,)


def _stub_benchmarks(key):  # noqa: ARG001 — the valuation's multiples table
    return {"industry_key_used": "generic", "industry_key_requested": key or "generic",
            "ev_ebitda": {"p25": 6.0, "p50": 8.0, "p75": 10.0, "source": "stub", "as_of_date": None},
            "ev_revenue": {"p25": 0.5, "p50": 0.8, "p75": 1.2, "source": "stub", "as_of_date": None}}


class _NoDb(object):
    def delete(self, *_a: Any, **_k: Any) -> None:
        return None

    def insert(self, *_a: Any, **_k: Any) -> None:
        return None


@contextlib.contextmanager
def _no_database() -> Iterator[_NoDb]:
    yield _NoDb()


def _get(book: Any) -> Dict[str, Any]:
    mp = MonkeyPatch()
    try:
        with ANCHOR._routed(book, mp) as (client, _db):
            resp = client.get("/api/period/%s" % book.period_id, headers=AUTH)
    finally:
        mp.undo()
    assert resp.status_code == 200, resp.text[:400]
    return resp.json()


def get_seeded(name: str, tables: Dict[str, List[Dict[str, Any]]]) -> Dict[str, Any]:
    """GET /api/period for one book with STORED rows seeded into the
    double (`calculated_metrics`, `valuations`, ...): what the route serves
    over rows a period persisted before a refusal existed. Each row is
    stamped with the book's period and org. The valuation multiples are
    the stub table (no benchmark I/O)."""
    from engine.api import _valuation as V

    book = _persisted(name)
    mp = MonkeyPatch()
    try:
        mp.setattr(V, "load_valuation_benchmarks", _stub_benchmarks)
        with ANCHOR._routed(book, mp) as (client, db):
            for table, rows in tables.items():
                db.tables[table] = [dict(r, period_id=book.period_id, org_id=book.org["id"])
                                    for r in rows]
            resp = client.get("/api/period/%s" % book.period_id, headers=AUTH)
    finally:
        mp.undo()
    assert resp.status_code == 200, resp.text[:400]
    return resp.json()


#: A derived book: (the corpus case it is copied from, the account prefix
#: whose rows are deleted from the copied FILE).
_DERIVED = {"realestate_no121": ("saga_10_col_realestate", "121")}
_DERIVED_DIRS: Dict[str, Path] = {}


def _derived_case_dir(name: str) -> Path:
    """A temporary copy of the corpus case with the rows whose account
    starts with the prefix deleted from the workbook itself. Nothing is
    written into the repository."""
    import shutil
    import tempfile

    import openpyxl

    if name in _DERIVED_DIRS:
        return _DERIVED_DIRS[name]
    case, prefix = _DERIVED[name]
    src = REPO / "corpus" / case
    out = Path(tempfile.mkdtemp(prefix="derived_%s_" % name)) / ("%s_%s" % (case, name))
    out.mkdir()
    shutil.copy(str(src / "meta.yaml"), str(out / "meta.yaml"))
    wb = openpyxl.load_workbook(str(src / "input.xlsx"))
    ws = wb.worksheets[0]
    doomed = [r for r in range(1, ws.max_row + 1)
              if str(ws.cell(r, 1).value or "").strip().startswith(prefix)]
    assert doomed, "%s: no %s row in %s — the witness would be vacuous" % (name, prefix, case)
    for r in reversed(doomed):
        ws.delete_rows(r)
    wb.save(str(out / "input.xlsx"))
    _DERIVED_DIRS[name] = out
    return out


def _persisted(name: str) -> Any:
    if name in CORPUS_BOOKS:
        case = "saga_10_col_%s" % name
        return ANCHOR._book(case, REPO / "corpus" / case)
    if name in _DERIVED:
        case_dir = _derived_case_dir(name)
        return ANCHOR._book(case_dir.name, case_dir)
    return N._persisted(name)


def _metric_rows(statements: Dict[str, Any]) -> List[Dict[str, Any]]:
    """The calculated_metrics rows the pipeline's stage_compute writes for
    these statements (the benchmark, the council and the ratio table's
    stored-row path read them)."""
    from engine.api import pipeline

    envelope = statements.get("assembled_canonical_v1") or {}
    original = pipeline._supabase.admin  # type: ignore[attr-defined]
    pipeline._supabase.admin = _no_database  # type: ignore[attr-defined]
    try:
        return pipeline.stage_compute(
            {"org_id": "gate"},
            {"statements": copy.deepcopy(statements),
             "source_data_quality": envelope.get("source_data_quality") or {}},
            "gate")
    finally:
        pipeline._supabase.admin = original  # type: ignore[attr-defined]


_CACHE: Dict[str, Any] = {}


def served(name: str) -> types.SimpleNamespace:
    """Everything the engine serves for one book, computed once."""
    if name in _CACHE:
        return _CACHE[name]
    from engine.api import _benchmark_engine as B
    from engine.api import _valuation as V
    from engine.confidence.reconciliation_checks import run_reconciliation_checks
    from engine.forecast.history import pl_history_from_payload
    from engine.ratios import credit_model
    from engine.serving.facts import FactsGateway

    persisted = _persisted(name)
    body = _get(persisted)
    statements = body["statements"]
    envelope = statements.get("assembled_canonical_v1") or {}
    rows = _metric_rows(statements)
    original = V.load_valuation_benchmarks
    V.load_valuation_benchmarks = _stub_benchmarks
    try:
        valuation = V.compute_valuation(industry_key=None, statements=copy.deepcopy(statements))
    finally:
        V.load_valuation_benchmarks = original
    bundle = types.SimpleNamespace(
        name=name,
        persisted=persisted,
        body=body,
        statements=statements,
        apl=statements["assembled_pl"],
        am=body.get("assembled_metrics") or {},
        envelope=envelope,
        rows=rows,
        metrics=dict((r["name"], r.get("value")) for r in rows),
        gateway=FactsGateway.from_envelope(envelope, currency=str(statements.get("currency") or "RON")),
        valuation=valuation,
        bench=B.compute_company_metrics(copy.deepcopy(rows), copy.deepcopy(body.get("line_items") or [])),
        history=pl_history_from_payload({"statements": statements, "envelope": envelope}),
        checks=dict((c.id, c) for c in run_reconciliation_checks(
            {"statements": statements, "assembled_canonical_v1": envelope,
             "source_data_quality": envelope.get("source_data_quality") or {}})),
        credit=credit_model.operating_figures(statements),
    )
    _CACHE[name] = bundle
    return bundle


def ratio_row(b: types.SimpleNamespace, key: str) -> Dict[str, Any]:
    for row in (b.am.get("ratio_table") or {}).get("rows") or []:
        if row.get("key") == key:
            return row
    raise AssertionError("%s: the served ratio table has no %r row" % (b.name, key))


def operand(b: types.SimpleNamespace, key: str, name: str) -> Any:
    for op in ratio_row(b, key).get("operands") or []:
        if op.get("name") == name:
            return op.get("value")
    raise AssertionError("%s: ratio row %r carries no operand %r" % (b.name, key, name))


def gateway_ebitda(b: types.SimpleNamespace) -> Any:
    from engine.serving.facts import MissingFactError

    try:
        return b.gateway.ebitda().amount_minor / 100.0
    except MissingFactError as err:  # RefusedFactError is a subclass
        return Refused(((getattr(err, "refusal", None) or {}).get("code")))


def methodology_ebitda(b: types.SimpleNamespace, view: str) -> Any:
    meth = b.envelope.get("methodology") or {}
    value = (meth.get("ebitda") or {}).get(view)
    if value is None:
        ref = (meth.get("refusals") or {}).get("ebitda.%s" % view) or {}
        return Refused(ref.get("code"))
    return value


def history_ebitda(b: types.SimpleNamespace) -> Any:
    h = b.history
    if h.ebitda is None:
        return Refused((h.ebitda_refusal or {}).get("code"))
    return h.ebitda / 100.0


def credit_ebitda(b: types.SimpleNamespace, key: str = "ebitda") -> Any:
    if b.credit.get(key) is None:
        return Refused(((b.credit.get("refusal") or {}).get("cause")
                        or (b.credit.get("refusal") or {}).get("code")))
    return b.credit[key]


def valuation_ebitda(b: types.SimpleNamespace) -> Any:
    if b.valuation.get("ebitda_used") is None:
        return Refused(((b.valuation.get("ebitda_refusal") or {}).get("cause")))
    return b.valuation["ebitda_used"]


def bench_ebitda(b: types.SimpleNamespace) -> Any:
    """The benchmark reads the STORED metric rows, which carry no refusal
    cause: it refuses with its own code `ebitda_refused` and a RO/EN
    sentence. Served as the refusal only when both are present."""
    if b.bench.get("ebitda") is None:
        ref = (b.bench.get("refusals") or {}).get("ebitda") or {}
        display = ref.get("display") or {}
        ok = ref.get("code") == "ebitda_refused" and display.get("ro") and display.get("en")
        return Refused("bench:ebitda_refused" if ok else None)
    return b.bench["ebitda"]


def recon_line(b: types.SimpleNamespace, key: str) -> Any:
    for line in (b.apl.get("ebitda_reconciliation") or {}).get("lines") or []:
        if line.get("key") == key:
            if line.get("value") is None:
                return Refused(((line.get("refusal") or {}).get("code")))
            return line["value"]
    raise AssertionError("%s: the reconciliation carries no %r line" % (b.name, key))


def bridge_sum(b: types.SimpleNamespace) -> Any:
    """The one-line bridge's parts, summed: EBITDA before + 711 + 72x."""
    parts = dict((p["key"], p.get("value")) for p in
                 ((b.apl.get("ebitda_reconciliation") or {}).get("bridge") or {}).get("parts") or [])
    if parts.get("ebitda") is None:
        return Refused("bridge_refused")
    return round(parts["ebitda_before_stock_variation"] + parts["inventory_variation"]
                 + parts["capitalized_own_work"], 2)


def rollup(b: types.SimpleNamespace) -> Any:
    """The engine's own reconciliation identity (confidence checks): the
    computed side of `ebitda_rollup`, or refused when no check is emitted
    (a refused EBITDA has no roll-up)."""
    check = b.checks.get("ebitda_rollup")
    if check is None:
        return Refused("no_rollup_on_a_refused_ebitda")
    return check.computed


#: Every ENGINE surface that carries THE ONE EBITDA (or its EBIT), read off
#: the same served period. The expectation is always `assembled_pl.ebitda`
#: (or `operating_result`) — never a figure computed here.
EBITDA_SURFACES: Tuple[Tuple[str, Callable[[types.SimpleNamespace], Any]], ...] = (
    ("assembled_pl.ebitda_statutory (alias)", lambda b: b.apl.get("ebitda_statutory")),
    ("assembled_pl.ebitda_operational (alias)", lambda b: b.apl.get("ebitda_operational")),
    ("assembled_pl.ebitda_operating_view (alias)", lambda b: b.apl.get("ebitda_operating_view")),
    ("assembled_pl.ebitda_cash (alias)", lambda b: b.apl.get("ebitda_cash")),
    ("assembled_pl.operating_ebitda (alias)", lambda b: b.apl.get("operating_ebitda")),
    ("assembled_pl.ebitda_reconciliation line 'ebitda'", lambda b: recon_line(b, "ebitda")),
    ("assembled_pl.ebitda_reconciliation bridge (before + 711 + 72x)", bridge_sum),
    ("GET /api/period assembled_metrics.pl.ebitda", lambda b: (b.am.get("pl") or {}).get("ebitda")),
    ("metric row 'ebitda' (stage_compute)", lambda b: b.metrics.get("ebitda")),
    ("metric row 'ebitda_statutory'", lambda b: b.metrics.get("ebitda_statutory")),
    ("metric row 'ebitda_cash'", lambda b: b.metrics.get("ebitda_cash")),
    ("ratio table Debt / EBITDA operand", lambda b: operand(b, "debt_to_ebitda", "ebitda")),
    ("credit model operating_figures.ebitda", credit_ebitda),
    ("methodology ebitda.reported", lambda b: methodology_ebitda(b, "reported")),
    ("methodology ebitda.cash", lambda b: methodology_ebitda(b, "cash")),
    ("FactsGateway.ebitda() (Capsule get_facts, advisory, radar)", gateway_ebitda),
    ("valuation ebitda_used (EV/EBITDA)", valuation_ebitda),
    ("benchmark company metrics ebitda (Section 9)", bench_ebitda),
    ("forecast PlHistory.ebitda (year 0)", history_ebitda),
    ("confidence check ebitda_rollup (computed side)", rollup),
)

EBIT_SURFACES: Tuple[Tuple[str, Callable[[types.SimpleNamespace], Any]], ...] = (
    ("assembled_pl.ebit", lambda b: b.apl.get("ebit")),
    ("assembled_pl.operating_ebit (alias)", lambda b: b.apl.get("operating_ebit")),
    ("assembled_pl.ebitda_reconciliation line 'operating_result'",
     lambda b: recon_line(b, "operating_result")),
    ("metric row 'operating_profit'", lambda b: b.metrics.get("operating_profit")),
    ("GET /api/period assembled_metrics.pl.ebit", lambda b: (b.am.get("pl") or {}).get("ebit")),
    ("credit model operating_figures.ebit", lambda b: credit_ebitda(b, "ebit")),
)
