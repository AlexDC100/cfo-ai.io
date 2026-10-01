#!/usr/bin/env python3
"""The two trial balances of the PUBLIC SAMPLE company — a FICTIONAL book.

WHAT THIS IS. The landing page links a sample anyone can inspect without
signing in: a trial balance and the complete report the engine produces
from it. The owner's rule (2026-10-01): fictional data only, never an
anonymised client book. So the company below does not exist. Its name says
so, its fiscal code fails the Romanian checksum (asserted here, so it can
match no real company), and every figure comes out of this file.

HOW THE BOOK IS BUILT. Not as a table of balances — as a LEDGER. A small
food manufacturer's year is written as aggregate double-entry postings
(sales, purchases, consumption, production stocked and sold, payroll, VAT,
depreciation, loans, provisions, income tax), posted onto an opening
balance sheet, and closed the way a Romanian accounting package closes it:
every class-6/7 account into account 121. The trial balance is then READ
off the ledger. Consequences, by construction and asserted below:

  · every column pair sums debit == credit;
  · the balance sheet closes;
  · account 121's closing balance is the year's net profit to the cent;
  · FY2025's opening column is FY2024's closing column, account by account;
  · account 711 is closed (cumulative debit == cumulative credit), so its
    NET — the stock variation — is not readable off the account, exactly
    as on a real closed book; it is what the engine's account-121 bridge
    has to recover (`STOCK_VARIATION` below states what it must find).

NO CLOCK, NO RANDOM STATE. The drivers are round planning figures; the
cents come from a hash of (SEED, year, key), so two runs — on any machine,
in any order — produce the same books. The workbook writer stores its zip
members uncompressed with fixed timestamps, so the .xlsx bytes are the
same on every platform (a deflate stream is not).

Run:  python scripts/build_public_sample_tb.py --out public/sample
      python scripts/build_public_sample_tb.py --print   (the ledger summary)

Python 3.9 — no ``match``, no ``X | Y``.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import re
import sys
import zipfile
from collections import OrderedDict, defaultdict
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

#: The fixed seed of the cents. Changing it changes every figure of the
#: public sample — and therefore every published file under public/sample.
SEED = 20261001

COMPANY_NAME = "CONSERVE EXEMPLU FICTIV SRL"
#: Deliberately INVALID: `assert_fiscal_code_is_invalid` proves the control
#: digit is wrong, so no real company can carry this code. All nines, like
#: the trade-register number below: the shape this repository's data-hygiene
#: gate (scripts/pdf_scrambler.py `is_placeholder_identifier`) recognises as
#: a fabricated placeholder rather than a registration.
FISCAL_CODE = "RO 99999999"
#: County code 99 does not exist; neither does the year.
TRADE_REGISTER = "J99/9999/2099"
CAEN = "1039"
CAEN_NAME_RO = "Prelucrarea și conservarea fructelor și legumelor n.c.a."
CAEN_NAME_EN = "Other processing and preserving of fruit and vegetables"
FICTIONAL_NOTICE_RO = (
    "EXEMPLU FICTIV - date fictive, generate pentru demonstrație; "
    "nu reprezintă nicio societate reală"
)

YEARS: Tuple[int, int] = (2024, 2025)

CENT = Decimal("0.01")
ZERO = Decimal("0.00")


def D(value) -> Decimal:
    return Decimal(str(value)).quantize(CENT, rounding=ROUND_HALF_UP)


def assert_fiscal_code_is_invalid(code: str = FISCAL_CODE) -> None:
    """The Romanian CUI control digit: weights 753217532 over the body
    (right-aligned), sum × 10 mod 11, 10 → 0. This code must FAIL it."""
    digits = re.sub(r"\D", "", code)
    body, control = digits[:-1], int(digits[-1])
    weights = "753217532"[-len(body):]
    total = sum(int(a) * int(b) for a, b in zip(body, weights))
    expected = (total * 10) % 11 % 10
    if expected == control:
        raise AssertionError(
            "the sample's fiscal code %s passes the CUI checksum — it could "
            "belong to a real company; pick a code that fails it" % code)


def _unit(year: int, key: str) -> float:
    """A stable number in [0, 1) for (SEED, year, key) — no RNG state."""
    digest = hashlib.sha256(("%d:%d:%s" % (SEED, year, key)).encode("utf-8")).digest()
    return int.from_bytes(digest[:8], "big") / float(1 << 64)


def amt(year: int, key: str, base: float, spread: float = 0.004) -> Decimal:
    """`base` moved by at most ±spread, to the cent — the planning figure
    with the irregular cents a real ledger has."""
    u = _unit(year, key)
    return D(base * (1.0 + (2.0 * u - 1.0) * spread))


# ── The chart ─────────────────────────────────────────────────────────
#
# Leaves only — the accounts a small manufacturer actually posts to. Names
# are the OMFP 1802/2014 chart's, in Romanian (never translated in a table).

CHART: "OrderedDict[str, str]" = OrderedDict([
    ("1012", "Capital subscris vărsat"),
    ("1061", "Rezerve legale"),
    ("1068", "Alte rezerve"),
    ("117", "Rezultatul reportat"),
    ("121", "Profit sau pierdere"),
    ("1518", "Alte provizioane"),
    ("1621", "Credite bancare pe termen lung"),
    ("167", "Alte împrumuturi și datorii asimilate (leasing)"),
    ("208", "Alte imobilizări necorporale (programe informatice)"),
    ("211", "Terenuri și amenajări de terenuri"),
    ("212", "Construcții"),
    ("2131", "Echipamente tehnologice (mașini, utilaje și instalații de lucru)"),
    ("2133", "Mijloace de transport"),
    ("214", "Mobilier, aparatură birotică, alte active corporale"),
    ("231", "Imobilizări corporale în curs de execuție"),
    ("2808", "Amortizarea altor imobilizări necorporale"),
    ("2812", "Amortizarea construcțiilor"),
    ("2813", "Amortizarea instalațiilor și mijloacelor de transport"),
    ("2814", "Amortizarea altor imobilizări corporale"),
    ("301", "Materii prime"),
    ("3022", "Combustibili"),
    ("3028", "Alte materiale consumabile"),
    ("303", "Materiale de natura obiectelor de inventar"),
    ("331", "Produse în curs de execuție"),
    ("345", "Produse finite"),
    ("371", "Mărfuri"),
    ("381", "Ambalaje"),
    ("401", "Furnizori"),
    ("404", "Furnizori de imobilizări"),
    ("408", "Furnizori - facturi nesosite"),
    ("409", "Furnizori - debitori"),
    ("4111", "Clienți"),
    ("4118", "Clienți incerți sau în litigiu"),
    ("419", "Clienți - creditori"),
    ("421", "Personal - salarii datorate"),
    ("4315", "Contribuția de asigurări sociale"),
    ("4316", "Contribuția de asigurări sociale de sănătate"),
    ("436", "Contribuția asiguratorie pentru muncă"),
    ("4411", "Impozitul pe profit curent"),
    ("4423", "TVA de plată"),
    ("4426", "TVA deductibilă"),
    ("4427", "TVA colectată"),
    ("444", "Impozitul pe venituri de natura salariilor"),
    ("446", "Alte impozite, taxe și vărsăminte asimilate"),
    ("457", "Dividende de plată"),
    ("462", "Creditori diverși"),
    ("471", "Cheltuieli înregistrate în avans"),
    ("4751", "Subvenții guvernamentale pentru investiții"),
    ("491", "Ajustări pentru deprecierea creanțelor - clienți"),
    ("5121", "Conturi la bănci în lei"),
    ("5124", "Conturi la bănci în valută"),
    ("5191", "Credite bancare pe termen scurt"),
    ("5311", "Casa în lei"),
    ("581", "Viramente interne"),
    ("601", "Cheltuieli cu materiile prime"),
    ("6022", "Cheltuieli privind combustibilii"),
    ("6028", "Cheltuieli privind alte materiale consumabile"),
    ("603", "Cheltuieli privind materialele de natura obiectelor de inventar"),
    ("604", "Cheltuieli privind materialele nestocate"),
    ("605", "Cheltuieli privind utilitățile"),
    ("607", "Cheltuieli privind mărfurile"),
    ("608", "Cheltuieli privind ambalajele"),
    ("611", "Cheltuieli cu întreținerea și reparațiile"),
    ("612", "Cheltuieli cu redevențele, locațiile de gestiune și chiriile"),
    ("613", "Cheltuieli cu primele de asigurare"),
    ("623", "Cheltuieli privind protocolul, reclama și publicitatea"),
    ("624", "Cheltuieli cu transportul de bunuri și personal"),
    ("625", "Cheltuieli cu deplasări, detașări și transferări"),
    ("626", "Cheltuieli poștale și taxe de telecomunicații"),
    ("627", "Cheltuieli cu serviciile bancare și asimilate"),
    ("628", "Alte cheltuieli cu serviciile executate de terți"),
    ("635", "Cheltuieli cu alte impozite, taxe și vărsăminte asimilate"),
    ("641", "Cheltuieli cu salariile personalului"),
    ("6422", "Cheltuieli cu tichetele de masă acordate salariaților"),
    ("646", "Cheltuieli privind contribuția asiguratorie pentru muncă"),
    ("654", "Pierderi din creanțe și debitori diverși"),
    ("6581", "Despăgubiri, amenzi și penalități"),
    ("6588", "Alte cheltuieli de exploatare"),
    ("665", "Cheltuieli din diferențe de curs valutar"),
    ("666", "Cheltuieli privind dobânzile"),
    ("6811", "Cheltuieli de exploatare privind amortizarea imobilizărilor"),
    ("6812", "Cheltuieli de exploatare privind provizioanele"),
    ("6814", "Cheltuieli de exploatare privind ajustările pentru deprecierea activelor circulante"),
    ("691", "Cheltuieli cu impozitul pe profit"),
    ("701", "Venituri din vânzarea produselor finite"),
    ("704", "Venituri din servicii prestate"),
    ("707", "Venituri din vânzarea mărfurilor"),
    ("709", "Reduceri comerciale acordate"),
    ("711", "Venituri aferente costurilor stocurilor de produse"),
    ("7584", "Venituri din subvenții pentru investiții"),
    ("7588", "Alte venituri din exploatare"),
    ("765", "Venituri din diferențe de curs valutar"),
    ("766", "Venituri din dobânzi"),
    ("7812", "Venituri din provizioane"),
    ("7814", "Venituri din ajustări pentru deprecierea activelor circulante"),
])

CLASS_TITLES = {
    "1": "Clasa 1 - Conturi de capitaluri, provizioane, împrumuturi și datorii asimilate",
    "2": "Clasa 2 - Conturi de imobilizări",
    "3": "Clasa 3 - Conturi de stocuri și producție în curs de execuție",
    "4": "Clasa 4 - Conturi de terți",
    "5": "Clasa 5 - Conturi de trezorerie",
    "6": "Clasa 6 - Conturi de cheltuieli",
    "7": "Clasa 7 - Conturi de venituri",
}

# ── The opening balance sheet, 1 January 2024 ─────────────────────────
#
# Signed: debit positive, credit negative. Balanced (asserted). Account 121
# opens with the FY2023 result, which the first year transfers to 117 —
# the entry the engine's guard G6 looks for.

OPENING_2024: Dict[str, str] = {
    "1012": "-400000.00",
    "1061": "-80000.00",
    "1068": "-210000.00",
    "117": "-1645944.27",
    "121": "-287356.73",
    "1621": "-1420000.00",
    "167": "-268400.00",
    "4751": "-252000.00",
    "208": "48200.00",
    "2808": "-30150.00",
    "211": "310000.00",
    "212": "1640000.00",
    "2812": "-410000.00",
    "2131": "2180460.00",
    "2133": "420300.00",
    "2813": "-1020760.00",
    "214": "96400.00",
    "2814": "-58450.00",
    "301": "240318.40",
    "3022": "6120.00",
    "3028": "21044.15",
    "303": "14210.30",
    "345": "1010462.85",
    "371": "48133.60",
    "381": "96710.70",
    "4111": "1180254.36",
    "4118": "34100.00",
    "491": "-20460.00",
    "409": "18000.00",
    "471": "12400.00",
    "5121": "285906.44",
    "5124": "38112.50",
    "5311": "5180.70",
    "5191": "-520000.00",
    "401": "-812344.18",
    "404": "-46210.00",
    "408": "-22000.00",
    "419": "-9000.00",
    "421": "-96120.00",
    "4315": "-38450.00",
    "4316": "-15380.00",
    "444": "-9410.00",
    "436": "-3460.00",
    "4411": "-14230.00",
    "4423": "-6380.00",
    "446": "-3790.00",
    "462": "-6018.82",
}

# ── The year's drivers (planning figures, before the cents) ───────────
#
# One column per year. Everything the ledger posts is one of these, a
# stated share of one, or a payment plugged to a stated closing balance.

DRIVERS: Dict[str, Tuple[float, float]] = {
    # sales, net of VAT
    "sales_701": (7_940_000, 8_870_000),
    "sales_707": (540_000, 610_000),
    "sales_704": (72_000, 86_000),
    "discounts_709": (168_000, 205_000),
    # purchases into stock
    "buy_301": (3_150_000, 3_540_000),
    "buy_381": (610_000, 668_000),
    "buy_3022": (118_000, 131_000),
    "buy_3028": (96_000, 104_000),
    "buy_303": (38_000, 41_000),
    "buy_371": (421_000, 478_000),
    # consumption out of stock
    "use_601": (3_120_000, 3_495_000),
    "use_608": (598_000, 655_000),
    "use_6028": (94_000, 107_000),
    "use_603": (36_000, 43_000),
    "use_607": (415_000, 468_000),
    # production: cost stocked into 345, cost of the goods sold out of it
    "production_cost": (5_760_000, 6_350_000),
    "cost_of_goods_sold": (5_798_000, 6_205_000),
    "wip_year_end": (24_000, 31_000),
    # services and other operating costs
    "exp_604": (22_000, 26_000),
    "exp_605": (418_000, 476_000),
    "exp_611": (96_000, 118_000),
    "exp_612": (132_000, 144_000),
    "exp_613": (36_000, 39_000),
    "exp_623": (38_000, 52_000),
    "exp_624": (262_000, 301_000),
    "exp_625": (14_000, 17_000),
    "exp_626": (16_000, 17_500),
    "exp_627": (19_000, 21_500),
    "exp_628": (148_000, 171_000),
    "exp_635": (41_000, 44_000),
    "payroll_641": (1_690_000, 1_912_000),
    "vouchers_6422": (142_000, 163_000),
    "exp_6581": (2_400, 3_100),
    "exp_6588": (9_000, 11_500),
    "depr_2812": (41_000, 41_000),
    "depr_2813": (318_000, 344_000),
    "depr_2814": (7_000, 8_500),
    "depr_2808": (6_000, 7_500),
    # provisions and allowances
    "provision_6812": (28_000, 40_000),
    "provision_reversed_7812": (0, 18_000),
    "allowance_6814": (22_000, 19_000),
    "doubtful_to_4118": (20_000, 16_000),
    "written_off_654": (0, 14_200),
    # other income, financial
    "subsidy_7584": (36_000, 36_000),
    "other_income_7588": (12_000, 15_500),
    "fx_loss_665": (5_200, 6_900),
    "fx_gain_765": (3_900, 4_300),
    "interest_666": (171_000, 164_000),
    "interest_income_766": (1_800, 2_600),
    # investment and financing
    "capex_231": (0, 508_000),
    "cip_to_2131": (0, 412_000),
    "capex_2131_direct": (186_000, 0),
    "capex_2133": (0, 118_000),
    "capex_214": (12_000, 0),
    "capex_208": (0, 14_000),
    "scrapped_2133": (0, 60_000),
    "loan_drawn_1621": (0, 420_000),
    "loan_repaid_1621": (184_000, 196_000),
    "lease_repaid_167": (64_000, 63_000),
    "dividends": (120_000, 100_000),
    # closing balances the payments are plugged to
    "close_4111": (1_262_000, 1_371_000),
    "close_401": (868_000, 934_000),
    "close_404": (31_000, 58_000),
    "close_408": (24_500, 27_500),
    "close_419": (12_500, 7_000),
    "close_4411": (18_400, 24_900),
    "close_4423": (7_900, 9_200),
    "close_446": (4_100, 4_600),
    "close_5191": (470_000, 450_000),
    "close_5311": (4_600, 5_900),
    "close_471": (13_200, 14_100),
    "close_409": (15_000, 21_000),
}

#: VAT actually charged, as a blended share of the year's base (the rates
#: changed on 1 August 2025; the ledger only needs the flow to balance).
VAT_FOOD = (0.09, 0.0983)
VAT_STANDARD = (0.19, 0.1983)
#: Dividend tax withheld.
DIVIDEND_TAX = (0.08, 0.10)
INCOME_TAX_RATE = 0.16


class Ledger:
    """A year of postings over an opening balance sheet."""

    def __init__(self, year: int, opening: Dict[str, Decimal]) -> None:
        self.year = year
        self.opening = {k: v for k, v in opening.items() if v != ZERO}
        self.debit: Dict[str, Decimal] = defaultdict(lambda: ZERO)
        self.credit: Dict[str, Decimal] = defaultdict(lambda: ZERO)

    def post(self, debit: str, credit: str, amount: Decimal) -> None:
        amount = D(amount)
        if amount < ZERO:
            raise ValueError("negative posting %s / %s %s" % (debit, credit, amount))
        for code in (debit, credit):
            if code not in CHART:
                raise KeyError("account %s is not in the sample's chart" % code)
        if amount == ZERO:
            return
        self.debit[debit] += amount
        self.credit[credit] += amount

    def balance(self, code: str) -> Decimal:
        """Signed: debit positive."""
        return self.opening.get(code, ZERO) + self.debit[code] - self.credit[code]

    def credit_balance(self, code: str) -> Decimal:
        return -self.balance(code)

    def accounts(self) -> List[str]:
        used = set(self.opening) | set(self.debit) | set(self.credit)
        return [c for c in CHART if c in used]

    def closing(self) -> Dict[str, Decimal]:
        return {c: self.balance(c) for c in self.accounts() if self.balance(c) != ZERO}


def _driver(year: int, key: str) -> float:
    return DRIVERS[key][YEARS.index(year)]


def run_year(year: int, opening: Dict[str, Decimal]) -> Tuple[Ledger, Dict[str, Decimal]]:
    """Post the year; return the ledger and the facts the book was built
    to carry (net profit, the stock variation the engine must recover)."""
    yi = YEARS.index(year)
    L = Ledger(year, opening)

    def a(key: str, spread: float = 0.004) -> Decimal:
        return amt(year, key, _driver(year, key), spread)

    vat_food, vat_std = VAT_FOOD[yi], VAT_STANDARD[yi]

    def buy(debit: str, base: Decimal, rate: float, supplier: str = "401") -> None:
        """A supplier invoice: the cost, and its deductible VAT."""
        L.post(debit, supplier, base)
        L.post("4426", supplier, D(float(base) * rate))

    # 1 ── last year's result leaves 121 (the entry guard G6 reads), then
    #      part of it is distributed.
    prior_result = L.credit_balance("121")
    L.post("121", "117", prior_result)
    dividends = a("dividends", 0.0)
    L.post("117", "457", dividends)
    dividend_tax = D(float(dividends) * DIVIDEND_TAX[yi])
    L.post("457", "446", dividend_tax)
    L.post("457", "5121", dividends - dividend_tax)

    # 2 ── sales, and the commercial reductions granted on them (709).
    s701, s707, s704 = a("sales_701"), a("sales_707"), a("sales_704")
    L.post("4111", "701", s701)
    L.post("4111", "707", s707)
    L.post("4111", "704", s704)
    L.post("4111", "4427", D(float(s701) * vat_food + float(s707 + s704) * vat_std))
    discounts = a("discounts_709")
    L.post("709", "4111", discounts)
    L.post("4427", "4111", D(float(discounts) * vat_food))

    # 3 ── customers: advances, doubtful balances, the allowance, a write-off.
    L.post("419", "4111", L.credit_balance("419"))            # last year's advances settle
    L.post("5121", "419", a("close_419", 0.0))                # this year's are received
    L.post("4118", "4111", a("doubtful_to_4118"))
    L.post("6814", "491", a("allowance_6814"))
    written_off = a("written_off_654")
    L.post("654", "4118", written_off)
    L.post("491", "7814", written_off)                        # its allowance is released

    # 4 ── purchases into stock, and what production took out of it.
    buy("301", a("buy_301"), vat_food)
    buy("381", a("buy_381"), vat_std)
    fuel = a("buy_3022")
    buy("3022", fuel, vat_std)
    buy("3028", a("buy_3028"), vat_std)
    buy("303", a("buy_303"), vat_std)
    buy("371", a("buy_371"), vat_food)
    L.post("601", "301", a("use_601"))
    L.post("608", "381", a("use_608"))
    L.post("6022", "3022", fuel)                              # bought and burnt in the year
    L.post("6028", "3028", a("use_6028"))
    L.post("603", "303", a("use_603"))
    L.post("607", "371", a("use_607"))

    # 5 ── production: the cost stocked (345 / 711), the cost of what was
    #      sold (711 / 345), and work in progress at the year end.
    opening_wip = L.balance("331")
    L.post("711", "331", opening_wip)
    L.post("345", "711", a("production_cost"))
    L.post("711", "345", a("cost_of_goods_sold"))
    L.post("331", "711", a("wip_year_end"))

    # 6 ── services and other running costs.
    buy("604", a("exp_604"), vat_std)
    utilities = a("exp_605")
    opening_accrual = L.credit_balance("408")                 # last December's invoice arrives
    L.post("408", "401", opening_accrual)
    L.post("4426", "401", D(float(opening_accrual) * vat_std))
    accrued = a("close_408", 0.0)                             # this December, not yet invoiced
    buy("605", utilities - accrued, vat_std)
    L.post("605", "408", accrued)
    buy("611", a("exp_611"), vat_std)
    buy("612", a("exp_612"), vat_std)
    insurance_paid = a("exp_613") + a("close_471", 0.0) - L.balance("471")
    L.post("471", "401", insurance_paid)                      # paid ahead, expensed over time
    L.post("613", "471", a("exp_613"))
    buy("623", a("exp_623"), vat_std)
    buy("624", a("exp_624"), vat_std)
    L.post("625", "5311", a("exp_625"))
    buy("626", a("exp_626"), vat_std)
    L.post("627", "5121", a("exp_627"))
    buy("628", a("exp_628"), vat_std)
    L.post("6422", "401", a("vouchers_6422"))
    L.post("635", "446", a("exp_635"))
    L.post("6581", "5121", a("exp_6581"))
    L.post("6588", "5121", a("exp_6588"))
    L.post("5121", "7588", a("other_income_7588"))

    # 7 ── payroll.
    gross = a("payroll_641")
    cas = D(float(gross) * 0.25)
    cass = D(float(gross) * 0.10)
    wage_tax = D(float(gross - cas - cass) * 0.094)
    cam = D(float(gross) * 0.0225)
    L.post("641", "421", gross)
    L.post("421", "4315", cas)
    L.post("421", "4316", cass)
    L.post("421", "444", wage_tax)
    L.post("646", "436", cam)
    twelfth = Decimal(1) / Decimal(12)
    for code, yearly in (("4315", cas), ("4316", cass), ("444", wage_tax), ("436", cam)):
        L.post(code, "5121", L.credit_balance(code) - D(yearly * twelfth))
    net_pay = gross - cas - cass - wage_tax
    L.post("421", "5121", L.credit_balance("421") - D(net_pay * twelfth))

    # 8 ── fixed assets: depreciation, investment, one scrapped vehicle.
    for code in ("2812", "2813", "2814", "2808"):
        L.post("6811", code, a("depr_%s" % code))
    buy("231", a("capex_231"), vat_std, supplier="404")
    L.post("2131", "231", a("cip_to_2131"))
    buy("2131", a("capex_2131_direct"), vat_std, supplier="404")
    buy("2133", a("capex_2133"), vat_std, supplier="404")
    buy("214", a("capex_214"), vat_std, supplier="404")
    buy("208", a("capex_208"), vat_std, supplier="404")
    L.post("2813", "2133", a("scrapped_2133", 0.0))           # fully depreciated
    L.post("4751", "7584", a("subsidy_7584", 0.0))            # the grant, over the asset's life

    # 9 ── provisions (a labour dispute, a customer claim).
    L.post("6812", "1518", a("provision_6812"))
    L.post("1518", "7812", a("provision_reversed_7812"))

    # 10 ── financial: interest, the EUR account, the loans.
    L.post("666", "5121", a("interest_666"))
    L.post("5121", "766", a("interest_income_766"))
    L.post("665", "5124", a("fx_loss_665"))
    L.post("5124", "765", a("fx_gain_765"))
    L.post("5121", "1621", a("loan_drawn_1621", 0.0))
    L.post("1621", "5121", a("loan_repaid_1621", 0.0))
    L.post("167", "5121", a("lease_repaid_167", 0.0))
    credit_line = L.credit_balance("5191")
    target_line = a("close_5191", 0.0)
    # the line revolves: drawn and repaid several times over in a year
    L.post("5121", "5191", D(float(target_line) * 2.6))
    L.post("5191", "5121", credit_line + D(float(target_line) * 2.6) - target_line)

    # 11 ── settle the trading accounts to their closing balances.
    advances_move = a("close_409", 0.0) - L.balance("409")
    if advances_move >= ZERO:
        L.post("409", "5121", advances_move)                  # paid ahead to suppliers
    else:
        L.post("401", "409", -advances_move)                  # an advance met its invoice
    cash_sales = D(float(s707) * 0.18)
    L.post("5311", "4111", cash_sales)
    L.post("5121", "4111", L.balance("4111") - a("close_4111"))
    banked = L.balance("5311") - a("close_5311")
    L.post("581", "5311", banked)
    L.post("5121", "581", banked)
    L.post("404", "5121", L.credit_balance("404") - a("close_404"))
    L.post("401", "5121", L.credit_balance("401") - a("close_401"))
    L.post("446", "5121", L.credit_balance("446") - a("close_446"))

    # 12 ── VAT: deductible against collected, the difference payable.
    collected = L.credit_balance("4427")
    deductible = L.balance("4426")
    if collected <= deductible:
        raise AssertionError("FY%d: VAT collected %s does not exceed deductible %s — "
                             "the sample has no 4424 account" % (year, collected, deductible))
    L.post("4427", "4426", deductible)
    L.post("4427", "4423", collected - deductible)
    L.post("4423", "5121", L.credit_balance("4423") - a("close_4423"))

    # 13 ── income tax on the year's profit (penalties are not deductible).
    pre_tax = -sum((L.balance(c) for c in L.accounts() if c[0] in "67"), ZERO)
    income_tax = D(float(pre_tax + L.balance("6581")) * INCOME_TAX_RATE)
    L.post("691", "4411", income_tax)
    L.post("4411", "5121", L.credit_balance("4411") - a("close_4411"))

    # 14 ── the closing entries: every class-6/7 account into 121.
    stock_variation = -L.balance("711")                       # credit − debit, before closing
    for code in [c for c in L.accounts() if c[0] in "67"]:
        bal = L.balance(code)
        if bal > ZERO:
            L.post("121", code, bal)
        elif bal < ZERO:
            L.post(code, "121", -bal)

    facts = {
        "net_profit": L.credit_balance("121"),
        "stock_variation_711": stock_variation,
        "pre_tax_profit": pre_tax,
        "income_tax": income_tax,
        "net_turnover": s701 + s707 + s704 - discounts,
        "prior_result_transferred": prior_result,
    }
    return L, facts


# ── The trial balance, read off the ledger ────────────────────────────

Row = Tuple[str, str, Decimal, Decimal, Decimal, Decimal, Decimal, Decimal, Decimal, Decimal]


def _sides(signed: Decimal) -> Tuple[Decimal, Decimal]:
    return (signed, ZERO) if signed >= ZERO else (ZERO, -signed)


def trial_balance(L: Ledger) -> List[Row]:
    rows: List[Row] = []
    for code in L.accounts():
        si_d, si_c = _sides(L.opening.get(code, ZERO))
        r_d, r_c = L.debit[code], L.credit[code]
        t_d, t_c = si_d + r_d, si_c + r_c
        sf_d, sf_c = _sides(t_d - t_c)
        if not any((si_d, si_c, r_d, r_c)):
            continue
        rows.append((code, CHART[code], si_d, si_c, r_d, r_c, t_d, t_c, sf_d, sf_c))
    return rows


def build_books() -> "OrderedDict[int, Dict[str, object]]":
    """Both years: {year: {"rows": [...], "facts": {...}, "ledger": Ledger}}."""
    assert_fiscal_code_is_invalid()
    opening = {code: D(value) for code, value in OPENING_2024.items()}
    if sum(opening.values(), ZERO) != ZERO:
        raise AssertionError("the 1 January 2024 balance sheet does not balance: %s"
                             % sum(opening.values(), ZERO))
    books: "OrderedDict[int, Dict[str, object]]" = OrderedDict()
    for year in YEARS:
        ledger, facts = run_year(year, opening)
        rows = trial_balance(ledger)
        _assert_book(year, ledger, rows, facts)
        books[year] = {"rows": rows, "facts": facts, "ledger": ledger}
        opening = ledger.closing()
    return books


#: Accounts that may only ever carry a balance on one side.
_DEBIT_NATURED = ("2131", "2133", "211", "212", "214", "208", "231", "301", "3022", "3028", "303",
                  "331", "345", "371", "381", "4111", "4118", "409", "471", "5121", "5124", "5311")
_CREDIT_NATURED = ("1012", "1061", "1068", "117", "1518", "1621", "167", "4751", "2808", "2812",
                   "2813", "2814", "401", "404", "408", "419", "421", "4315", "4316", "436", "444",
                   "4411", "4423", "446", "462", "491", "5191")


def _assert_book(year: int, L: Ledger, rows: Sequence[Row], facts: Dict[str, Decimal]) -> None:
    def col(i: int) -> Decimal:
        return sum((r[i] for r in rows), ZERO)

    for name, d_i in (("opening", 2), ("movement", 4), ("cumulative", 6), ("closing", 8)):
        if col(d_i) != col(d_i + 1):
            raise AssertionError("FY%d %s columns: debit %s != credit %s"
                                 % (year, name, col(d_i), col(d_i + 1)))
    for code in _DEBIT_NATURED:
        if L.balance(code) < ZERO:
            raise AssertionError("FY%d: %s closes on the credit side (%s)" % (year, code, L.balance(code)))
    for code in _CREDIT_NATURED:
        if L.balance(code) > ZERO:
            raise AssertionError("FY%d: %s closes on the debit side (%s)" % (year, code, L.balance(code)))
    for code in L.accounts():
        if code[0] in "67" and (L.balance(code) != ZERO or L.debit[code] != L.credit[code]):
            raise AssertionError("FY%d: %s is not closed into 121" % (year, code))
    for code in ("4426", "4427", "581", "457"):
        if L.balance(code) != ZERO:
            raise AssertionError("FY%d: %s should close at zero (%s)" % (year, code, L.balance(code)))
    if facts["net_profit"] <= ZERO:
        raise AssertionError("FY%d: the sample company should be profitable" % year)


# ── The workbook ──────────────────────────────────────────────────────

SHEET_NAME = "Balanta"
GROUP_HEADER = (None, None, "Solduri inițiale an", None, "Rulaje cumulate", None,
                "Sume totale", None, "Solduri finale", None)
COLUMN_HEADER = ("Cont", "Denumire cont", "Debit", "Credit", "Debit", "Credit",
                 "Debit", "Credit", "Debit", "Credit")


def title_rows(year: int, notice: str = FICTIONAL_NOTICE_RO, width: int = 10) -> List[Tuple[object, ...]]:
    pad = (None,) * (width - 1)
    return [
        (COMPANY_NAME,) + pad,
        ("CUI %s   Nr. Reg. Com. %s   CAEN %s" % (FISCAL_CODE, TRADE_REGISTER, CAEN),) + pad,
        ("Balanță de verificare   Perioada: 01.01.%d - 31.12.%d" % (year, year),) + pad,
        (notice,) + pad,
        (None,) * width,
    ]


def sheet_rows(year: int, rows: Sequence[Row], notice: str = FICTIONAL_NOTICE_RO) -> List[Tuple[object, ...]]:
    """Every row of the four-pair sheet, top to bottom (titles, headers,
    class headings, accounts, the totals line)."""
    out: List[Tuple[object, ...]] = title_rows(year, notice)
    out.append(GROUP_HEADER)
    out.append(COLUMN_HEADER)
    current = ""
    for row in rows:
        cls = row[0][0]
        if cls != current:
            current = cls
            out.append((CLASS_TITLES[cls],) + (None,) * 9)
        out.append((row[0], row[1]) + tuple(float(v) for v in row[2:]))
    totals = tuple(float(sum((r[i] for r in rows), ZERO)) for i in range(2, 10))
    out.append((None, "TOTALURI") + totals)
    return out


# The COMPACT layout: three Debitoare/Creditoare pairs (opening, movement,
# cumulative) and no closing-balance pair — a reader derives the closing
# balance from the cumulative pair. A "Total sume clasa N" line closes each
# class and "Totaluri:" closes the sheet.
COMPACT_GROUP_HEADER = (None, None, "Solduri inițiale an", None, "Rulaje perioada", None,
                        "Sume totale", None)
COMPACT_COLUMN_HEADER = ("Cont", "Denumirea contului", "Debitoare", "Creditoare",
                         "Debitoare", "Creditoare", "Debitoare", "Creditoare")


def compact_sheet_rows(year: int, rows: Sequence[Row],
                       notice: str = FICTIONAL_NOTICE_RO) -> List[Tuple[object, ...]]:
    out: List[Tuple[object, ...]] = title_rows(year, notice, width=8)
    out.append(COMPACT_GROUP_HEADER)
    out.append(COMPACT_COLUMN_HEADER)

    def class_total(cls: str) -> Tuple[object, ...]:
        members = [r for r in rows if r[0][0] == cls]
        return ("Total sume clasa %s" % cls, None, None, None, None, None,
                float(sum((r[6] for r in members), ZERO)), float(sum((r[7] for r in members), ZERO)))

    current = ""
    for row in rows:
        cls = row[0][0]
        if current and cls != current:
            out.append(class_total(current))
        current = cls
        out.append((row[0], row[1]) + tuple(float(v) for v in row[2:8]))
    if current:
        out.append(class_total(current))
    totals = tuple(float(sum((r[i] for r in rows), ZERO)) for i in range(2, 8))
    out.append(("Totaluri:", None) + totals)
    return out


#: Fixed document dates — the workbook says when the BOOK closes, not when
#: this script ran.
_DOC_DATE = "%d-12-31T00:00:00Z"


def write_workbook(year: int, rows: Sequence[Row], *, notice: str = FICTIONAL_NOTICE_RO,
                   compact: bool = False) -> bytes:
    """The .xlsx, byte-stable: openpyxl writes the parts, then the archive
    is rebuilt uncompressed with fixed member timestamps and the document
    properties' dates pinned. `compact` writes the three-pair layout."""
    import datetime

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    grid = compact_sheet_rows(year, rows, notice) if compact else sheet_rows(year, rows, notice)
    width = 8 if compact else 10
    for line in grid:
        ws.append(list(line))
    header_row = len(title_rows(year)) + 2
    bold = Font(bold=True)
    for cell in ws[1] + ws[header_row - 1] + ws[header_row] + ws[len(grid)]:
        cell.font = bold
    for r_i, line in enumerate(grid, start=1):
        if r_i > header_row and line[1] is None and line[0]:
            ws.cell(row=r_i, column=1).font = bold
        for c_i in range(3, width + 1):
            cell = ws.cell(row=r_i, column=c_i)
            if isinstance(cell.value, float):
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right")
    for c_i, col_width in enumerate((9, 62) + (16,) * (width - 2), start=1):
        ws.column_dimensions[get_column_letter(c_i)].width = col_width
    for first in range(3, width, 2):
        ws.merge_cells(start_row=header_row - 1, start_column=first,
                       end_row=header_row - 1, end_column=first + 1)
        ws.cell(row=header_row - 1, column=first).alignment = Alignment(horizontal="center")
    ws.freeze_panes = ws.cell(row=header_row + 1, column=3)
    stamp = datetime.datetime(year, 12, 31)
    wb.properties.creator = "CFO AI - exemplu fictiv"
    wb.properties.lastModifiedBy = "CFO AI - exemplu fictiv"
    wb.properties.title = "%s - balanță de verificare %d (date fictive)" % (COMPANY_NAME, year)
    wb.properties.created = stamp
    wb.properties.modified = stamp
    raw = io.BytesIO()
    wb.save(raw)
    return _stable_zip(raw.getvalue(), _DOC_DATE % year)


def _stable_zip(data: bytes, doc_date: str) -> bytes:
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_STORED) as dst:
        for name in sorted(src.namelist(), key=lambda n: (n != "[Content_Types].xml", n)):
            body = src.read(name)
            if name == "docProps/core.xml":
                text = body.decode("utf-8")
                text = re.sub(r"(<dcterms:(created|modified)[^>]*>)[^<]*(</dcterms:\2>)",
                              lambda m: m.group(1) + doc_date + m.group(3), text)
                body = text.encode("utf-8")
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_STORED
            info.external_attr = 0o644 << 16
            info.create_system = 3
            dst.writestr(info, body)
    return out.getvalue()


def workbook_filename(year: int) -> str:
    return "balanta_exemplu_fictiv_31.12.%d.xlsx" % year


# ── CLI ───────────────────────────────────────────────────────────────

def _summary(books) -> str:
    lines = []
    for year, book in books.items():
        L: Ledger = book["ledger"]  # type: ignore[assignment]
        facts: Dict[str, Decimal] = book["facts"]  # type: ignore[assignment]
        lines.append("FY%d  accounts %d" % (year, len(book["rows"])))  # type: ignore[arg-type]
        for key in ("net_turnover", "pre_tax_profit", "income_tax", "net_profit",
                    "stock_variation_711", "prior_result_transferred"):
            lines.append("  %-26s %16s" % (key, "{:,.2f}".format(facts[key])))
        for code in ("5121", "5124", "5311", "5191", "1621", "167", "4111", "401", "345", "301", "117"):
            lines.append("  closing %-6s %18s" % (code, "{:,.2f}".format(L.balance(code))))
    return "\n".join(lines)


def main(argv: Optional[List[str]] = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--out", help="directory to write the two workbooks into")
    parser.add_argument("--print", action="store_true", dest="show",
                        help="print the ledger summary")
    args = parser.parse_args(argv)
    books = build_books()
    if args.show or not args.out:
        print(_summary(books))
    if args.out:
        out_dir = Path(args.out)
        out_dir.mkdir(parents=True, exist_ok=True)
        for year, book in books.items():
            path = out_dir / workbook_filename(year)
            path.write_bytes(write_workbook(year, book["rows"]))  # type: ignore[arg-type]
            print("wrote %s" % path)
    return 0


if __name__ == "__main__":
    sys.exit(main())
