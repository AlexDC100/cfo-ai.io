"""FIRM COCKPIT — the file-request flow, per-client cadence, digest
preferences + cron, and the firm email queue (engine.api._firm_requests).

THE REQUEST-FILE FLOW, and the one property it must keep
  An accountant asks a client for ONE period's trial balance. The ask is
  a SIGNED, EXPIRING, SINGLE-USE link bound to (request, client, period).
  The link does not create a side channel: the file it delivers lands
  through the NORMAL upload pipeline — the same `documents` row shape
  `lib/supabase.ts::uploadDocument` writes, the same storage path, the
  same `_admin_set_status("queued") + _enqueue(document_id)` the
  `/api/pipeline/run` route calls — with the request's period as the
  human-confirmed `period_end_hint` (the accountant chose the month when
  minting the link), so `stage_persist`'s detection-vs-hint record and
  the PERIOD_MISMATCH surface work exactly as for any other upload.

  On top of the pipeline's own guards, the landing adds one the pipeline
  never had: an ENTITY guard. The document's preamble (the rows above
  the column header the parser locates) is read for a company name and a
  CUI and compared to the identity the request was minted against. A
  MISMATCH refuses the landing (409, recorded on the request); an UNKNOWN
  (the file carries no identity — anonymised exports have none) is
  recorded and lets the file through: absence of evidence is not
  evidence of the wrong company. FC7 (tests/engine/test_firm_gates.py)
  plants a wrong-period and a wrong-entity file and proves both guards
  fire.

TOKENS
  base64url(json claims) + "." + base64url(HMAC-SHA256(claims)) under
  FIRM_REQUEST_SIGNING_KEY. Only the token's SHA-256 is stored, so the
  table cannot mint an upload. No key → 503, fail closed, on every route
  that mints or verifies (the purge-cron precedent in _org.py).

EMAIL is a QUEUED STUB — `firm_email_queue`, the renewal_email_queue
  pattern — drained by POST /api/firm/email/drain through the branded
  templates below; nothing here calls a mail provider inline.

DIGEST preferences are per user (+ firm); `enabled` defaults to FALSE and
  the cron queues nothing for a user who never opted in.

TWO WALLS, NEVER ONE — ON READS AND ON WRITES (the tenancy lane's
  invariant, held here too). Every route a signed-in user reads client
  data through — /requests/list, /requests/{id}/revoke, /cadence/status,
  /digest/prefs, and the brief's report provider — checks
  `authorize_client` (the Python wall: firm membership, the role's cell,
  client-of-this-firm) AND reads through the CALLER's own Supabase
  client (`_supabase.per_user`), so the firm-read policies in
  schema_phase_firm_requests.sql are the second wall. Every WRITE a
  signed-in user makes — POST /requests (the request row), the
  revocation, PUT /cadence, PUT /digest/prefs — goes through the CALLER's
  own client too, under the INSERT / UPDATE policies of the same
  migration; an RLS refusal is answered 403 by `write_as_caller`, never
  swallowed, never a 500. A firm-B owner whose Python wall was removed
  by hand cannot overwrite firm-A's client's cadence: the policy pins the
  row to `firm_of_org(client_org_id)` and the role's `request_file` cell
  (W1). The service-role client is used only where NO user exists: the
  two crons (ENGINE_API_TOKEN bearer, fail closed), the public token
  landing (the signed token is the auth), the admin-only e-mail drain,
  the secret `firm_file_request_tokens` table and the e-mail queue (no
  policy at all — written only AFTER the caller's own insert of the
  request row was admitted). tests/engine/test_firm_tenancy.py drives
  every one of these routes cross-firm against the policies parsed out
  of the migration, with the Python wall removed by hand for every
  write, and shows the write lands only once the policy is loosened.

Python 3.9 — no `match`, no `X | Y` unions.
"""
from __future__ import annotations

import base64
import csv
import hashlib
import hmac
import html as _html
import io
import json
import logging
import os
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta, timezone
from typing import Any, Callable, Dict, List, Optional, Sequence, Tuple

from fastapi import APIRouter, File, Header, HTTPException, Query, UploadFile
from pydantic import BaseModel, Field

from engine.firm import cadence as CAD
from engine.firm import digest as DG

from . import _email, _email_templates, _org, _period_detect, _supabase
# EVERY list read below walks pages (critic D11): a bare select answers at
# most the deployment's db-max-rows, silently — the crons and the brief
# covered 1,000 clients of a 1,200-client firm. The walk stops on an EMPTY
# page (engine.api._paging), never on the cap it does not know.
from ._paging import chunked, select_all

logger = logging.getLogger(__name__)

#: Contract version for the /api/firm request / cadence / digest surface.
FIRM_REQUESTS_API_VERSION = "firm-d1"

#: Route string the frontend renders for a request link (it owns the page).
ROUTE_UPLOAD_REQUEST = "/upload/request/{token}"

SIGNING_KEY_ENV = "FIRM_REQUEST_SIGNING_KEY"
TOKEN_VERSION = "fr1"
DOC_BUCKET = "documents"
MAX_UPLOAD_BYTES = 25 * 1024 * 1024

STATUS_REQUESTED = "requested"
STATUS_REMINDED = "reminded"
STATUS_RECEIVED = "received"
STATUS_EXPIRED = "expired"
STATUS_REVOKED = "revoked"
OPEN_STATUSES = (STATUS_REQUESTED, STATUS_REMINDED)

VERDICT_MATCH = "match"
VERDICT_MISMATCH = "mismatch"
VERDICT_UNKNOWN = "unknown"

EMAIL_KIND_REQUEST = "file_request"
EMAIL_KIND_REMINDER = "file_request_reminder"
EMAIL_KIND_DIGEST = "digest"

#: The nil uuid the migration folds a null firm onto (firm_digest_prefs).
NIL_UUID = "00000000-0000-0000-0000-000000000000"


def _now() -> datetime:
    return datetime.now(timezone.utc)


def _iso(dt: datetime) -> str:
    return dt.isoformat()


def _app_url() -> str:
    return (os.environ.get("APP_URL") or "https://cfo-ai.io").rstrip("/")


# ══════════════════════════════════════════════════════════════════════════
# 1. TOKENS — signed, expiring, single-use
# ══════════════════════════════════════════════════════════════════════════


class SigningKeyMissing(RuntimeError):
    """FIRM_REQUEST_SIGNING_KEY is not configured. Fail closed."""


class TokenError(ValueError):
    """A token that cannot be honoured. `kind` names why:
    malformed | bad_signature | expired | version."""

    def __init__(self, kind: str, message: str) -> None:
        super().__init__(message)
        self.kind = kind


@dataclass(frozen=True)
class TokenClaims:
    request_id: str
    client_org_id: str
    period_end: str            # ISO date
    expires_at: str            # ISO timestamp (UTC)
    nonce: str
    version: str = TOKEN_VERSION

    def to_payload(self) -> Dict[str, Any]:
        return {"v": self.version, "r": self.request_id, "c": self.client_org_id,
                "p": self.period_end, "e": self.expires_at, "n": self.nonce}

    def expired(self, now: Optional[datetime] = None) -> bool:
        return _parse_ts(self.expires_at) <= (now or _now())


def _parse_ts(value: Any) -> datetime:
    text = str(value or "").replace("Z", "+00:00")
    dt = datetime.fromisoformat(text)
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt


def signing_key(env: Optional[Dict[str, str]] = None) -> bytes:
    raw = (env if env is not None else os.environ).get(SIGNING_KEY_ENV) or ""
    if len(raw.strip()) < 16:
        raise SigningKeyMissing(
            "%s is not configured (needs 16+ characters); refusing to mint or "
            "verify a file-request link against a key nobody set." % SIGNING_KEY_ENV)
    return raw.strip().encode("utf-8")


def _b64(data: bytes) -> str:
    return base64.urlsafe_b64encode(data).decode("ascii").rstrip("=")


def _unb64(text: str) -> bytes:
    return base64.urlsafe_b64decode(text + "=" * (-len(text) % 4))


def mint_token(claims: TokenClaims, key: bytes) -> str:
    body = json.dumps(claims.to_payload(), sort_keys=True, separators=(",", ":")).encode("utf-8")
    sig = hmac.new(key, body, hashlib.sha256).digest()
    return "%s.%s" % (_b64(body), _b64(sig))


def parse_token(token: str, key: bytes, now: Optional[datetime] = None) -> TokenClaims:
    """Verify the signature FIRST, then the version, then the expiry —
    an unverified token tells us nothing, not even its expiry."""
    parts = (token or "").strip().split(".")
    if len(parts) != 2 or not parts[0] or not parts[1]:
        raise TokenError("malformed", "the link is not a file-request link")
    try:
        body = _unb64(parts[0])
        sig = _unb64(parts[1])
    except Exception:  # noqa: BLE001 — any decode failure is malformed
        raise TokenError("malformed", "the link is not a file-request link")
    expected = hmac.new(key, body, hashlib.sha256).digest()
    if not hmac.compare_digest(sig, expected):
        raise TokenError("bad_signature", "the link's signature does not verify")
    try:
        payload = json.loads(body.decode("utf-8"))
    except Exception:  # noqa: BLE001
        raise TokenError("malformed", "the link's claims are unreadable")
    if not isinstance(payload, dict) or payload.get("v") != TOKEN_VERSION:
        raise TokenError("version", "the link was minted by a different version")
    claims = TokenClaims(
        request_id=str(payload.get("r") or ""), client_org_id=str(payload.get("c") or ""),
        period_end=str(payload.get("p") or ""), expires_at=str(payload.get("e") or ""),
        nonce=str(payload.get("n") or ""), version=TOKEN_VERSION)
    if not (claims.request_id and claims.client_org_id and claims.period_end
            and claims.expires_at):
        raise TokenError("malformed", "the link's claims are incomplete")
    if claims.expired(now):
        raise TokenError("expired", "the link has expired")
    return claims


def token_hash(token: str) -> str:
    return hashlib.sha256((token or "").encode("utf-8")).hexdigest()


# ══════════════════════════════════════════════════════════════════════════
# 2. DOCUMENT IDENTITY and the ENTITY GUARD
# ══════════════════════════════════════════════════════════════════════════


@dataclass(frozen=True)
class Preamble:
    """The rows ABOVE the column header the trial-balance parser located
    — where a SAGA / WinMentor export prints the company, its CUI and
    'Balanta de verificare la data de …'. `header_row_index` is the
    parser's own; -1 when the parser could not place a header (the first
    rows are then read as a best-effort preamble)."""

    text: str
    lines: Tuple[str, ...]
    sheet: Optional[str]
    header_row_index: int
    source_format: Optional[str]
    parse_ok: bool
    parse_error: Optional[str]


def _rows_xlsx(content: bytes, sheet: Optional[str], limit: int) -> List[List[str]]:
    import openpyxl  # local: heavy, and only reached for spreadsheets

    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.worksheets[0]
    rows = []  # type: List[List[str]]
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=limit, values_only=True)):
        rows.append([str(c) for c in row if c not in (None, "")])
    return rows


def _rows_csv(content: bytes, limit: int) -> List[List[str]]:
    text = content.decode("utf-8", errors="replace")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        delim = dialect.delimiter
    except csv.Error:
        delim = ";" if sample.count(";") >= sample.count(",") else ","
    rows = []  # type: List[List[str]]
    for i, row in enumerate(csv.reader(io.StringIO(text), delimiter=delim)):
        if i >= limit:
            break
        rows.append([c for c in row if c not in (None, "")])
    return rows


def _is_xlsx(content: bytes, filename: str) -> bool:
    return content[:4] == b"PK\x03\x04" or filename.lower().endswith((".xlsx", ".xlsm"))


def default_jurisdiction() -> str:
    """The jurisdiction a request-link upload is parsed under when the
    client workspace declares none — configuration, read once here."""
    return (os.environ.get("ENGINE_DEFAULT_JURISDICTION") or "RO").strip().upper()


def read_preamble(content: bytes, filename: str, max_rows: int = 12,
                  jurisdiction: Optional[str] = None) -> Preamble:
    """Locate the header row with the REAL parser of the client's
    jurisdiction pack (the same one the pipeline runs), then read the rows
    above it on the same sheet. A jurisdiction with no registered pack
    yields a best-effort preamble and a recorded parse gap."""
    sheet = None  # type: Optional[str]
    header_idx = -1
    source_format = None  # type: Optional[str]
    parse_ok = False
    parse_error = None  # type: Optional[str]
    try:
        from engine.core.country_pack_registry import get_pack
        import engine.country_packs.ro_romania  # noqa: F401 — registers the packs

        pack = get_pack(jurisdiction or default_jurisdiction())
        if content[:5] == b"%PDF-":
            rows_result = pack.parse_trial_balance(content, filename)
        elif _is_xlsx(content, filename):
            rows_result = pack.parse_trial_balance(content, filename)
        else:
            rows_result = pack.parse_trial_balance_csv(content, filename)
        extraction = getattr(rows_result, "extraction", None) or {}
        sheet = extraction.get("sheet")
        header_idx = int(extraction.get("header_row_index", -1))
        source_format = extraction.get("source_format")
        parse_ok = len(rows_result) > 0
    except Exception as exc:  # noqa: BLE001 — a parse failure is recorded, not raised
        parse_error = "%s: %s" % (type(exc).__name__, str(exc)[:200])
    lines = []  # type: List[str]
    try:
        if content[:5] == b"%PDF-":
            rows = []  # type: List[List[str]]
        elif _is_xlsx(content, filename):
            rows = _rows_xlsx(content, sheet, max_rows)
        else:
            rows = _rows_csv(content, max_rows)
        limit = header_idx if header_idx >= 0 else min(len(rows), 6)
        for row in rows[:limit]:
            line = " ".join(row).strip()
            if line:
                lines.append(line)
    except Exception as exc:  # noqa: BLE001
        parse_error = parse_error or ("preamble: %s" % type(exc).__name__)
    return Preamble(text="\n".join(lines), lines=tuple(lines), sheet=sheet,
                    header_row_index=header_idx, source_format=source_format,
                    parse_ok=parse_ok, parse_error=parse_error)


_CUI_RX = re.compile(
    r"(?:\bC\.?U\.?I\.?|\bC\.?I\.?F\.?|\bcod(?:ul)?\s+fiscal|\bc\.?f\.?)\s*[:.\-]?\s*(RO)?\s*(\d{2,10})\b",
    re.IGNORECASE)
_LEGAL_FORM_RX = re.compile(
    r"\b(S\.?\s?C\.?|S\.?R\.?L\.?|S\.?A\.?|S\.?N\.?C\.?|S\.?C\.?S\.?|P\.?F\.?A\.?|I\.?I\.?|I\.?F\.?|SRL-D|"
    r"LTD|LLC|GMBH|KFT|ZRT|BT|SA)\b\.?", re.IGNORECASE)
_LEGAL_FORM_HINT_RX = re.compile(
    r"\b(S\.?\s?C\.?|S\.?R\.?L\.?|S\.?A\.?|S\.?N\.?C\.?|S\.?C\.?S\.?|P\.?F\.?A\.?|SRL-D|LTD|LLC|GMBH|KFT|ZRT)\b",
    re.IGNORECASE)
_BALANCE_LINE_RX = re.compile(r"balan[tț][aă]|trial\s+balance", re.IGNORECASE)
_NOISE_WORDS = frozenset(["si", "and", "de", "the"])


def normalize_cui(raw: Any) -> Optional[str]:
    """Digits only; the RO prefix and punctuation are not identity."""
    if raw is None:
        return None
    digits = re.sub(r"\D", "", str(raw))
    return digits if 2 <= len(digits) <= 10 else None


def normalize_name(raw: Any) -> Tuple[str, ...]:
    """Lower-case ASCII tokens with legal forms, punctuation and noise
    words removed — what is left is the name itself."""
    if raw is None:
        return ()
    text = unicodedata.normalize("NFKD", str(raw))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = _LEGAL_FORM_RX.sub(" ", text)
    text = re.sub(r"[^A-Za-z0-9]+", " ", text).lower()
    return tuple(t for t in text.split() if t and t not in _NOISE_WORDS)


@dataclass(frozen=True)
class DocumentIdentity:
    name: Optional[str]
    cui: Optional[str]
    evidence: Tuple[str, ...]

    def to_payload(self) -> Dict[str, Any]:
        return {"name": self.name, "cui": self.cui, "evidence": list(self.evidence)}


def extract_identity(preamble_text: str) -> DocumentIdentity:
    """Company name + CUI from preamble text. The name is the first line
    that carries a legal form, else the first non-empty line before the
    'balanta' line; the CUI is the first fiscal-code pattern. Either may
    be absent — absence is reported, never filled."""
    lines = [ln.strip() for ln in (preamble_text or "").splitlines() if ln.strip()]
    cui = None  # type: Optional[str]
    evidence = []  # type: List[str]
    for ln in lines:
        m = _CUI_RX.search(ln)
        if m:
            cui = normalize_cui(m.group(2))
            evidence.append("cui: %s" % ln[:120])
            break
    name = None  # type: Optional[str]
    for ln in lines:
        if _BALANCE_LINE_RX.search(ln):
            break
        if _LEGAL_FORM_HINT_RX.search(ln) and not _CUI_RX.search(ln):
            name = ln
            evidence.append("name: %s" % ln[:120])
            break
    if name is None:
        for ln in lines:
            if _BALANCE_LINE_RX.search(ln):
                break
            if _CUI_RX.search(ln):
                continue
            if normalize_name(ln):
                name = ln
                evidence.append("name (first line): %s" % ln[:120])
                break
    return DocumentIdentity(name=name, cui=cui, evidence=tuple(evidence))


@dataclass(frozen=True)
class EntityVerdict:
    verdict: str
    reason: str
    document: Dict[str, Any]
    expected: Dict[str, Any]

    def to_payload(self) -> Dict[str, Any]:
        return {"verdict": self.verdict, "reason": self.reason,
                "document": dict(self.document), "expected": dict(self.expected)}


def entity_guard(document: DocumentIdentity, expected: Dict[str, Any]) -> EntityVerdict:
    """Compare the document's identity to the request's. CUI decides when
    both sides carry one; otherwise names decide when both carry one;
    otherwise the verdict is UNKNOWN — recorded, never treated as a
    match and never as a mismatch."""
    exp_cui = normalize_cui(expected.get("cui"))
    exp_name = normalize_name(expected.get("name"))
    doc_cui = normalize_cui(document.cui)
    doc_name = normalize_name(document.name)
    exp_payload = {"name": expected.get("name"), "cui": exp_cui}
    if exp_cui and doc_cui:
        if exp_cui == doc_cui:
            return EntityVerdict(VERDICT_MATCH, "the document's CUI equals the client's",
                                 document.to_payload(), exp_payload)
        return EntityVerdict(VERDICT_MISMATCH,
                             "the document's CUI %s is not the client's %s" % (doc_cui, exp_cui),
                             document.to_payload(), exp_payload)
    if exp_name and doc_name:
        a, b = set(exp_name), set(doc_name)
        overlap = len(a & b) / float(len(a | b)) if (a | b) else 0.0
        contained = a <= b or b <= a
        if contained or overlap >= 0.5:
            return EntityVerdict(VERDICT_MATCH,
                                 "the document names the client (%s)" % " ".join(doc_name),
                                 document.to_payload(), exp_payload)
        return EntityVerdict(VERDICT_MISMATCH,
                             "the document names %r, the request was minted for %r"
                             % (" ".join(doc_name), " ".join(exp_name)),
                             document.to_payload(), exp_payload)
    missing = []
    if not (doc_cui or doc_name):
        missing.append("the document carries no company name or CUI in its preamble")
    if not (exp_cui or exp_name):
        missing.append("the request carries no expected identity")
    return EntityVerdict(VERDICT_UNKNOWN, "; ".join(missing) or "identity could not be compared",
                         document.to_payload(), exp_payload)


def period_precheck(preamble_text: str, filename: str, bound_period_end: str) -> Dict[str, Any]:
    """The document's OWN period evidence (preamble text as the
    in-document channel, plus the filename) against the period the
    request is bound to. Advisory at landing; the AUTHORITATIVE record is
    the pipeline's `resolve_period_end_for_persist` at persist time."""
    detected = _period_detect.detect_period(
        extracted={"header_text": preamble_text} if preamble_text else None,
        filename=filename)
    proposed = detected.get("proposed_period_end")
    agrees = None  # type: Optional[bool]
    if proposed:
        agrees = proposed == bound_period_end
    return {
        "bound": bound_period_end,
        "proposed": proposed,
        "signal": detected.get("signal_used"),
        "confidence": detected.get("confidence"),
        "evidence": detected.get("evidence_snippet"),
        "agrees": agrees,
    }


def detected_type_for(filename: str, parse_ok: bool, mime: str = "") -> str:
    """Mirror of lib/supabase.ts::detectFromName, with one addition: a
    file the real parser accepted IS a trial balance."""
    if parse_ok:
        return "trial_balance"
    n = (filename or "").lower()
    m = (mime or "").lower()
    if re.search(r"balanta\s*verificare|trial[\s_-]?balance|^tb_", n):
        return "trial_balance"
    if re.search(r"bilan[tț]|balance[\s_-]?sheet|^bs_", n):
        return "bilant"
    if re.search(r"factur|invoice|saft|d406|smartbill", n):
        return "invoice"
    if re.search(r"p[\s_-]?l|profit[\s_-]?loss|cont[\s_-]?profit|^pl_", n):
        return "pl"
    if re.search(r"anual|annual[\s_-]?report", n):
        return "annual_report"
    if "spreadsheet" in m or re.search(r"\.xlsx?$", n):
        return "xlsx_workbook"
    if m == "text/csv" or n.endswith(".csv"):
        return "csv"
    if m.startswith("image/"):
        return "image"
    return "unknown"


@dataclass(frozen=True)
class UploadInspection:
    preamble: Preamble
    identity: DocumentIdentity
    entity: EntityVerdict
    period: Dict[str, Any]
    detected_type: str

    def to_payload(self) -> Dict[str, Any]:
        return {
            "identity": self.identity.to_payload(),
            "entity_guard": self.entity.to_payload(),
            "period_guard": dict(self.period),
            "detected_type": self.detected_type,
            "parse_ok": self.preamble.parse_ok,
            "parse_error": self.preamble.parse_error,
            "source_format": self.preamble.source_format,
        }


def inspect_upload(content: bytes, filename: str, expected_identity: Dict[str, Any],
                   bound_period_end: str, mime: str = "",
                   jurisdiction: Optional[str] = None) -> UploadInspection:
    preamble = read_preamble(content, filename, jurisdiction=jurisdiction)
    identity = extract_identity(preamble.text)
    return UploadInspection(
        preamble=preamble,
        identity=identity,
        entity=entity_guard(identity, expected_identity or {}),
        period=period_precheck(preamble.text, filename, bound_period_end),
        detected_type=detected_type_for(filename, preamble.parse_ok, mime),
    )


# ══════════════════════════════════════════════════════════════════════════
# 3. THE LANDING — through the normal pipeline
# ══════════════════════════════════════════════════════════════════════════


class LandingRefused(Exception):
    """The landing was refused. `status` is the HTTP status the route
    answers with; `detail` the payload."""

    def __init__(self, status: int, detail: Dict[str, Any]) -> None:
        super().__init__(str(detail.get("message") or detail.get("code") or "refused"))
        self.status = status
        self.detail = detail


@dataclass
class LandingDeps:
    """Every side effect of a landing, injectable. Defaults are the
    PRODUCTION functions — the same ones /api/pipeline/run uses — so the
    request link is not a second pipeline."""

    upload_object: Callable[[str, str, bytes, str], None]
    insert_document: Callable[[Dict[str, Any]], Dict[str, Any]]
    set_status: Callable[[str, str, Optional[str]], None]
    enqueue: Callable[[str], None]
    reserve: Callable[[str], Any]
    record_usage: Callable[[str, str], None]
    now: Callable[[], datetime] = _now


def _prod_upload_object(bucket: str, path: str, content: bytes, content_type: str) -> None:
    with _supabase.admin() as ac:
        ac.upload_object(bucket, path, content, content_type=content_type)


def _prod_insert_document(row: Dict[str, Any]) -> Dict[str, Any]:
    """Insert with the same graceful degrades lib/supabase.ts applies:
    an unmigrated optional column must never cost the client its
    upload."""
    with _supabase.admin() as ac:
        body = dict(row)
        for optional in ("content_hash", "period_end_hint", "jurisdiction_hint"):
            try:
                inserted = ac.insert("documents", body, returning=True)
                return inserted[0] if inserted else body
            except RuntimeError as exc:
                if optional in body and optional in str(exc):
                    logger.info("[firm] documents.%s rejected — retrying without it", optional)
                    body.pop(optional, None)
                    continue
                raise
        inserted = ac.insert("documents", body, returning=True)
        return inserted[0] if inserted else body


def _prod_set_status(doc_id: str, status: str, pipeline_started_at: Optional[str]) -> None:
    from . import pipeline as _pipeline
    _pipeline._admin_set_status(doc_id, status, pipeline_started_at=pipeline_started_at)


def _prod_enqueue(doc_id: str) -> None:
    from . import pipeline as _pipeline
    _pipeline._enqueue(doc_id)


def _prod_reserve(user_id: str) -> Any:
    from . import _usage_gate as _ug
    return _ug.reserve_document(user_id)


def _prod_record_usage(user_id: str, kind: str) -> None:
    from . import _usage_limits
    _usage_limits.record_usage(user_id, kind)


def production_deps() -> LandingDeps:
    return LandingDeps(
        upload_object=_prod_upload_object, insert_document=_prod_insert_document,
        set_status=_prod_set_status, enqueue=_prod_enqueue, reserve=_prod_reserve,
        record_usage=_prod_record_usage, now=_now)


@dataclass(frozen=True)
class LandingResult:
    document_id: str
    document_row: Dict[str, Any]
    storage_path: str
    inspection: UploadInspection
    was_extra: bool

    def to_payload(self) -> Dict[str, Any]:
        return {
            "document_id": self.document_id,
            "storage_path": self.storage_path,
            "period_end_hint": self.document_row.get("period_end_hint"),
            "original_filename": self.document_row.get("original_filename"),
            "inspection": self.inspection.to_payload(),
            "was_extra": self.was_extra,
        }


def _ext_of(filename: str) -> str:
    name = (filename or "").rsplit("/", 1)[-1]
    if "." in name:
        ext = name.rsplit(".", 1)[1].lower()
        if 1 <= len(ext) <= 8 and ext.isalnum():
            return ext
    return "bin"


def _quota_or_refuse(user_id: str, deps: LandingDeps) -> bool:
    """The SAME reservation `/api/pipeline/run` makes, mapped to the same
    HTTP shapes, charged to the accountant who minted the request."""
    decision = deps.reserve(user_id)
    kind = getattr(decision, "kind", "disabled")
    if kind == "blocked":
        raise LandingRefused(429, {
            "code": "doc_quota_blocked", "plan_key": getattr(decision, "plan_key", None),
            "docs_used": getattr(decision, "used", None),
            "docs_included": getattr(decision, "cap", None),
            "message": getattr(decision, "message", "document quota reached"),
            "upgrade_url": "/pricing"})
    if kind == "extra_required":
        raise LandingRefused(402, {
            "code": "extra_doc_confirmation_required",
            "plan_key": getattr(decision, "plan_key", None),
            "docs_used": getattr(decision, "used", None),
            "docs_included": getattr(decision, "cap", None),
            "extra_doc_eur": getattr(decision, "extra_doc_eur", None),
            "message": getattr(decision, "message", "an extra document must be confirmed"),
            "confirm_url": "/api/plan/confirm-extra-doc"})
    return bool(getattr(decision, "was_extra", False))


def land_file(request_row: Dict[str, Any], content: bytes, filename: str, mime: str,
              deps: LandingDeps, document_id: Optional[str] = None) -> LandingResult:
    """Land one file for one open request THROUGH THE NORMAL PIPELINE.

    Order matters and is the order the browser path takes: inspect
    (entity + period guards) → reserve the quota → write the blob → insert
    the documents row (period_end_hint = the request's period) →
    status queued → enqueue. The pipeline's own `stage_persist` then
    resolves the period from the hint and records its detection
    verdict exactly as for any other document.
    """
    if not content:
        raise LandingRefused(400, {"code": "empty_file", "message": "the file is empty"})
    if len(content) > MAX_UPLOAD_BYTES:
        raise LandingRefused(413, {"code": "file_too_large",
                                   "message": "the file exceeds the upload limit"})
    client_org_id = str(request_row.get("client_org_id") or "")
    period_end = str(request_row.get("period_end") or "")[:10]
    requested_by = request_row.get("requested_by")
    expected = request_row.get("expected_identity") or {}
    inspection = inspect_upload(content, filename, expected, period_end, mime,
                                jurisdiction=(request_row.get("jurisdiction") or None))
    if inspection.entity.verdict == VERDICT_MISMATCH:
        raise LandingRefused(409, {
            "code": "entity_mismatch",
            "message": "This file does not appear to belong to the company the "
                       "request was issued for: %s." % inspection.entity.reason,
            "entity_guard": inspection.entity.to_payload(),
            "period_guard": inspection.period,
        })

    was_extra = _quota_or_refuse(str(requested_by), deps) if requested_by else False

    import uuid

    doc_id = document_id or str(uuid.uuid4())
    ext = _ext_of(filename)
    storage_path = "%s/uploads/%s.%s" % (client_org_id, doc_id, ext)
    content_type = mime or "application/octet-stream"
    deps.upload_object(DOC_BUCKET, storage_path, content, content_type)

    row = {
        "id": doc_id,
        "org_id": client_org_id,
        "uploaded_by": requested_by,
        "storage_path": storage_path,
        "original_filename": filename or ("upload.%s" % ext),
        "mime_type": content_type,
        "size_bytes": len(content),
        "detected_type": inspection.detected_type,
        "status": "queued",
        "scope": "financial",
        "content_hash": hashlib.sha256(content).hexdigest(),
        # The CONFIRMATION channel (W-law): the accountant chose this
        # month when minting the request. A document that disagrees is
        # recorded as a mismatch by stage_persist, never silently refiled.
        "period_end_hint": period_end,
    }
    if was_extra:
        row["metered_extra"] = True
    inserted = deps.insert_document(row)
    started = _iso(deps.now())
    deps.set_status(doc_id, "queued", started)
    deps.enqueue(doc_id)
    if requested_by:
        try:
            deps.record_usage(str(requested_by), "upload")
        except Exception:  # noqa: BLE001 — the legacy counter is soft
            logger.exception("[firm] record_usage failed (soft counter)")
    return LandingResult(document_id=doc_id, document_row=dict(inserted or row),
                         storage_path=storage_path, inspection=inspection,
                         was_extra=was_extra)


# ══════════════════════════════════════════════════════════════════════════
# 4. EMAIL — queued stub + branded templates
# ══════════════════════════════════════════════════════════════════════════


def email_file_request(*, client_name: str, firm_name: str, period_end: str,
                       upload_url: str, expires_at: str, note: str = "",
                       reminder: bool = False, days_to_deadline: Optional[int] = None) -> str:
    """The branded request / reminder body (engine.api._email_templates
    shell). Every figure in it is a label the engine wrote."""
    T = _email_templates
    esc = _html.escape
    when = ("in %d day(s)" % days_to_deadline if isinstance(days_to_deadline, int)
            and days_to_deadline > 0 else "soon")
    eyebrow = "Reminder" if reminder else "Document request"
    title = ("%s still needs the trial balance for %s" % (firm_name, period_end)
             if reminder else
             "%s asks for the trial balance for %s" % (firm_name, period_end))
    detail = """
      <table role="presentation" width="100%%" cellpadding="0" cellspacing="0" border="0"
             style="margin:0 0 24px 0;background:%s;border:1px solid %s;border-radius:12px;">
        <tr>
          <td style="padding:16px 20px;font-family:%s;font-size:12px;letter-spacing:0.06em;text-transform:uppercase;color:%s;">Company</td>
          <td align="right" style="padding:16px 20px;font-family:%s;font-size:18px;color:%s;">%s</td>
        </tr>
        <tr><td colspan="2" style="padding:0 20px;"><div style="height:1px;background:%s;font-size:0;line-height:0;">&nbsp;</div></td></tr>
        <tr>
          <td style="padding:16px 20px;font-family:%s;font-size:12px;letter-spacing:0.06em;text-transform:uppercase;color:%s;">Period</td>
          <td align="right" style="padding:16px 20px;font-family:%s;font-size:18px;color:%s;">%s</td>
        </tr>
        <tr><td colspan="2" style="padding:0 20px;"><div style="height:1px;background:%s;font-size:0;line-height:0;">&nbsp;</div></td></tr>
        <tr>
          <td style="padding:16px 20px;font-family:%s;font-size:12px;letter-spacing:0.06em;text-transform:uppercase;color:%s;">Link valid until</td>
          <td align="right" style="padding:16px 20px;font-family:%s;font-size:18px;color:%s;">%s</td>
        </tr>
      </table>""" % (T.PAGE, T.BORDER, T.SANS, T.MUTED, T.SERIF, T.INK, esc(client_name),
                     T.BORDER, T.SANS, T.MUTED, T.SERIF, T.INK, esc(period_end),
                     T.BORDER, T.SANS, T.MUTED, T.SERIF, T.INK, esc(expires_at[:10]))
    lead = ("Your accountant is waiting for this file to close the period %s." % when
            if reminder else
            "Upload the trial balance (balanță de verificare) for this period using "
            "the secure link below. The link is personal, single-use, and expires.")
    body = "\n".join([
        T._eyebrow(eyebrow), T._h1(title), T._p(esc(lead), gap=22), detail,
        (T._p("<em>%s</em>" % esc(note), muted=True, small=True) if note else ""),
        T._button("Upload the file", upload_url), T._fallback_link(upload_url),
        T._p("The file lands directly in %s's workspace at CFO AI; the period is "
             "already selected. If the file is for a different month, say so to "
             "your accountant instead of uploading it here." % esc(client_name),
             muted=True, small=True, gap=0),
    ])
    return T._layout(title=title, body_html=body,
                     footer_html="Sent on behalf of %s." % esc(firm_name),
                     preheader="%s — trial balance for %s" % (client_name, period_end))


def render_queued_email(row: Dict[str, Any]) -> Tuple[str, str]:
    """(subject, html) for a queued row. The digest carries its own
    pre-rendered body (deterministic from its payload); request kinds are
    rendered here from `vars`."""
    payload = row.get("payload") or {}
    kind = row.get("kind")
    subject = str(payload.get("subject") or "CFO AI")
    if kind == EMAIL_KIND_DIGEST:
        return subject, str(payload.get("html") or "")
    v = payload.get("vars") or {}
    html = email_file_request(
        client_name=str(v.get("client_name") or ""), firm_name=str(v.get("firm_name") or "CFO AI"),
        period_end=str(v.get("period_end") or ""), upload_url=str(v.get("upload_url") or ""),
        expires_at=str(v.get("expires_at") or ""), note=str(v.get("note") or ""),
        reminder=(kind == EMAIL_KIND_REMINDER),
        days_to_deadline=(int(v["days_to_deadline"]) if isinstance(v.get("days_to_deadline"), int)
                          else None))
    return subject, html


def queue_email(client: Any, *, kind: str, to_email: str, template: str,
                payload: Dict[str, Any], firm_id: Optional[str] = None,
                client_org_id: Optional[str] = None, request_id: Optional[str] = None,
                user_id: Optional[str] = None, send_at: Optional[str] = None) -> Optional[str]:
    """One queued row. Never raises: a missing table logs the payload so
    an operator can wire delivery later (the billing precedent)."""
    row = {
        "kind": kind, "to_email": to_email, "template": template, "payload": payload,
        "firm_id": firm_id, "client_org_id": client_org_id, "request_id": request_id,
        "user_id": user_id, "send_at": send_at or _iso(_now()), "status": "queued",
    }
    try:
        inserted = client.insert("firm_email_queue", row, returning=True)
        return str(inserted[0].get("id")) if inserted else None
    except Exception:  # noqa: BLE001
        logger.info("[firm] email pending (queue unavailable): %s", json.dumps(
            {"kind": kind, "to": to_email, "subject": payload.get("subject")}))
        return None


# ══════════════════════════════════════════════════════════════════════════
# 5. AUTH — firm context when a firm is named, client membership otherwise
# ══════════════════════════════════════════════════════════════════════════


def _require_jwt(authorization: Optional[str]) -> str:
    from . import _test_mode
    if _test_mode.is_test_mode():
        try:
            return _test_mode.get_test_user_jwt()
        except Exception:  # noqa: BLE001
            return _test_mode.JWT_BYPASS_PLACEHOLDER
    if not authorization or not authorization.lower().startswith("bearer "):
        raise HTTPException(401, "Missing Bearer token.")
    return authorization.split(" ", 1)[1].strip()


def _user_id(jwt: str) -> str:
    from . import _test_mode
    if _test_mode.is_bypass_token(jwt):
        return _test_mode.test_user_id()
    return _org.resolve_user_id(jwt)


def authorize_client(jwt: str, firm_id: Optional[str], client_org_id: str,
                     action: str) -> Tuple[str, Optional[str]]:
    """(user_id, firm_id). With a firm: the tenancy lane's two walls
    (member of the firm, role holds `action`, client belongs to the
    firm). Without one: a direct membership in the client workspace —
    the pre-firm tenancy primitive, validated the way every route
    validates it."""
    firm = (firm_id or "").strip()
    if firm:
        from . import _firm
        ctx = _firm.resolve_firm(jwt, firm)
        _firm.require(ctx, action)
        _firm.require_client(ctx, client_org_id)
        return ctx.user_id, firm
    user_id = _user_id(jwt)
    if not _org.user_is_member(user_id, client_org_id):
        raise HTTPException(403, "Not a member of the requested workspace.")
    return user_id, None


#: What PostgREST says when a per-user write meets a policy that refuses
#: it: HTTP 401/403 with SQLSTATE 42501 ("new row violates row-level
#: security policy"). `SupabaseClient.insert` folds the body into a
#: RuntimeError; `upsert` / `update` raise httpx.HTTPStatusError with the
#: response attached. Both shapes are read here.
_RLS_REFUSAL_MARKERS = ("42501", "row-level security", "HTTP 401", "HTTP 403")


def rls_refused(exc: BaseException) -> bool:
    """True when the database wall refused a per-user write — not when
    the write failed for any other reason (those must surface)."""
    response = getattr(exc, "response", None)
    status = getattr(response, "status_code", None)
    if status in (401, 403):
        return True
    text = str(exc)
    return any(marker in text for marker in _RLS_REFUSAL_MARKERS)


def write_as_caller(fn: Callable[[], Any], table: str) -> Any:
    """Run ONE per-user write. A refusal by row-level security is the
    SECOND WALL answering — reported as 403 naming the table, never
    swallowed (a write that silently did not happen is a lie to the
    caller) and never a 500 (a wall is not a crash)."""
    try:
        return fn()
    except HTTPException:
        raise
    except Exception as exc:  # noqa: BLE001 — classified below
        if rls_refused(exc):
            raise HTTPException(
                403, "Write to %s refused by row-level security: the caller's role "
                     "does not hold it on this client." % table)
        raise


def _live_membership_org_ids(ac: Any, user_id: str) -> List[str]:
    out = []  # type: List[str]
    rows = select_all(ac, "memberships", order="created_at.asc,org_id.asc",
                      filters={"user_id": "eq.%s" % user_id}, columns="org_id")
    for row in rows or []:
        org = ac.select("organizations", filters={"id": "eq.%s" % row["org_id"]},
                        columns="id,archived_at", limit=1)
        if org and org[0].get("archived_at") is None:
            out.append(str(row["org_id"]))
    return sorted(set(out))


def client_org_ids_for(user_id: str, firm_id: Optional[str], jwt: Optional[str] = None,
                       ac: Optional[Any] = None) -> List[str]:
    """Every client the caller may see: the firm's clients when a firm is
    named (membership validated), else every live workspace the user
    holds a membership in. `ac` is the service-role client to read with
    (the crons pass theirs); None opens one."""
    firm = (firm_id or "").strip()
    if firm:
        from . import _firm
        if jwt is not None:
            _firm.resolve_firm(jwt, firm)
        elif _firm.firm_role_for_user(user_id, firm) is None:
            raise HTTPException(403, "Not a member of the requested firm.")
        return sorted(_firm.client_org_ids(firm))
    if ac is not None:
        return _live_membership_org_ids(ac, user_id)
    with _supabase.admin() as opened:
        return _live_membership_org_ids(opened, user_id)


def _scheduler_token_or_503(authorization: Optional[str]) -> None:
    """ENGINE_API_TOKEN bearer, FAIL CLOSED (the purge-cron precedent):
    a cron that sends mail must not run anonymously."""
    token = os.environ.get("ENGINE_API_TOKEN")
    if not token:
        raise HTTPException(503, "ENGINE_API_TOKEN is not configured; refusing to run a cron.")
    if not authorization or not authorization.lower().startswith("bearer "):
        raise HTTPException(401, "Missing Bearer token.")
    if authorization.split(" ", 1)[1].strip() != token:
        raise HTTPException(401, "Invalid scheduler token.")


# ══════════════════════════════════════════════════════════════════════════
# 6. REQUEST ROWS — helpers the routes and the crons share
# ══════════════════════════════════════════════════════════════════════════


def _org_row(ac: Any, org_id: str) -> Dict[str, Any]:
    rows = ac.select("organizations", filters={"id": "eq.%s" % org_id}, limit=1)
    if not rows:
        raise HTTPException(404, "Workspace not found.")
    return rows[0]


def _cadence_row(ac: Any, client_org_id: str,
                 firm_id: Optional[str] = None) -> Optional[Dict[str, Any]]:
    """The client's stored cadence, or None (pack default). When the
    caller acts THROUGH A FIRM, a row another firm stored (its `firm_id`
    is neither null nor this firm's) is not theirs — the route pins as
    the `firm_client_cadence firm read` policy pins (W3), so either wall
    alone leaves firm B nothing of firm A's era."""
    try:
        rows = ac.select("firm_client_cadence",
                         filters={"client_org_id": "eq.%s" % client_org_id}, limit=1)
    except Exception:  # noqa: BLE001 — table not migrated yet: pack default
        return None
    row = rows[0] if rows else None
    if row is not None and firm_id and row.get("firm_id") not in (None, firm_id):
        return None
    return row


def _filed_period_ends(ac: Any, client_org_id: str) -> List[str]:
    rows = select_all(ac, "financial_periods", order="id.asc",
                      filters={"org_id": "eq.%s" % client_org_id},
                      columns="period_end,source_document_id")
    return sorted(set(str(r.get("period_end"))[:10] for r in rows or []
                      if r.get("period_end")))


def request_token(ac: Any, request_id: str) -> Optional[str]:
    """The original link token (service role only — see the migration's
    §2b). None when the secret row is missing: a reminder then records
    that it could not carry a link rather than inventing one."""
    try:
        rows = ac.select("firm_file_request_tokens",
                         filters={"request_id": "eq.%s" % request_id}, limit=1)
    except Exception:  # noqa: BLE001
        return None
    return str(rows[0].get("token")) if rows and rows[0].get("token") else None


def request_status_label(row: Dict[str, Any], as_of: date) -> str:
    """'requested 2d ago' / 'reminded 1d ago' / 'received' / 'expired' /
    'revoked' — engine text for the list."""
    status = str(row.get("status") or "")
    if status in (STATUS_RECEIVED, STATUS_EXPIRED, STATUS_REVOKED):
        return status
    anchor = row.get("last_reminded_at") if status == STATUS_REMINDED else row.get("requested_at")
    try:
        days = (as_of - date.fromisoformat(str(anchor)[:10])).days
    except ValueError:
        return status
    if days <= 0:
        return "%s today" % status
    return "%s %dd ago" % (status, days)


def public_request_view(row: Dict[str, Any], as_of: date) -> Dict[str, Any]:
    return {
        "id": row.get("id"), "firm_id": row.get("firm_id"),
        "client_org_id": row.get("client_org_id"),
        "period_end": str(row.get("period_end") or "")[:10],
        "status": row.get("status"), "status_label": request_status_label(row, as_of),
        "requested_at": row.get("requested_at"), "expires_at": row.get("expires_at"),
        "reminder_count": row.get("reminder_count"), "last_reminded_at": row.get("last_reminded_at"),
        "refusal_count": row.get("refusal_count"), "consumed_at": row.get("consumed_at"),
        "document_id": row.get("document_id"), "to_email": row.get("to_email"),
        "note": row.get("note"), "entity_guard": row.get("entity_guard"),
        "period_guard": row.get("period_guard"),
        "route": DG.ROUTE_REQUESTS.format(client_org_id=row.get("client_org_id")),
    }


def open_request_items(rows: Sequence[Dict[str, Any]], client_names: Dict[str, str],
                       as_of: date) -> List[DG.AttentionItem]:
    """FILE_REQUESTED items for every open request — what the digest and
    the brief add to the board's own items."""
    out = []  # type: List[DG.AttentionItem]
    for row in rows:
        if str(row.get("status")) not in OPEN_STATUSES:
            continue
        name = client_names.get(str(row.get("client_org_id")), "")
        if not name:
            continue
        out.append(DG.item_for_open_request(row, name, as_of))
    return out


# ══════════════════════════════════════════════════════════════════════════
# 7. THE NUDGE CRON — reminders N days before the deadline, expiry
# ══════════════════════════════════════════════════════════════════════════


def run_nudge_cron(as_of: date, now: datetime, client: Any,
                   firm_names: Optional[Dict[str, str]] = None) -> Dict[str, Any]:
    """For every open request: expire it when its link has expired;
    otherwise queue a reminder for each cadence nudge day that has come
    and was not sent. Idempotent: `reminders_sent` on the row records
    (period_end, days_before) pairs."""
    pack = CAD.load_cadence_pack()
    rows = select_all(client, "firm_file_requests", order="id.asc",
                      filters={"status": "in.(%s)" % ",".join(OPEN_STATUSES)})
    queued = expired = 0
    for row in rows or []:
        rid = str(row.get("id"))
        try:
            if _parse_ts(row.get("expires_at")) <= now:
                client.update("firm_file_requests", {"status": STATUS_EXPIRED},
                              filters={"id": "eq.%s" % rid})
                expired += 1
                continue
        except (TypeError, ValueError):
            pass
        client_org_id = str(row.get("client_org_id") or "")
        cadence = CAD.resolve_client_cadence(client_org_id, _cadence_row(client, client_org_id), pack)
        period_end = str(row.get("period_end") or "")[:10]
        sent = [(str(k[0]), int(k[1])) for k in (row.get("reminders_sent") or [])
                if isinstance(k, (list, tuple)) and len(k) == 2]
        # Only nudges for THIS request's period; the request is the ask.
        due = [n for n in CAD.nudges_due(cadence, [], as_of, sent)
               if n.period_end.isoformat() == period_end]
        if not due:
            continue
        org = _org_row(client, client_org_id)
        to_email = row.get("to_email")
        token = request_token(client, rid)
        for nudge in due:
            if to_email and token:
                queue_email(client, kind=EMAIL_KIND_REMINDER, to_email=str(to_email),
                            template=EMAIL_KIND_REMINDER, payload={
                                "subject": "Reminder: trial balance for %s" % period_end,
                                "vars": {
                                    "client_name": org.get("name"),
                                    "firm_name": (firm_names or {}).get(str(row.get("firm_id")), "CFO AI"),
                                    "period_end": period_end,
                                    "upload_url": "%s%s" % (_app_url(), ROUTE_UPLOAD_REQUEST.format(
                                        token=token)),
                                    "expires_at": str(row.get("expires_at") or ""),
                                    "days_to_deadline": (nudge.deadline - as_of).days,
                                    "note": row.get("note") or "",
                                }},
                            firm_id=row.get("firm_id"), client_org_id=client_org_id,
                            request_id=rid)
                queued += 1
            sent.append((period_end, nudge.days_before))
        client.update("firm_file_requests", {
            "status": STATUS_REMINDED, "reminder_count": int(row.get("reminder_count") or 0) + len(due),
            "last_reminded_at": _iso(now), "reminders_sent": [list(k) for k in sent],
        }, filters={"id": "eq.%s" % rid})
    return {"as_of": as_of.isoformat(), "open": len(rows or []), "reminders_queued": queued,
            "expired": expired}


# ══════════════════════════════════════════════════════════════════════════
# 8. THE DIGEST CRON
# ══════════════════════════════════════════════════════════════════════════

#: (client_org_ids, as_of) -> an object with `.items()` and `.rows` (a
#: FirmAttentionReport) — the board computation, injected so the cron and
#: its tests never construct one here.
ReportProvider = Callable[[Sequence[str], date], Any]


class AttentionUnavailable(RuntimeError):
    """No report provider could compute the board (the board lane's
    loader is not wired on this host). The cron records the gap and
    queues nothing — an empty digest is not a digest."""


def default_report_provider(client_factory: Optional[Callable[[], Any]] = None,
                            firm_id: Optional[str] = None) -> ReportProvider:
    """The board's own runner over the persisted periods of the given
    clients. Statements (for CRITICAL_FINDING) are not loaded here — the
    board records that as a gap per period; wiring the statements
    provider is the board lane's route work.

    `client_factory` opens the Supabase client the periods are read
    with. A ROUTE passes the caller's own (`_supabase.per_user(jwt)`) so
    RLS is the second wall under the Python one; only the digest CRON —
    where no user JWT exists — takes the default service-role client.
    `firm_id` pins the cadence read to the firm the caller acts through
    (see `_cadence_row`)."""
    from engine.firm import attention as A, model as M

    opener = client_factory or _supabase.admin

    def _provider(client_org_ids: Sequence[str], as_of: date) -> Any:
        clients = []  # type: List[M.ClientRecord]
        with opener() as ac:
            for org_id in client_org_ids:
                org = ac.select("organizations", filters={"id": "eq.%s" % org_id}, limit=1)
                if not org:
                    continue
                periods = select_all(
                    ac, "financial_periods", order="period_end.desc,id.asc",
                    filters={"org_id": "eq.%s" % org_id},
                    columns="id,period_end,period_start,currency,source_document_id,"
                            "updated_at,assembled_canonical_v1")
                records = []  # type: List[M.PeriodRecord]
                for p in periods or []:
                    env = p.get("assembled_canonical_v1")
                    records.append(M.PeriodRecord(
                        period_id=str(p.get("id")), period_end=str(p.get("period_end"))[:10],
                        currency=str(p.get("currency") or "RON"),
                        period_start=(str(p.get("period_start"))[:10] if p.get("period_start") else None),
                        source_document_id=p.get("source_document_id"),
                        updated_at=str(p.get("updated_at") or ""),
                        envelope=(env if isinstance(env, dict) and env else None)))
                clients.append(M.ClientRecord(
                    client_id=str(org_id), client_name=str(org[0].get("name") or org_id),
                    jurisdiction=str(org[0].get("country") or os.environ.get(
                        "ENGINE_DEFAULT_JURISDICTION") or "RO"),
                    periods=tuple(records), cadence_row=_cadence_row(ac, org_id, firm_id)))
        return A.compute_firm_attention(clients, as_of=as_of)

    return _provider


def _pref_due(pref: Dict[str, Any], as_of: date, now: datetime) -> bool:
    if not pref.get("enabled"):
        return False
    hour = int(pref.get("send_hour_utc") if pref.get("send_hour_utc") is not None else 7)
    if now.hour < hour:
        return False
    last = pref.get("last_sent_at")
    if not last:
        return True
    try:
        last_day = _parse_ts(last).date()
    except (TypeError, ValueError):
        return True
    if str(pref.get("frequency") or "daily") == "weekly":
        return (as_of - last_day).days >= 7
    return last_day < as_of


def run_digest_cron(as_of: date, now: datetime, client: Any,
                    report_provider: Optional[ReportProvider] = None,
                    email_of: Optional[Callable[[str], Optional[str]]] = None,
                    app_url: Optional[str] = None) -> Dict[str, Any]:
    """One pass over every ENABLED preference. Nothing is queued for a
    user who did not opt in, for a user whose last digest is too recent,
    or when nothing is new since that digest."""
    prefs = select_all(client, "firm_digest_prefs", order="id.asc",
                       filters={"enabled": "eq.true"})
    provider = report_provider
    base = (app_url if app_url is not None else _app_url())
    queued = skipped = empty = gaps = 0
    notes = []  # type: List[str]
    for pref in prefs or []:
        user_id = str(pref.get("user_id") or "")
        firm_id = pref.get("firm_id")
        firm_key = str(firm_id or user_id)
        if not _pref_due(pref, as_of, now):
            skipped += 1
            continue
        try:
            ids = client_org_ids_for(user_id, firm_id, ac=client)
        except HTTPException as exc:
            notes.append("user %s: %s" % (user_id, exc.detail))
            gaps += 1
            continue
        if provider is None:
            try:
                provider = default_report_provider()
            except Exception as exc:  # noqa: BLE001
                raise AttentionUnavailable("the board runner could not be loaded: %s" % exc)
        report = provider(ids, as_of)
        names = dict((str(r.client_id), str(r.client_name)) for r in getattr(report, "rows", ()))
        requests = []  # type: List[Dict[str, Any]]
        for chunk in chunked(ids):
            requests.extend(select_all(client, "firm_file_requests", order="id.asc",
                                       filters={"client_org_id": "in.(%s)" % ",".join(chunk)}))
        raws = list(report.items()) + [it.to_payload() for it in
                                       open_request_items(requests or [], names, as_of)]
        digest = DG.build_digest(
            raws, as_of, seen_ids=list(pref.get("last_item_ids") or []),
            user_id=user_id, firm_key=firm_key, client_names=names,
            kind_order=DG.kind_order_from_report(report))
        if digest.is_empty:
            empty += 1
            continue
        to_email = (email_of or _default_email_of)(user_id)
        if not to_email:
            notes.append("user %s: no email address" % user_id)
            gaps += 1
            continue
        log_row = {"user_id": user_id, "firm_id": firm_id, "sent_for_date": as_of.isoformat(),
                   "item_set_hash": digest.item_set_hash, "item_count": digest.counts.get("shown", 0)}
        try:
            existing = client.select("firm_digest_log", filters={
                "user_id": "eq.%s" % user_id, "sent_for_date": "eq.%s" % as_of.isoformat(),
                "firm_id": ("is.null" if firm_id is None else "eq.%s" % firm_id)}, limit=1)
        except Exception:  # noqa: BLE001
            existing = []
        if existing:
            skipped += 1
            continue
        email_id = queue_email(client, kind=EMAIL_KIND_DIGEST, to_email=to_email,
                               template=EMAIL_KIND_DIGEST, payload={
                                   "subject": DG.digest_subject(digest),
                                   "html": DG.render_digest_email(digest, base),
                                   "text": DG.render_digest_text(digest, base),
                                   "digest": digest.to_payload(),
                               }, firm_id=firm_id, user_id=user_id)
        log_row["queued_email_id"] = email_id
        try:
            client.insert("firm_digest_log", log_row, returning=False)
        except Exception:  # noqa: BLE001
            logger.exception("[firm] digest log insert failed")
        client.update("firm_digest_prefs", {
            "last_sent_at": _iso(now), "last_item_set_hash": digest.item_set_hash,
            "last_item_ids": list(digest.item_ids),
        }, filters={"id": "eq.%s" % pref.get("id")})
        queued += 1
    return {"as_of": as_of.isoformat(), "prefs": len(prefs or []), "queued": queued,
            "skipped": skipped, "nothing_new": empty, "gaps": gaps, "notes": notes}


def _default_email_of(user_id: str) -> Optional[str]:
    from . import _billing
    return _billing._user_email(user_id)


def open_request_for(ac: Any, token: str, now: Optional[datetime] = None
                     ) -> Tuple[TokenClaims, Dict[str, Any]]:
    """Resolve a token to its OPEN request row, or refuse with the HTTP
    status the route answers with: 503 no key, 403 bad signature / a row
    that does not match its claims, 410 expired / used / revoked, 404
    unknown. Signature first, row second, state last."""
    try:
        key = signing_key()
    except SigningKeyMissing as exc:
        raise HTTPException(503, str(exc))
    try:
        claims = parse_token(token, key, now)
    except TokenError as exc:
        status = {"expired": 410}.get(exc.kind, 403)
        raise HTTPException(status, str(exc))
    rows = ac.select("firm_file_requests",
                     filters={"token_hash": "eq.%s" % token_hash(token)}, limit=1)
    if not rows:
        raise HTTPException(404, "This request link is not known.")
    row = rows[0]
    if str(row.get("id")) != claims.request_id or \
            str(row.get("client_org_id")) != claims.client_org_id or \
            str(row.get("period_end"))[:10] != claims.period_end:
        raise HTTPException(403, "This request link does not match its record.")
    status = str(row.get("status") or "")
    if status == STATUS_RECEIVED or row.get("consumed_at"):
        raise HTTPException(410, "This request link was already used.")
    if status == STATUS_REVOKED:
        raise HTTPException(410, "This request link was revoked.")
    if status == STATUS_EXPIRED:
        raise HTTPException(410, "This request link has expired.")
    # THE ERA PIN (D3): a signed link is not an era pin. The request was
    # minted by `row.firm_id`; when that firm no longer serves the client
    # (detached since, or attached elsewhere) the link is dead — 410, like
    # a revoked one — whatever its signature and expiry still say. A row a
    # workspace member minted with no firm (firm_id null) is the client's
    # own and outlives any firm. Read as `ac` (the landing's service-role
    # client): the pin is the CURRENT firm, which no policy would hide.
    minted_by = row.get("firm_id")
    if minted_by is not None:
        try:
            orgs = ac.select("organizations", filters={"id": "eq.%s" % claims.client_org_id},
                             columns="id,firm_id", limit=1) or []
        except Exception as exc:  # noqa: BLE001 — fail CLOSED: an unverifiable era is not an open link
            raise HTTPException(503, "This request link cannot be verified against the client's "
                                     "current firm (%s)." % (type(exc).__name__,))
        current = orgs[0].get("firm_id") if orgs else None
        if current is None or str(current) != str(minted_by):
            raise HTTPException(410, "This request link was issued by a firm that no longer "
                                     "serves this client.")
    return claims, row


# ══════════════════════════════════════════════════════════════════════════
# 9. ROUTES
# ══════════════════════════════════════════════════════════════════════════


class CreateRequestBody(BaseModel):
    client_org_id: str = Field(..., min_length=1)
    period_end: str = Field(..., min_length=10, max_length=10)
    firm_id: Optional[str] = None
    to_email: Optional[str] = None
    note: Optional[str] = Field(None, max_length=1000)
    expected_cui: Optional[str] = None
    firm_name: Optional[str] = None


class CadenceBody(BaseModel):
    client_org_id: str = Field(..., min_length=1)
    firm_id: Optional[str] = None
    cadence: str = Field(..., min_length=1)
    deadline_days_after_period_end: Optional[int] = Field(None, ge=1, le=120)
    fiscal_year_end_month: Optional[int] = Field(None, ge=1, le=12)
    nudge_days_before: Optional[List[int]] = None


class DigestPrefsBody(BaseModel):
    firm_id: Optional[str] = None
    enabled: bool = False
    frequency: str = "daily"
    send_hour_utc: int = Field(7, ge=0, le=23)


def _period_end_or_400(value: str) -> str:
    try:
        d = date.fromisoformat(value)
    except ValueError:
        raise HTTPException(400, "period_end must be an ISO date (YYYY-MM-DD).")
    if d != CAD.month_end(d.year, d.month):
        raise HTTPException(400, "period_end must be the last day of a month.")
    if d.year < 2000 or d.year > 2035:
        raise HTTPException(400, "period_end is outside the plausible range.")
    return d.isoformat()


def build_router() -> APIRouter:
    router = APIRouter(prefix="/api/firm", tags=["firm-requests"])

    # ── requests ────────────────────────────────────────────────────────

    @router.post("/requests", status_code=201)
    def create_request(body: CreateRequestBody,
                       authorization: Optional[str] = Header(None)) -> Dict[str, Any]:
        jwt = _require_jwt(authorization)
        user_id, firm_id = authorize_client(jwt, body.firm_id, body.client_org_id, "request_file")
        period_end = _period_end_or_400(body.period_end)
        try:
            key = signing_key()
        except SigningKeyMissing as exc:
            raise HTTPException(503, str(exc))
        pack = CAD.load_cadence_pack()
        now = _now()
        expires = now + timedelta(days=pack.request_link_ttl_days)
        import uuid

        # The client's identity is read as the CALLER (`organizations firm
        # read` / membership) and the request row is INSERTED AS THE CALLER
        # — `firm_file_requests firm write` (RLS) is the second wall on the
        # write (W1). Only once that insert was admitted does the backend
        # write the secret token and queue the e-mail, service-role.
        expected = None  # type: Optional[Dict[str, Any]]
        with _supabase.per_user(jwt) as client:
            org = _org_row(client, body.client_org_id)
            expected = {"name": org.get("name"),
                        "cui": normalize_cui(body.expected_cui or org.get("cui"))}
            request_id = str(uuid.uuid4())
            claims = TokenClaims(request_id=request_id, client_org_id=body.client_org_id,
                                 period_end=period_end, expires_at=_iso(expires),
                                 nonce=uuid.uuid4().hex)
            token = mint_token(claims, key)
            row = {
                "id": request_id, "firm_id": firm_id, "client_org_id": body.client_org_id,
                "period_end": period_end, "requested_by": user_id, "requested_at": _iso(now),
                "expires_at": _iso(expires), "token_hash": token_hash(token),
                "to_email": body.to_email, "note": body.note, "status": STATUS_REQUESTED,
                "expected_identity": expected, "reminders_sent": [],
            }

            def _insert_request() -> None:
                try:
                    client.insert("firm_file_requests", row, returning=False)
                except RuntimeError as exc:
                    if "firm_file_requests_open_uidx" in str(exc) or "23505" in str(exc):
                        raise HTTPException(409, "An open request for this client and period "
                                                 "already exists; remind or revoke it instead.")
                    raise

            write_as_caller(_insert_request, "firm_file_requests")
        with _supabase.admin() as ac:
            # The secret half, service-role only (schema_phase_firm_requests
            # §2b): the reminder cron needs the link, members only see the hash.
            ac.insert("firm_file_request_tokens", {"request_id": request_id, "token": token},
                      returning=False)
            route = ROUTE_UPLOAD_REQUEST.format(token=token)
            upload_url = "%s%s" % (_app_url(), route)
            queued_email_id = None
            if body.to_email:
                queued_email_id = queue_email(ac, kind=EMAIL_KIND_REQUEST, to_email=body.to_email,
                                              template=EMAIL_KIND_REQUEST, payload={
                                                  "subject": "%s asks for the trial balance for %s"
                                                             % (body.firm_name or "Your accountant",
                                                                period_end),
                                                  "vars": {
                                                      "client_name": org.get("name"),
                                                      "firm_name": body.firm_name or "Your accountant",
                                                      "period_end": period_end,
                                                      "upload_url": upload_url,
                                                      "expires_at": _iso(expires),
                                                      "note": body.note or "",
                                                  }}, firm_id=firm_id,
                                              client_org_id=body.client_org_id,
                                              request_id=request_id, user_id=user_id)
        view = public_request_view(row, now.date())
        view.update({"token": token, "upload_route": route, "upload_url": upload_url,
                     "queued_email_id": queued_email_id, "expected_identity": expected,
                     "api_version": FIRM_REQUESTS_API_VERSION})
        return view

    # Two literal segments on every READ: the tenancy router (mounted
    # first) owns `GET /api/firm/{firm_id}`, which would capture a
    # single-segment `GET /api/firm/requests` as a firm id. Asserted by
    # tests/engine/test_firm_gates.py against the real routers.
    @router.get("/requests/list")
    def list_requests(client_org_id: str = Query(...), firm_id: Optional[str] = Query(None),
                      status: Optional[str] = Query(None),
                      authorization: Optional[str] = Header(None)) -> Dict[str, Any]:
        """The client's requests as THIS firm sees them: pinned to the
        firm (a workspace detached from firm X and attached to firm Y
        does not carry X's request notes to Y — C3) and read through the
        caller's own client, so `firm_file_requests firm read` (RLS) is
        the wall even if the Python guard above it were removed."""
        jwt = _require_jwt(authorization)
        _user, firm = authorize_client(jwt, firm_id, client_org_id, "read")
        filters = {"client_org_id": "eq.%s" % client_org_id}
        if firm:
            filters["firm_id"] = "eq.%s" % firm
        if status:
            filters["status"] = "eq.%s" % status
        with _supabase.per_user(jwt) as client:
            rows = select_all(client, "firm_file_requests", order="requested_at.desc,id.asc",
                              filters=filters)
        today = _now().date()
        return {"requests": [public_request_view(r, today) for r in rows or []],
                "as_of": today.isoformat(), "api_version": FIRM_REQUESTS_API_VERSION}

    @router.post("/requests/{request_id}/revoke")
    def revoke_request(request_id: str, firm_id: Optional[str] = Query(None),
                       authorization: Optional[str] = Header(None)) -> Dict[str, Any]:
        jwt = _require_jwt(authorization)
        if firm_id:
            # Membership in the named firm is decided BEFORE any read, with
            # the one message every firm route gives a non-member.
            from . import _firm
            _firm.resolve_firm(jwt, firm_id)
        # The row is read as the CALLER (RLS): a request of another firm's
        # client is not found, which is the same answer as a request that
        # does not exist — a probe learns nothing. The revocation is the
        # CALLER's update too — `firm_file_requests firm update` (RLS) is
        # the second wall — and it is re-read as the caller afterwards: an
        # update RLS filtered to zero rows is a silent no-op at PostgREST,
        # and a route that then said "revoked" would be lying.
        with _supabase.per_user(jwt) as client:
            rows = client.select("firm_file_requests", filters={"id": "eq.%s" % request_id},
                                 limit=1)
            if not rows:
                raise HTTPException(404, "Request not found.")
            row = rows[0]
            authorize_client(jwt, firm_id or row.get("firm_id"), str(row.get("client_org_id")),
                             "request_file")
            if str(row.get("status")) not in OPEN_STATUSES:
                raise HTTPException(409, "Only an open request can be revoked.")
            write_as_caller(lambda: client.update("firm_file_requests", {"status": STATUS_REVOKED},
                                                  filters={"id": "eq.%s" % request_id}),
                            "firm_file_requests")
            after = client.select("firm_file_requests",
                                  filters={"id": "eq.%s" % request_id,
                                           "status": "eq.%s" % STATUS_REVOKED}, limit=1)
        if not after:
            raise HTTPException(403, "Write to firm_file_requests refused by row-level "
                                     "security: the revocation did not apply.")
        row["status"] = STATUS_REVOKED
        return public_request_view(row, _now().date())

    # ── the public half: resolve + upload (the token is the auth) ────────

    @router.get("/requests/resolve/{token}")
    def resolve_request(token: str) -> Dict[str, Any]:
        """What the upload page prefills: the client, the period (the
        period picker is locked to it), and what the link still allows."""
        with _supabase.admin() as ac:
            claims, row = open_request_for(ac, token)
            org = _org_row(ac, claims.client_org_id)
        return {
            "request_id": claims.request_id, "client_org_id": claims.client_org_id,
            "client_name": org.get("name"), "period_end": claims.period_end,
            "expires_at": claims.expires_at, "status": row.get("status"),
            "note": row.get("note"), "single_use": True,
            "upload_route": "/api/firm/requests/%s/upload" % token,
            "api_version": FIRM_REQUESTS_API_VERSION,
        }

    @router.post("/requests/{token}/upload", status_code=202)
    async def upload_for_request(token: str, file: UploadFile = File(...)) -> Dict[str, Any]:
        content = await file.read()
        filename = file.filename or "upload.bin"
        mime = file.content_type or "application/octet-stream"
        with _supabase.admin() as ac:
            claims, row = open_request_for(ac, token)
            org = _org_row(ac, claims.client_org_id)
            row = dict(row, jurisdiction=(org.get("country") or default_jurisdiction()))
            try:
                result = land_file(row, content, filename, mime, production_deps())
            except LandingRefused as exc:
                detail = dict(exc.detail)
                if detail.get("code") == "entity_mismatch":
                    ac.update("firm_file_requests", {
                        "refusal_count": int(row.get("refusal_count") or 0) + 1,
                        "entity_guard": detail.get("entity_guard"),
                        "period_guard": detail.get("period_guard"),
                    }, filters={"id": "eq.%s" % row["id"]})
                raise HTTPException(exc.status, detail)
            ac.update("firm_file_requests", {
                "status": STATUS_RECEIVED, "consumed_at": _iso(_now()),
                "document_id": result.document_id,
                "entity_guard": result.inspection.entity.to_payload(),
                "period_guard": result.inspection.period,
            }, filters={"id": "eq.%s" % row["id"]})
        payload = result.to_payload()
        payload.update({"request_id": claims.request_id, "client_org_id": claims.client_org_id,
                        "period_end": claims.period_end, "status": "queued"})
        return payload

    # ── cadence ─────────────────────────────────────────────────────────

    @router.get("/cadence/status")
    def get_cadence(client_org_id: str = Query(...), firm_id: Optional[str] = Query(None),
                    as_of: Optional[str] = Query(None),
                    authorization: Optional[str] = Header(None)) -> Dict[str, Any]:
        jwt = _require_jwt(authorization)
        _user, firm = authorize_client(jwt, firm_id, client_org_id, "read")
        day = date.fromisoformat(as_of) if as_of else _now().date()
        pack = CAD.load_cadence_pack()
        # Both reads as the CALLER: `firm_client_cadence firm read` and
        # `financial_periods firm read` are the second wall; the cadence
        # row is pinned to the firm the caller acts through (W3).
        with _supabase.per_user(jwt) as client:
            stored = _cadence_row(client, client_org_id, firm)
            filed = _filed_period_ends(client, client_org_id)
        try:
            cadence = CAD.resolve_client_cadence(client_org_id, stored, pack)
            status = CAD.assess(cadence, filed, day, pack=pack)
        except CAD.CadenceInputError as exc:
            raise HTTPException(422, str(exc))
        nudges = CAD.nudges_due(cadence, filed, day)
        return {"cadence": cadence.to_payload(), "status": status.to_payload(),
                "filed_period_ends": filed, "nudges_due": [n.to_payload() for n in nudges],
                "pack": {"source": pack.source, "default_cadence": pack.default_cadence,
                         "cadences": sorted(pack.cadences)},
                "api_version": FIRM_REQUESTS_API_VERSION}

    @router.put("/cadence")
    def put_cadence(body: CadenceBody, authorization: Optional[str] = Header(None)) -> Dict[str, Any]:
        jwt = _require_jwt(authorization)
        user_id, firm_id = authorize_client(jwt, body.firm_id, body.client_org_id, "request_file")
        pack = CAD.load_cadence_pack()
        row = {
            "client_org_id": body.client_org_id, "firm_id": firm_id, "cadence": body.cadence,
            "deadline_days_after_period_end": body.deadline_days_after_period_end,
            "fiscal_year_end_month": body.fiscal_year_end_month or 12,
            "nudge_days_before": body.nudge_days_before, "set_by": user_id,
        }
        try:
            cadence = CAD.resolve_client_cadence(body.client_org_id, row, pack)
        except CAD.CadenceInputError as exc:
            raise HTTPException(422, str(exc))
        # The CALLER's upsert: `firm_client_cadence firm write` / `firm
        # update` (RLS) pin the row to the client's CURRENT firm and the
        # caller's `request_file` cell — the second wall on the one write
        # the critics drove cross-firm (W1).
        with _supabase.per_user(jwt) as client:
            write_as_caller(lambda: client.upsert("firm_client_cadence", row,
                                                  on_conflict="client_org_id"),
                            "firm_client_cadence")
        return {"cadence": cadence.to_payload(), "api_version": FIRM_REQUESTS_API_VERSION}

    # ── digest preferences ──────────────────────────────────────────────

    @router.get("/digest/prefs")
    def get_digest_prefs(firm_id: Optional[str] = Query(None),
                         authorization: Optional[str] = Header(None)) -> Dict[str, Any]:
        jwt = _require_jwt(authorization)
        user_id = _user_id(jwt)
        if firm_id:
            from . import _firm
            _firm.resolve_firm(jwt, firm_id)
        with _supabase.per_user(jwt) as client:
            rows = client.select("firm_digest_prefs", filters={
                "user_id": "eq.%s" % user_id,
                "firm_id": ("is.null" if not firm_id else "eq.%s" % firm_id)}, limit=1)
        pref = rows[0] if rows else {"user_id": user_id, "firm_id": firm_id, "enabled": False,
                                     "frequency": "daily", "send_hour_utc": 7,
                                     "last_sent_at": None, "last_item_ids": []}
        return {"prefs": {k: pref.get(k) for k in ("user_id", "firm_id", "enabled", "frequency",
                                                    "send_hour_utc", "last_sent_at")},
                "api_version": FIRM_REQUESTS_API_VERSION}

    @router.put("/digest/prefs")
    def put_digest_prefs(body: DigestPrefsBody,
                         authorization: Optional[str] = Header(None)) -> Dict[str, Any]:
        jwt = _require_jwt(authorization)
        user_id = _user_id(jwt)
        if body.frequency not in ("daily", "weekly"):
            raise HTTPException(422, "frequency must be daily or weekly.")
        if body.firm_id:
            from . import _firm
            _firm.resolve_firm(jwt, body.firm_id)
        # Own row, written AS THE CALLER: `firm_digest_prefs own insert /
        # own update` (RLS) require user_id = auth.uid() AND a firm the
        # caller is a member of — the second wall under resolve_firm.
        patch = {"enabled": body.enabled, "frequency": body.frequency,
                 "send_hour_utc": body.send_hour_utc}
        with _supabase.per_user(jwt) as client:
            rows = client.select("firm_digest_prefs", filters={
                "user_id": "eq.%s" % user_id,
                "firm_id": ("is.null" if not body.firm_id else "eq.%s" % body.firm_id)}, limit=1)
            if rows:
                write_as_caller(lambda: client.update("firm_digest_prefs", patch,
                                                      filters={"id": "eq.%s" % rows[0]["id"]}),
                                "firm_digest_prefs")
            else:
                write_as_caller(lambda: client.insert(
                    "firm_digest_prefs", dict(patch, user_id=user_id, firm_id=body.firm_id),
                    returning=False), "firm_digest_prefs")
        return {"prefs": dict(patch, user_id=user_id, firm_id=body.firm_id),
                "api_version": FIRM_REQUESTS_API_VERSION}

    # ── crons + drain ───────────────────────────────────────────────────

    @router.post("/requests/cron/nudge")
    def cron_nudge(as_of: Optional[str] = Query(None),
                   authorization: Optional[str] = Header(None)) -> Dict[str, Any]:
        _scheduler_token_or_503(authorization)
        now = _now()
        day = date.fromisoformat(as_of) if as_of else now.date()
        with _supabase.admin() as ac:
            return run_nudge_cron(day, now, ac)

    @router.post("/digest/cron/run")
    def cron_digest(as_of: Optional[str] = Query(None),
                    authorization: Optional[str] = Header(None)) -> Dict[str, Any]:
        _scheduler_token_or_503(authorization)
        now = _now()
        day = date.fromisoformat(as_of) if as_of else now.date()
        with _supabase.admin() as ac:
            try:
                return run_digest_cron(day, now, ac)
            except AttentionUnavailable as exc:
                return {"as_of": day.isoformat(), "queued": 0,
                        "gap": "attention items unavailable: %s" % exc}

    @router.post("/email/drain")
    def drain_email(limit: int = Query(200, ge=1, le=1000),
                    authorization: Optional[str] = Header(None)) -> Dict[str, Any]:
        jwt = _require_jwt(authorization)
        # `_user_id` → `_org.resolve_user_id` → a VERIFIED identity
        # (engine.api._jwt: ES256 against Supabase's JWKS). Until
        # 2026-09-05 this was an unverified payload decode, so the
        # allowlist below was gated by a `sub` anyone could type (critic
        # D5). Nothing here ever sends the bearer to PostgREST — this
        # verifier IS the wall.
        user_id = _user_id(jwt)
        from . import _newsletter
        if not _newsletter._is_admin(user_id):
            raise HTTPException(403, "Admin-only. Add the user_id to PRICING_ADMIN_USER_IDS.")
        drained = failed = 0
        with _supabase.admin() as ac:
            try:
                pending = ac.select("firm_email_queue", filters={"status": "eq.queued"},
                                    limit=limit, order="send_at.asc")
            except Exception:  # noqa: BLE001
                return {"drained": 0, "failed": 0, "note": "firm_email_queue unavailable"}
            for row in pending or []:
                to = row.get("to_email")
                if not to:
                    continue
                subject, html = render_queued_email(row)
                result = _email.send_email(to=to, subject=subject, html=html)
                if result.get("ok"):
                    ac.update("firm_email_queue", {"status": "sent", "sent_at": _iso(_now())},
                              filters={"id": "eq.%s" % row["id"]})
                    drained += 1
                else:
                    ac.update("firm_email_queue",
                              {"status": "failed",
                               "error": str(result.get("error") or result.get("reason"))},
                              filters={"id": "eq.%s" % row["id"]})
                    failed += 1
        return {"drained": drained, "failed": failed}

    return router


__all__ = [
    "FIRM_REQUESTS_API_VERSION", "ROUTE_UPLOAD_REQUEST", "SIGNING_KEY_ENV",
    "TOKEN_VERSION", "DOC_BUCKET", "MAX_UPLOAD_BYTES", "OPEN_STATUSES",
    "STATUS_REQUESTED", "STATUS_REMINDED", "STATUS_RECEIVED", "STATUS_EXPIRED",
    "STATUS_REVOKED", "VERDICT_MATCH", "VERDICT_MISMATCH", "VERDICT_UNKNOWN",
    "EMAIL_KIND_REQUEST", "EMAIL_KIND_REMINDER", "EMAIL_KIND_DIGEST",
    "SigningKeyMissing", "TokenError", "TokenClaims", "signing_key", "mint_token",
    "parse_token", "token_hash", "Preamble", "read_preamble", "normalize_cui",
    "normalize_name", "DocumentIdentity", "extract_identity", "EntityVerdict",
    "entity_guard", "period_precheck", "detected_type_for", "UploadInspection",
    "inspect_upload", "LandingRefused", "LandingDeps", "production_deps",
    "LandingResult", "land_file", "email_file_request", "render_queued_email",
    "queue_email", "authorize_client", "client_org_ids_for", "open_request_for",
    "rls_refused", "write_as_caller",
    "default_jurisdiction", "request_status_label",
    "request_token", "public_request_view", "open_request_items", "run_nudge_cron",
    "AttentionUnavailable", "default_report_provider", "run_digest_cron",
    "build_router",
]
