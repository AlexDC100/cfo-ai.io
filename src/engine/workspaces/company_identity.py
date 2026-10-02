"""Which company, and which period, is a stored document about?

Read from the document's OWN bytes — the header rows and sheet names of a
workbook, the first page of a PDF — with the evidence for every field kept
next to the value (``sources``), so a reviewer can find each answer in the
document.

SIGNALS, strongest first
------------------------
cui
  1. ``document_header_cui``  — a fiscal code printed in the header block
     ("Cod fiscal: 12345674", "C.U.I: RO12345674", "c.f. 12345674"), whose
     Romanian control digit (key 753217532) checks out. A number that fails
     the checksum is NOT a CUI and is ignored.
  2. ``registry_name_match``  — the company name the document prints (a
     "Societate:" line, an "... SRL" title line, a non-generic sheet name)
     matches exactly ONE registered company after normalisation — decided
     over EVERY registered name (``registry_match_name``), never over one
     capped page of prefix hits — AND that company files financial
     statements (at least one filing in the registry). A registered name
     with no filing (dormant, struck off, never filed: "CARNIPROD SRL"
     4705349 for the brand word on a sheet named "Carniprod", which six
     registered companies carry) is recorded as ``sources["cui_hint"]``
     and mints nothing — unless the document prints that CUI's digits.
     When the document LABELS its company ("Societate: ...", "Firma: ..."),
     that labelled name is the only one looked up: unresolved (no match,
     several, a namesake that files nothing) leaves the company unresolved
     (no CUI; the labelled name) — never the CUI of another string in the
     header, the sheet or the filename.
  3. ``filename_registry_match`` — the same, from a name read out of the
     filename. The weakest signal: only names of at least six letters, only
     an unambiguous exact match, and only when the DOCUMENT corroborates it
     (the CUI's digits appear in its header block). Otherwise the match is
     recorded as ``sources["cui_hint"]`` for the operator and mints nothing:
     a filename alone never makes a company key.
  An operator-verified identity (``apply_known_identity``) is applied on top
  by the caller; it never overrides a CUI the document itself prints.

company_name
  the registry's registered name for the CUI, else the name the document
  prints, else a non-generic sheet name. A name read ONLY from the filename
  is recorded (signal ``filename``) but never becomes a company key: a
  filename is typed by a person and says nothing the document confirms.

period_end
  ``engine.api._period_detect.detect_period`` over the header text, so the
  engine's own ranking applies: a period the document states, then a date
  beside closing-balance vocabulary, then the filename. The filename never
  overrides the document. A period RANGE ("pe perioada 01.01.2025 -
  31.12.2025", "Ianuarie - Decembrie 2024", in the header or the filename)
  resolves to its END — the month the book closes — never to its start:
  the range is recognised lexically and only its end is handed to the
  engine's reader.

caen_code / industry_key
  the registry's CAEN for the CUI (else a "CAEN: nnnn" printed in the
  header), mapped to an industry through the repo's own catalogue
  (``engine/api/seed/caen_industry_mappings.yaml`` — the file production's
  ``caen_industry_mappings`` table is loaded from).

Nothing here writes anything — beyond the registry's derived name index
(``engine.workspaces.registry_names``, a file beside the registry) — calls a
network or reads a clock (beyond the engine's period helper refusing "today"
as evidence).

Python 3.9 — no ``match``, no ``X | Y`` unions.
"""
from __future__ import annotations

import fnmatch
import functools
import io
import logging
import re
import unicodedata
import weakref
from dataclasses import dataclass, field, replace
from datetime import date, datetime
from typing import Any, Dict, Iterable, List, Mapping, Optional, Sequence, Tuple

logger = logging.getLogger(__name__)

# ── CUI ────────────────────────────────────────────────────────────────

#: The Romanian CUI control key (ANAF). The body (all digits but the last)
#: is left-padded to 9 digits, multiplied digit-wise by this key, summed,
#: times 10 mod 11 (10 -> 0) and compared with the last digit.
CUI_CONTROL_KEY = "753217532"


def cui_control_digit(body: str) -> int:
    """Control digit for a CUI body (the CUI without its last digit)."""
    digits = str(body).rjust(len(CUI_CONTROL_KEY), "0")
    if len(digits) != len(CUI_CONTROL_KEY) or not digits.isdigit():
        raise ValueError("a CUI body has at most 9 digits: %r" % (body,))
    total = sum(int(d) * int(k) for d, k in zip(digits, CUI_CONTROL_KEY))
    check = (total * 10) % 11
    return 0 if check == 10 else check


def normalize_cui(raw: Any) -> Optional[str]:
    """Digits-only CUI when ``raw`` is a valid one, else None.

    Accepts the printed shapes ("RO 12345674", "12345674", "ro12345674").
    Refuses anything whose control digit does not check out — a wrong
    digit is a different (or no) company, never a near miss."""
    if raw is None:
        return None
    s = re.sub(r"\s+", "", str(raw)).upper()
    if s.startswith("RO"):
        s = s[2:]
    if not s.isdigit():
        return None
    s = s.lstrip("0")
    if not 2 <= len(s) <= 10:
        return None
    if cui_control_digit(s[:-1]) != int(s[-1]):
        return None
    return s


# ── company names ──────────────────────────────────────────────────────

_LEGAL_FORM_SUFFIXES = (
    ("S", "R", "L", "D"), ("S", "R", "L"), ("S", "A"), ("S", "N", "C"),
    ("S", "C", "S"), ("S", "C", "A"), ("SRLD",), ("SRL",), ("SA",), ("SNC",),
    ("SCS",), ("SCA",), ("RA",), ("PFA",), ("II",), ("IF",),
)


def normalize_company_name(name: Any) -> str:
    """Upper-case, diacritics and punctuation folded, legal form dropped.

    "OMEGA'S FOOD FACTORY SRL" and "Omegas Food Factory" both give
    "OMEGAS FOOD FACTORY"; "S.C. Alfa Food S.R.L." gives "ALFA FOOD".
    Deterministic and total: any input gives a (possibly empty) string."""
    if name is None:
        return ""
    s = unicodedata.normalize("NFKD", str(name))
    s = "".join(ch for ch in s if not unicodedata.combining(ch)).upper()
    s = re.sub(r"[.'`’´]", "", s)          # S.R.L. -> SRL, AGRA'S -> AGRAS
    s = re.sub(r"[^A-Z0-9]+", " ", s).strip()
    tokens = s.split()
    changed = True
    while changed and tokens:
        changed = False
        for form in _LEGAL_FORM_SUFFIXES:
            n = len(form)
            if len(tokens) > n and tuple(tokens[-n:]) == form:
                tokens = tokens[:-n]
                changed = True
                break
    if len(tokens) > 1 and tokens[0] == "SC":           # "S.C. X S.R.L."
        tokens = tokens[1:]
    return " ".join(tokens)


# ── the identity ───────────────────────────────────────────────────────

#: Signals whose company NAME is strong enough to key a company that has
#: no CUI. The filename is deliberately absent.
NAME_KEY_SIGNALS = frozenset({
    "document_header_label", "document_header_title", "sheet_name",
    "operator_verified", "registry",
})

DOCUMENT_KINDS = ("trial_balance", "not_a_balance", "uncertain", "unreadable")


@dataclass(frozen=True)
class CompanyIdentity:
    """Who a document is about. ``sources`` maps each field to
    ``{"signal": ..., "evidence": ...}`` — the literal text that produced it."""

    cui: Optional[str] = None
    company_name: Optional[str] = None
    period_end: Optional[str] = None
    caen_code: Optional[str] = None
    industry_key: Optional[str] = None
    sources: Dict[str, Dict[str, str]] = field(default_factory=dict, hash=False)
    #: "trial_balance" (account rows read), "not_a_balance" (readable text,
    #: no account rows AND none of the signals a balance carries — a printed
    #: CUI, a balance title, ledger columns; e.g. a delegation itinerary),
    #: "uncertain" (too few rows read — a layout the reader refused is
    #: never condemned as "not a balance"), "unreadable" (no bytes /
    #: unparseable / image-only).
    document_kind: str = "uncertain"

    @property
    def company_key(self) -> Optional[str]:
        """``cui:<digits>`` — or ``name:<normalized>`` for a company whose
        name the DOCUMENT (or an operator) states but no CUI is known — or
        None when nothing but a filename speaks for it."""
        if self.cui:
            return "cui:" + self.cui
        signal = (self.sources.get("company_name") or {}).get("signal")
        if self.company_name and signal in NAME_KEY_SIGNALS:
            norm = normalize_company_name(self.company_name)
            if norm:
                return "name:" + norm
        return None

    def to_dict(self) -> Dict[str, Any]:
        return {
            "cui": self.cui,
            "company_name": self.company_name,
            "period_end": self.period_end,
            "caen_code": self.caen_code,
            "industry_key": self.industry_key,
            "document_kind": self.document_kind,
            "company_key": self.company_key,
            "sources": {k: dict(v) for k, v in sorted(self.sources.items())},
        }

    @classmethod
    def from_dict(cls, d: Mapping[str, Any]) -> "CompanyIdentity":
        return cls(
            cui=d.get("cui"),
            company_name=d.get("company_name"),
            period_end=d.get("period_end"),
            caen_code=d.get("caen_code"),
            industry_key=d.get("industry_key"),
            sources={k: dict(v) for k, v in (d.get("sources") or {}).items()},
            document_kind=d.get("document_kind") or "uncertain",
        )


# ── reading the bytes ──────────────────────────────────────────────────

#: A trial-balance account code: RAS classes 1-8, a 3-6 digit synthetic
#: account, optional analytic suffixes ("4111.00123", "1012.01").
_ACCOUNT_CELL_RE = re.compile(r"^[1-8]\d{2,5}(?:[.\-/]\d{1,6}){0,3}$")
_ACCOUNT_LINE_RE = re.compile(r"^[1-8]\d{2,5}(?:[.\-]\d{1,6}){0,3}\s+\S.*\d")
_NUMERIC_CELL_RE = re.compile(r"^-?[\d\s.,]+$")

#: Rows / lines read. Enough for the title block and to count account rows;
#: never the whole ledger.
MAX_SHEET_ROWS = 1500
MAX_HEADER_LINES = 25
PDF_PAGES = 2

#: Thresholds for ``document_kind``.
BALANCE_MIN_ACCOUNT_LINES = 10
NOT_A_BALANCE_MAX_ACCOUNT_LINES = 2
NOT_A_BALANCE_MIN_CHARS = 80


@dataclass(frozen=True)
class DocumentText:
    """What the identity is read from."""

    readable: bool
    header_lines: Tuple[str, ...] = ()
    sheet_names: Tuple[str, ...] = ()
    account_lines: int = 0
    text_chars: int = 0
    read_error: Optional[str] = None

    @property
    def header_text(self) -> str:
        return "\n".join(self.header_lines)

    @property
    def document_kind(self) -> str:
        if not self.readable:
            return "unreadable"
        if self.account_lines >= BALANCE_MIN_ACCOUNT_LINES:
            return "trial_balance"
        if (self.account_lines <= NOT_A_BALANCE_MAX_ACCOUNT_LINES
                and self.text_chars >= NOT_A_BALANCE_MIN_CHARS):
            return "not_a_balance"
        return "uncertain"


def _cell_text(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, float):
        if v != v:  # NaN
            return ""
        if v.is_integer():
            return str(int(v))
        return repr(v)
    if isinstance(v, datetime):
        return v.strftime("%d.%m.%Y")
    if isinstance(v, date):
        return v.strftime("%d.%m.%Y")
    return str(v).strip()


def _rows_to_text(rows: Iterable[Sequence[Any]]) -> Tuple[List[str], int, int]:
    """(header lines, account rows, text chars) from a sheet's rows.

    Header lines are the rows ABOVE the first account row (the title block
    and the column header). Account names further down routinely carry
    OTHER companies' names (analytic customer / supplier accounts), so the
    identity is never read from them."""
    header: List[str] = []
    accounts = 0
    chars = 0
    seen_account = False
    for i, row in enumerate(rows):
        if i >= MAX_SHEET_ROWS:
            break
        cells = [_cell_text(v) for v in row]
        cells = [c for c in cells if c]
        if not cells:
            continue
        line = " ".join(cells)
        chars += len(line)
        first = cells[0]
        numeric_others = sum(1 for c in cells[1:] if _NUMERIC_CELL_RE.match(c))
        if _ACCOUNT_CELL_RE.match(first) and numeric_others >= 2:
            accounts += 1
            seen_account = True
            continue
        if not seen_account and len(header) < MAX_HEADER_LINES:
            header.append(line)
    return header, accounts, chars


def _read_xlsx(content: bytes) -> DocumentText:
    import openpyxl  # noqa: WPS433 — heavy, imported on use

    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        names = tuple(str(n) for n in wb.sheetnames)
        best = None  # type: Optional[Tuple[List[str], int, int]]
        for ws in wb.worksheets[:3]:
            got = _rows_to_text(ws.iter_rows(values_only=True))
            if best is None or got[1] > best[1]:
                best = got
            if got[1] >= BALANCE_MIN_ACCOUNT_LINES:
                break
    finally:
        wb.close()
    header, accounts, chars = best or ([], 0, 0)
    return DocumentText(True, tuple(header), names, accounts, chars)


def _read_xls(content: bytes) -> DocumentText:
    import xlrd  # noqa: WPS433

    book = xlrd.open_workbook(file_contents=content, on_demand=True)
    try:
        names = tuple(str(n) for n in book.sheet_names())
        best = None  # type: Optional[Tuple[List[str], int, int]]
        for idx in range(min(3, book.nsheets)):
            sh = book.sheet_by_index(idx)
            rows = (sh.row_values(r) for r in range(min(sh.nrows, MAX_SHEET_ROWS)))
            got = _rows_to_text(rows)
            if best is None or got[1] > best[1]:
                best = got
            if got[1] >= BALANCE_MIN_ACCOUNT_LINES:
                break
    finally:
        book.release_resources()
    header, accounts, chars = best or ([], 0, 0)
    return DocumentText(True, tuple(header), names, accounts, chars)


def _read_pdf(content: bytes) -> DocumentText:
    import pdfplumber  # noqa: WPS433

    lines: List[str] = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for page in pdf.pages[:PDF_PAGES]:
            lines.extend((page.extract_text() or "").splitlines())
    header: List[str] = []
    accounts = 0
    chars = 0
    seen_account = False
    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        chars += len(line)
        if _ACCOUNT_LINE_RE.match(line):
            accounts += 1
            seen_account = True
            continue
        if not seen_account and len(header) < MAX_HEADER_LINES:
            header.append(line)
    if chars == 0:
        # An image-only (scanned) PDF: nothing to read is not evidence of
        # "not a balance".
        return DocumentText(False, read_error="no text layer")
    return DocumentText(True, tuple(header), (), accounts, chars)


def extract_document_text(content: bytes, filename: Optional[str] = None,
                          mime: Optional[str] = None) -> DocumentText:
    """Header lines, sheet names and an account-row count, or an
    ``unreadable`` result naming why. Never raises.

    READ BY REAL TYPE — the BYTES pick the reader, never the name (review
    2026-10-02, HIGH). This used to test ``name.endswith(".pdf")`` first, so
    a workbook named balanta.pdf — a file the upload card and the pipeline
    both READ since the one upload policy — went to the PDF text reader,
    came back ``unreadable`` with no CUI, and (a) the card filed it in the
    company on screen and (b) ``pipeline._document_company_cui`` found no CUI
    to hold against the month's owner, so the cross-company takeover wall
    let another company's book replace a served month.

    THE DISPATCH IS THE PIPELINE'S OWN LABEL, not a second reading of the
    magic bytes (review round 4, 2026-10-02). ``_upload_type.sniff_container``
    names the container; a workbook label (``XLSX`` / ``OOXML_UNKNOWN``) goes
    to the workbook reader, ``XLS_OLE2`` to the legacy one, and bytes the
    pipeline reads as a PDF on the branch the NAME AND THE DECLARED MIME
    TYPE lead to (``reads_as_pdf(classify(filename, mime), …)``) go to the
    PDF reader. Two byte shapes had the round-3 dispatch answer differently
    from the pipeline, and both walked through the takeover wall:

    · it tested ``content[:2] == b"PK"``, where the sniffer demands the
      four-byte zip signature. ``b"PK"`` + junk + a balance PDF is a PDF to
      the pipeline under every name; here it went to the workbook reader and
      came back ``unreadable``.
    · it classified with no MIME type. A PDF whose header sits past the
      sniffer's window and whose trailer sits before the last 2,048 bytes is
      a PDF to the pipeline exactly where the branch is ``pdf`` or ``xlsx``
      — and the branch of a name with no extension is the declared MIME
      type's. Named ``balanta`` with ``application/pdf`` the pipeline read
      it and this function answered ``unreadable``.

    ``mime`` is the type the upload declared (the routes hold it; a stored
    document carries it as ``mime_type``). A ``.pdf`` name is consulted LAST,
    for bytes nothing above names: the PDF reader gets its try and its
    failure is the ``unreadable`` answer."""
    if not content:
        return DocumentText(False, read_error="no bytes")
    name = (filename or "").lower()
    try:
        from engine.api import _upload_type as _ut  # noqa: WPS433 — lazy (pipeline cycle); pure
        real = _ut.sniff_container(content)
        if real in (_ut.XLSX, _ut.OOXML_UNKNOWN):
            return _read_xlsx(content)
        if real == _ut.XLS_OLE2:
            return _read_xls(content)
        if _ut.reads_as_pdf(_ut.classify(filename, mime), real, content) or name.endswith(".pdf"):
            return _read_pdf(content)
    except Exception as exc:  # noqa: BLE001 — an unreadable file is an answer
        return DocumentText(False, read_error="%s: %s" % (type(exc).__name__, str(exc)[:160]))
    return DocumentText(False, read_error="unsupported format")


# ── header parsing ─────────────────────────────────────────────────────

_CUI_RE = re.compile(
    r"(?:\bC\.?\s*U\.?\s*I\b\.?|\bC\.?\s*I\.?\s*F\b\.?|\bc\.\s*f\.|"
    r"\bcod\s+(?:de\s+identificare\s+)?fiscal[aă]?\b|"
    r"\bcod\s+unic(?:\s+de\s+[iî]nregistrare)?\b)"
    r"\s*[:.]?\s*(?:RO\s*)?(\d{2,10})\b",
    re.IGNORECASE,
)
_CAEN_RE = re.compile(r"(?i)\b(?:cod\s+)?CAEN\b\s*[:.]?\s*(\d{3,4})\b")
_NAME_LABEL_RE = re.compile(
    r"(?i)\b(?:societatea|societate|unitatea|firma|denumirea\s+(?:firmei|societ[aă][tț]ii|unit[aă][tț]ii)|"
    r"denumire\s+(?:firm[aă]|societate)|company|entity)\s*[:]\s*(.+)"
)
#: "<Name> SRL" / "<NAME> S.R.L." / "<Name> SA" — the legal form is matched
#: upper-case only, so the Romanian word "sa" never reads as a company.
_TITLE_RE = re.compile(
    r"(?<![\w&])([A-Z0-9ĂÂÎȘŞȚŢ]"
    r"[\w&.'`’\- ]{1,70}?\s(?:S\.?\s?R\.?\s?L\.?|S\.?\s?A\.?))(?=[\s,;]|$)"
)
_NAME_CUTS = re.compile(
    r"(?i)\s+(?:adresa|adres[aă]|str\.|strada|sediul|c\.?\s*u\.?\s*i|c\.?\s*i\.?\s*f|c\.\s*f\.|"
    r"cod\s+fiscal|nr\.?\s*reg|r\.\s*c\.|j\d{2}|capital)\b.*$"
)
_GENERIC_SHEET_RE = re.compile(
    r"(?i)^(?:sheet|foaie|foaia|page|pagina|tabel|table|data|date|export|report|raport|"
    r"worksheet|tb|bal|balanta|balanța|balanta\s+de\s+verificare|trial\s*balance|"
    r"document(?:_ch)?|bs|pl|pnl|cont|conturi|sintetic|analitic|analitica|sintetica)?[\s_\-]*\d*$"
)
_FILENAME_STOP = frozenset("""
balanta balanța balance trial tb verificare de si și analitica analitică sintetica sintetică
lv extern externa final finala copy copie fy an anul luna the of and export raport report
ianuarie februarie martie aprilie mai iunie iulie august septembrie octombrie noiembrie decembrie
ian feb mar apr iun iul aug sep sept oct noi nov dec jan january february march april may june
july august september october november december
""".split())


def _clean_name(raw: str) -> str:
    s = _NAME_CUTS.sub("", raw).strip(" \t,;:.-")
    return re.sub(r"\s{2,}", " ", s)


def _header_cui(lines: Sequence[str]) -> Optional[Tuple[str, str]]:
    for line in lines:
        for m in _CUI_RE.finditer(line):
            cui = normalize_cui(m.group(1))
            if cui:
                return cui, line.strip()[:160]
    return None


#: A balance names itself ("Balanta de verificare", "Balanță", "trial
#: balance") and its columns (sold / rulaj / debit / credit / cont).
_BALANCE_TITLE_RE = re.compile(r"(?i)\bbalan[tţț][aăâ]?\b|\btrial\s*balance\b")
_LEDGER_WORDS_RE = re.compile(r"(?i)\b(?:sold(?:uri)?|rulaj(?:e)?|debit(?:oare)?|credit(?:oare)?|cont(?:uri)?)\b")


def balance_signals(lines: Sequence[str], filename: Optional[str], printed_cui: Optional[str]) -> List[str]:
    """Why readable text with no account rows is still NOT "demonstrably
    not a trial balance": the header prints a CUI, the header (or the file)
    names a balance in its title, or the header carries the ledger's column
    vocabulary. Any of these means the row READER refused the layout — the
    2026-09-26 case: a balanță whose rows the deterministic readers could
    not parse, whose header printed the company and its CUI, was classified
    not_a_balance, archived and moved into a holding archive."""
    out: List[str] = []
    if printed_cui:
        out.append("prints CUI %s" % printed_cui)
    text = "\n".join(lines)
    if _BALANCE_TITLE_RE.search(text):
        out.append("a balance title in the header")
    elif filename and _BALANCE_TITLE_RE.search(str(filename)):
        out.append("a balance title in the filename")
    words = sorted({m.group(0).lower()[:4] for m in _LEDGER_WORDS_RE.finditer(text)})
    if len(words) >= 2:
        out.append("ledger columns (%s)" % ", ".join(words))
    return out


def _digits_in(lines: Sequence[str], cui: str) -> bool:
    """The CUI printed as a standalone number (RO prefix allowed) anywhere
    in the header block."""
    pat = re.compile(r"(?<!\d)(?:RO\s*)?0*%s(?!\d)" % re.escape(str(cui)), re.IGNORECASE)
    return any(pat.search(line) for line in lines)


def _header_caen(lines: Sequence[str]) -> Optional[Tuple[str, str]]:
    for line in lines:
        m = _CAEN_RE.search(line)
        if m:
            return _normalize_caen(m.group(1)), line.strip()[:160]
    return None


_TITLE_FORM_RE = re.compile(r"^(.*?)\s+((?:S\.?\s?R\.?\s?L\.?|S\.?\s?A\.?))$")
#: Lowercase words a company name carries between capitalised ones
#: ("Casa de Ajutor Reciproc").
_TITLE_CONNECTORS = frozenset({"de", "si", "și", "şi", "and", "of", "la", "din", "pentru", "&", "-"})
#: Words that introduce a name rather than belong to it ("Firma ALFA FOOD SRL").
_TITLE_LABELS = frozenset({"firma", "societatea", "societate", "unitatea", "company", "entity", "client",
                           "furnizor", "beneficiar", "emitent"})
TITLE_MAX_TOKENS = 6


def _trim_title(capture: str) -> Optional[str]:
    """The company name inside a "<words> SRL" capture: the run of
    capitalised words (ALL-CAPS words when the name is written in capitals)
    directly before the legal form, through lowercase connectors, without a
    leading label word, at most TITLE_MAX_TOKENS words. "Welcome dinner
    hosted by Scandia Food S.R.L" -> "Scandia Food S.R.L"; "Firma ALFA FOOD
    SRL" -> "ALFA FOOD SRL". None when no such run exists."""
    m = _TITLE_FORM_RE.match(capture.strip())
    if not m:
        return None
    tokens, form = m.group(1).split(), m.group(2)
    if not tokens:
        return None

    def caps(tok: str) -> bool:
        letters = [c for c in tok if c.isalpha()]
        return bool(letters) and all(c.isupper() for c in letters)

    all_caps = caps(tokens[-1])

    def ok(tok: str) -> bool:
        if tok == "&" or tok.isdigit():
            return True
        return caps(tok) if all_caps else (tok[:1].isupper() or tok[:1].isdigit())

    run: List[str] = []
    i = len(tokens) - 1
    while i >= 0:
        tok = tokens[i]
        if ok(tok):
            run.insert(0, tok)
        elif run and tok.lower() in _TITLE_CONNECTORS and i > 0 and ok(tokens[i - 1]):
            run.insert(0, tok)
        else:
            break
        i -= 1
    while run and run[0].lower().strip(".:,") in _TITLE_LABELS:
        run.pop(0)
    while run and run[0].lower() in _TITLE_CONNECTORS:
        run.pop(0)
    if not run:
        return None
    return " ".join(run[-TITLE_MAX_TOKENS:] + [form])


def _header_names(lines: Sequence[str]) -> List[Tuple[str, str, str]]:
    """[(name, signal, evidence)] in reading order: labelled names first."""
    out: List[Tuple[str, str, str]] = []
    for line in lines:
        m = _NAME_LABEL_RE.search(line)
        if m:
            name = _clean_name(m.group(1))
            if len(normalize_company_name(name)) >= 2:
                out.append((name, "document_header_label", line.strip()[:160]))
    for line in lines:
        for m in _TITLE_RE.finditer(line):
            trimmed = _trim_title(m.group(1))
            if not trimmed:
                continue
            name = _clean_name(trimmed)
            if len(normalize_company_name(name)) >= 2:
                out.append((name, "document_header_title", line.strip()[:160]))
    return out


def _sheet_names(names: Sequence[str]) -> List[str]:
    out = []
    for n in names:
        s = str(n).strip()
        if not s or _GENERIC_SHEET_RE.match(s):
            continue
        if len(re.sub(r"[^A-Za-z]", "", s)) < 3:
            continue
        out.append(s)
    return out


def filename_company_name(filename: Optional[str]) -> Optional[str]:
    """The words of a filename that are neither dates nor balance
    vocabulary ("Balanta Alfa Food_FY2025.xls" -> "Alfa Food"), or
    None. A reading aid and a weak registry key — never a company key."""
    if not filename:
        return None
    stem = re.sub(r"\.[A-Za-z0-9]{2,5}$", "", str(filename))
    stem = re.sub(r"\d{1,2}[.\-_]\d{1,2}[.\-_]\d{2,4}", " ", stem)
    words = re.split(r"[\s_\-.()\[\]]+", stem)
    kept = []
    for w in words:
        lw = w.lower()
        if not w or lw in _FILENAME_STOP:
            continue
        if re.fullmatch(r"(?i)(?:fy|v|an)?\d+", w):
            continue
        kept.append(w)
    name = " ".join(kept).strip()
    return name or None


# ── registry ───────────────────────────────────────────────────────────

def _normalize_caen(caen: Any) -> Optional[str]:
    digits = "".join(ch for ch in str(caen or "") if ch.isdigit())
    if not digits:
        return None
    if len(digits) == 3:
        digits = "0" + digits
    return digits[:4]


def _registry_company(registry: Any, cui: str) -> Optional[Dict[str, Any]]:
    if registry is None:
        return None
    try:
        row = registry.get_company(int(cui))
    except Exception:  # noqa: BLE001 — a registry outage is "not found"
        return None
    return dict(row) if row else None


#: Rows asked of ``search_companies`` per query (its own cap is 100). A
#: query that comes back FULL may have cut off the match, so a full page
#: without a unique exact match is "unknown", never a guess.
REGISTRY_SEARCH_LIMIT = 100
FILENAME_NAME_MIN_CHARS = 6


#: Per-registry memo of name searches: a 600k-row spine answers a
#: case-insensitive prefix query by scanning, and a migration asks the same
#: few names once per document.
_SEARCH_MEMO: "weakref.WeakKeyDictionary[Any, Dict[str, List[Dict[str, Any]]]]" = weakref.WeakKeyDictionary()


def _search(registry: Any, q: str) -> List[Dict[str, Any]]:
    try:
        memo = _SEARCH_MEMO.setdefault(registry, {})
    except TypeError:  # not weak-referenceable: no memo
        memo = {}
    if q not in memo:
        memo[q] = [dict(r) for r in (registry.search_companies(q, limit=REGISTRY_SEARCH_LIMIT) or [])]
    return memo[q]


#: Per-registry index: normalized name -> the one CUI registered under it,
#: or AMBIGUOUS when several are — decided over every registered name, the
#: only way "exactly one" is a fact rather than a guess about what a capped
#: prefix page left out. Building it is a full pass over the spine's names
#: (~1M rows through the normalizer: 5-20 s), so for a registry with a file
#: it is PERSISTED beside it and read with one indexed query per name
#: (``engine.workspaces.registry_names``; live walkthrough, 2026-09-26: the
#: upload route opens a registry per request, and every identify — and the
#: commit after it — rebuilt it: 15-20 s per upload in production). A
#: registry without a file (a test double) keeps the in-memory build. The
#: memo holds the answer per registry object.
_NAME_INDEX_MEMO: "weakref.WeakKeyDictionary[Any, Mapping[str, int]]" = weakref.WeakKeyDictionary()
AMBIGUOUS = -1


def _code_shape(code: Any) -> Tuple[Any, ...]:
    """A code object's bytecode, names and constants — nested code objects
    (the generator inside the normalizer) by their own shape, never by a
    repr that carries a memory address and changes every process."""
    consts = tuple(_code_shape(c) if hasattr(c, "co_code") else repr(c) for c in code.co_consts)
    return (code.co_code, code.co_names, consts)


@functools.lru_cache(maxsize=1)
def normalizer_tag() -> str:
    """A fingerprint of the normalizer itself — its code, its constants and
    the legal-form table — so a persisted index built by another version of
    ``normalize_company_name`` is never read as this one's. Stable across
    processes (the persisted index outlives each one)."""
    import hashlib

    blob = repr((_code_shape(normalize_company_name.__code__), _LEGAL_FORM_SUFFIXES)).encode("utf-8")
    return hashlib.sha256(blob).hexdigest()[:16]


def _normalized_name_index(registry: Any) -> Optional[Mapping[str, int]]:
    """The registry's names, normalized — or None when the registry cannot
    list its names (then ``registry_match_name`` falls back to search)."""
    names = getattr(registry, "iter_company_names", None)
    if names is None:
        return None
    try:
        cached = _NAME_INDEX_MEMO.get(registry)
    except TypeError:  # not weak-referenceable
        cached = None
    if cached is not None:
        return cached
    from engine.workspaces.registry_names import build_name_map, registry_name_index

    index = registry_name_index(registry, normalize_company_name, normalizer_tag())
    if index is None:
        index = build_name_map(names(), normalize_company_name)
    try:
        _NAME_INDEX_MEMO[registry] = index
    except TypeError:
        pass
    return index


def warm_name_index(registry: Any) -> str:
    """Make the registry's persisted name index current now (the operator
    CLI after an ingest, the server at start) so no request pays for the
    build. Returns a one-line status. Never raises."""
    try:
        index = _normalized_name_index(registry)
    except Exception as exc:  # noqa: BLE001
        logger.exception("[company_identity] name index warm-up failed")
        return "name index: failed (%s)" % type(exc).__name__
    if index is None:
        return "name index: the registry cannot list its names"
    return "name index: %d names, %s" % (len(index), "persisted" if hasattr(index, "fingerprint") else "in memory")



def registry_files(registry: Any, cui: str, row: Optional[Mapping[str, Any]]) -> Optional[bool]:
    """Does the registered company file financial statements? The row's
    ``has_filings`` when the registry gives it, else its filings
    (``get_filings``); None when the registry cannot say."""
    if row is not None and "has_filings" in row:
        return bool(row["has_filings"])
    getter = getattr(registry, "get_filings", None)
    if getter is None:
        return None
    try:
        return bool(getter(int(cui)))
    except Exception:  # noqa: BLE001 — a registry that cannot answer is "unsure"
        return None


def registry_match_name(registry: Any, name: str) -> Optional[Tuple[str, Dict[str, Any]]]:
    """(cui, registry row) when EXACTLY one registered company has this
    normalized name AND it files financial statements; otherwise None (no
    match, several, unsure — or a company with no filing).

    A registered name with no filing is not an identity: the index covers
    every registered name, publishable or not (so that "exactly one" is a
    fact), and the one company whose normalized name is a bare brand word
    is typically a dormant or struck-off namesake. Verifier finding
    (2026-09-21, p10_realstore.py / p11_realstore_plan.py): the sheet
    "Carniprod" resolved to CARNIPROD SRL 4705349 — not publishable, no
    filing, and not the book's company by the operator's own filing check
    — and the plan wrote that stranger's CUI into another user's org_prefs.
    ``identify_text`` records such a match as a hint instead.
    """
    hit = _unique_registered_name(registry, name)
    if hit is None or hit[2] is not True:
        return None
    return hit[0], hit[1]


def _unique_registered_name(registry: Any, name: str
                            ) -> Optional[Tuple[str, Dict[str, Any], Optional[bool]]]:
    """(cui, registry row, files) when EXACTLY one registered company has
    this normalized name; otherwise None (no match, several, or unsure).
    ``files`` is ``registry_files`` for that company.

    "Exactly one" is decided over EVERY registered name
    (``iter_company_names``, normalized once per registry): "ALFA FOOD SRL"
    and "ALFA-FOOD S.R.L." are the same name, and a second company with it
    is found however many other "ALFA …" companies sort before it. Until
    2026-09-21 uniqueness was decided over the prefix hits of the printed
    spelling on pages capped at REGISTRY_SEARCH_LIMIT — a full page with
    one hit was accepted, and punctuation variants were never seen.

    A registry that cannot list its names (a test double, another store)
    is asked through ``search_companies`` and ANY full page is "unsure"."""
    if registry is None or not name:
        return None
    target = normalize_company_name(name)
    if len(target) < 2:
        return None
    try:
        index = _normalized_name_index(registry)
    except Exception:  # noqa: BLE001 — a registry that fails to list is "unsure"
        return None
    if index is not None:
        try:
            cui_int = index.get(target)
        except Exception:  # noqa: BLE001 — an index that cannot answer is "unsure"
            logger.exception("[company_identity] name index lookup failed")
            return None
        if cui_int is None or cui_int == AMBIGUOUS:
            return None
        cui = normalize_cui(cui_int)
        row = _registry_company(registry, cui) if cui else None
        return (cui, row, registry_files(registry, cui, row)) if cui and row else None
    queries = []
    for q in (str(name).strip(), target, target.split()[0]):
        if q and q not in queries:
            queries.append(q)
    hits: Dict[int, Dict[str, Any]] = {}
    for q in queries:
        try:
            rows = _search(registry, q)
        except Exception:  # noqa: BLE001
            return None
        if len(rows) >= REGISTRY_SEARCH_LIMIT:
            # A full page may have cut off a second company with this
            # name: unsure, whatever this page happened to contain.
            return None
        for r in rows:
            if normalize_company_name(r.get("name")) == target:
                hits[int(r["cui"])] = dict(r)
    if len(hits) != 1:
        return None
    cui_int, row = next(iter(hits.items()))
    cui = normalize_cui(cui_int)
    # search_companies answers only companies it may publish, and a company
    # with no filing never qualifies (compliance.publishable_reason: "no
    # filings (thin row)") — a hit here files by the search's own contract.
    files = registry_files(registry, cui, row) if cui else None
    return (cui, row, True if files is None else files) if cui else None


# ── CAEN -> industry (the repo's own catalogue) ────────────────────────

@functools.lru_cache(maxsize=1)
def _caen_catalogue() -> Dict[str, str]:
    """caen_code -> industry_key, from the YAML production's
    ``caen_industry_mappings`` table is seeded from. Empty on any failure:
    an industry is a label, never worth failing an identification."""
    try:
        from engine.api.seed.load_industry_catalog import _MAPPINGS_YAML, _load_yaml
        rows = _load_yaml(_MAPPINGS_YAML)
    except Exception:  # noqa: BLE001
        return {}
    out: Dict[str, str] = {}
    for r in rows:
        code = _normalize_caen(r.get("caen_code"))
        key = r.get("industry_key")
        if code and key:
            out[code] = str(key)
    return out


def industry_key_for_caen(caen: Any) -> Optional[str]:
    code = _normalize_caen(caen)
    return _caen_catalogue().get(code) if code else None


@functools.lru_cache(maxsize=1)
def _industry_names() -> Dict[str, str]:
    """industry_key -> display_name, from the YAML production's
    ``industry_profiles`` table is seeded from. Empty on any failure."""
    try:
        from engine.api.seed.load_industry_catalog import _INDUSTRIES_YAML, _load_yaml
        rows = _load_yaml(_INDUSTRIES_YAML)
    except Exception:  # noqa: BLE001
        return {}
    return {str(r["key"]): str(r["display_name"]) for r in rows if r.get("key") and r.get("display_name")}


def industry_display_name(industry_key: Any) -> Optional[str]:
    return _industry_names().get(str(industry_key)) if industry_key else None


# ── period ranges ──────────────────────────────────────────────────────
#
# "Balanta de verificare pe perioada 01.01.2025 - 31.12.2025" states a
# period, and a period is named by its END — the month the book closes. The
# engine's date reader takes the FIRST date it finds, so the whole line read
# as January 2025 (and "Ianuarie - Decembrie 2024" as January 2024). A range
# is recognised here, LEXICALLY, and only its end is handed to the engine's
# own reader (``detect_period``), so month-end convention, sanity bounds and
# the "today is never evidence" rule stay where they live (_period_detect W5).

_MONTH_WORDS: Tuple[Tuple[int, Tuple[str, ...]], ...] = (
    (1, ("ianuarie", "january", "ian", "jan")),
    (2, ("februarie", "february", "feb")),
    (3, ("martie", "march", "mar")),
    (4, ("aprilie", "april", "apr")),
    (5, ("mai", "may")),
    (6, ("iunie", "june", "iun", "jun")),
    (7, ("iulie", "july", "iul", "jul")),
    (8, ("august", "aug")),
    (9, ("septembrie", "september", "sept", "sep")),
    (10, ("octombrie", "october", "oct")),
    (11, ("noiembrie", "november", "noi", "nov")),
    (12, ("decembrie", "december", "dec")),
)
_MONTH_BY_WORD = {w: m for m, words in _MONTH_WORDS for w in words}
_LETTERS = "A-Za-zĂÂÎȘŞȚŢăâîșşțţ"
_MONTH_ALT = "|".join(sorted(_MONTH_BY_WORD, key=len, reverse=True))
_YEAR_RE = r"(?:19|20)\d{2}"


def _side(p: str) -> str:
    """One end of a range, its parts in named groups prefixed ``p``."""
    month = r"(?<![%s])(?P<%smw{n}>%s)\.?(?![%s])" % (_LETTERS, p, _MONTH_ALT, _LETTERS)
    return "(?:" + "|".join((
        r"(?P<{p}d1>\d{1,2})[./-](?P<{p}m1>\d{1,2})[./-](?P<{p}y1>%s)" % _YEAR_RE,     # 31.12.2025
        r"(?P<{p}y2>%s)-(?P<{p}m2>\d{1,2})-(?P<{p}d2>\d{1,2})" % _YEAR_RE,              # 2025-12-31
        r"(?P<{p}d3>\d{1,2})\s+" + month.replace("{n}", "3") + r"\s+(?P<{p}y3>%s)" % _YEAR_RE,  # 31 dec 2025
        month.replace("{n}", "4") + r"(?:\s+(?P<{p}y4>%s))?" % _YEAR_RE,                     # decembrie [2025]
        r"(?P<{p}m5>\d{1,2})[./-](?P<{p}y5>%s)" % _YEAR_RE,                                # 12.2025
    )).replace("{p}", p) + ")"


_RANGE_SEP = (r"\s*(?:[\-‐‑–—]{1,2}|\bp[aâ]n[aă](?:\s+la)?\b|\bto\b|\bla\b"
              r"|\bthrough\b|\bthru\b)\s*")
_RANGE_RE = re.compile(r"(?<![\d.])" + _side("a_") + _RANGE_SEP + _side("b_") + r"(?![\d])", re.IGNORECASE)


def _range_side(m: "re.Match", p: str) -> Optional[Tuple[Optional[int], int, Optional[int]]]:
    """(year or None, month, day or None) of one side of a range match."""
    g = m.groupdict()
    for n in "12345":
        mw = g.get("%smw%s" % (p, n))
        num = g.get("%sm%s" % (p, n))
        if mw is None and num is None:
            continue
        month = _MONTH_BY_WORD[mw.lower()] if mw else int(num)
        year = g.get("%sy%s" % (p, n))
        day = g.get("%sd%s" % (p, n))
        if not 1 <= month <= 12:
            return None
        return (int(year) if year else None, month, int(day) if day else None)
    return None


def _range_end(m: "re.Match") -> Optional[str]:
    """The END of a period range as text the engine's reader understands
    ("31.12.2025", or "<month tag> <year>" for a month), or None when the
    match is not a period range (no year on either side, an impossible
    month)."""
    from engine.api._period_detect import _MONTH_TAGS  # noqa: WPS433 — lazy (pipeline cycle)

    start, end = _range_side(m, "a_"), _range_side(m, "b_")
    if start is None or end is None:
        return None
    year, month, day = end
    if year is None:
        if start[0] is None:
            return None
        # "Octombrie 2024 - Martie": the end month follows the start month
        year = start[0] + (1 if month < start[1] else 0)
    if day is not None:
        return "%02d.%02d.%04d" % (day, month, year)
    return "%s %d" % (_MONTH_TAGS[month], year)


def _first_period_range(lines: Sequence[str]) -> Optional[Tuple[str, str]]:
    """(end text, the literal line) of the first period range in reading
    order, or None."""
    for line in lines:
        for m in _RANGE_RE.finditer(line):
            end = _range_end(m)
            if end:
                return end, line.strip()[:160]
    return None


def collapse_period_ranges(text: Optional[str]) -> Optional[str]:
    """``text`` with every period range replaced by its end ("perioada
    01.01.2025 - 31.12.2025" -> "perioada 31.12.2025"): a lexical rewrite,
    so no reader downstream can take a range's START for the period."""
    if not text:
        return text

    def _sub(m: "re.Match") -> str:
        end = _range_end(m)
        return end if end else m.group(0)

    return _RANGE_RE.sub(_sub, str(text))


# ── identification ─────────────────────────────────────────────────────

def _period(header_text: str, filename: Optional[str]) -> Tuple[Optional[str], Optional[Dict[str, str]]]:
    from engine.api._period_detect import detect_period  # noqa: WPS433 — lazy (pipeline cycle)

    lines = [ln for ln in str(header_text or "").split("\n") if ln.strip()]
    found = _first_period_range(lines)
    if found is not None:
        # A period the header states as a range: its END, read by the
        # engine's own reader; the evidence is the literal line.
        got = detect_period(extracted={"closing_balance_date": found[0]}, filename=None)
        end = got.get("proposed_period_end")
        if end:
            return end, {"signal": "closing_balance", "evidence": found[1]}
    text = "\n".join(collapse_period_ranges(ln) or "" for ln in lines)
    fname = collapse_period_ranges(filename)
    got = detect_period(
        extracted={"header_text": text} if text else None,
        filename=fname,
    )
    end = got.get("proposed_period_end")
    if not end:
        return None, None
    signal = str(got.get("signal_used"))
    evidence = str(got.get("evidence_snippet") or "")
    if signal == "filename" and fname != filename:
        evidence = re.sub(r"\s+", " ", str(filename)).strip()[:160]   # the literal filename
    return end, {"signal": signal, "evidence": evidence}


def filename_period_end(filename: Optional[str]) -> Optional[str]:
    """The period a FILENAME alone names ("…_31.12.2025.xlsx", "…FY2025…"),
    through the engine's own detector — or None. A corroborating signal,
    never an identity."""
    if not filename:
        return None
    end, src = _period("", filename)
    return end if src and src.get("signal") == "filename" else None


def identify_text(doc: DocumentText, filename: Optional[str], *, registry: Any = None) -> CompanyIdentity:
    """The identity from already-extracted text. Pure given ``registry``."""
    sources: Dict[str, Dict[str, str]] = {}
    lines = list(doc.header_lines)
    cui: Optional[str] = None
    reg_row: Optional[Dict[str, Any]] = None

    header_names = _header_names(lines)
    sheet_names = _sheet_names(doc.sheet_names)
    file_name = filename_company_name(filename)

    found = _header_cui(lines)
    if found:
        cui = found[0]
        sources["cui"] = {"signal": "document_header_cui", "evidence": found[1]}
        reg_row = _registry_company(registry, cui)

    if cui is None:
        labelled = [(n, ev) for n, s, ev in header_names if s == "document_header_label"]
        if labelled:
            # The document names ITSELF ("Societate: ..."): that name is the
            # only one looked up. Unresolved (no match, several, a namesake
            # that files nothing) means the company is unresolved — never
            # the CUI of another string in the header (a title line, the
            # sheet, the filename). Verifier finding, 2026-09-21.
            candidates = [(labelled[0][0], "registry_name_match", labelled[0][1])]
        else:
            candidates = [(n, "registry_name_match", ev) for n, _s, ev in header_names]
            candidates += [(n, "registry_name_match", "sheet: " + n) for n in sheet_names]
            if file_name and len(normalize_company_name(file_name)) >= FILENAME_NAME_MIN_CHARS:
                candidates.append((file_name, "filename_registry_match", "filename: %s" % filename))
        for name, signal, evidence in candidates:
            hit = _unique_registered_name(registry, name)
            if not hit:
                continue
            printed = _digits_in(lines, hit[0])
            if signal == "filename_registry_match" and not printed:
                # A filename is typed by a person: without the document
                # printing that CUI somewhere in its header, it is a hint.
                sources.setdefault("cui_hint", {
                    "signal": signal, "cui": hit[0],
                    "evidence": "%s == registry %r; the document does not print %s"
                                % (evidence, hit[1].get("name"), hit[0])})
                break
            if hit[2] is not True and not printed:
                # The one registered company with this name files nothing
                # (or the registry cannot say): a namesake, not an identity.
                sources.setdefault("cui_hint", {
                    "signal": signal, "cui": hit[0],
                    "evidence": "%s == registry %r, which has %s; the document does not print %s"
                                % (evidence, hit[1].get("name"),
                                   "no filing" if hit[2] is False else "no known filing", hit[0])})
                continue
            cui, reg_row = hit[0], hit[1]
            sources["cui"] = {"signal": signal,
                              "evidence": "%s == registry %r" % (evidence, reg_row.get("name"))}
            break

    company_name: Optional[str] = None
    if reg_row and reg_row.get("name"):
        company_name = str(reg_row["name"])
        sources["company_name"] = {"signal": "registry", "evidence": "registry CUI %s" % cui}
    elif header_names:
        company_name, signal, evidence = header_names[0]
        sources["company_name"] = {"signal": signal, "evidence": evidence}
    elif sheet_names:
        company_name = sheet_names[0]
        sources["company_name"] = {"signal": "sheet_name", "evidence": "sheet: " + sheet_names[0]}
    elif file_name:
        company_name = file_name
        sources["company_name"] = {"signal": "filename", "evidence": str(filename)}

    caen: Optional[str] = None
    if reg_row and reg_row.get("caen"):
        caen = _normalize_caen(reg_row["caen"])
        if caen:
            sources["caen_code"] = {"signal": "registry", "evidence": "registry CUI %s" % cui}
    if caen is None:
        got_caen = _header_caen(lines)
        if got_caen and got_caen[0]:
            caen = got_caen[0]
            sources["caen_code"] = {"signal": "document_header", "evidence": got_caen[1]}
    industry = industry_key_for_caen(caen) if caen else None
    if industry:
        sources["industry_key"] = {"signal": "caen_catalogue", "evidence": "CAEN %s" % caen}

    period_end, period_src = _period(doc.header_text, filename)
    if period_src:
        sources["period_end"] = period_src

    # "not_a_balance" is asserted ONLY when the bytes were parsed and are
    # demonstrably not a trial balance: readable text, no account rows, and
    # none of the signals a balance carries. A balance whose row layout the
    # reader refused is "uncertain" — the failure is the reader's.
    kind = doc.document_kind
    if kind == "not_a_balance":
        signals = balance_signals(lines, filename, found[0] if found else None)
        if signals:
            kind = "uncertain"
            sources["document_kind"] = {
                "signal": "unparsed_balance",
                "evidence": "%d account row(s) read, yet %s — the reader refused the layout, the document "
                            "is not condemned" % (doc.account_lines, "; ".join(signals))}
        else:
            sources["document_kind"] = {
                "signal": "not_a_balance",
                "evidence": "readable text (%d chars), %d account row(s), no CUI, no balance title, no ledger "
                            "columns" % (doc.text_chars, doc.account_lines)}

    return CompanyIdentity(
        cui=cui,
        company_name=company_name,
        period_end=period_end,
        caen_code=caen,
        industry_key=industry,
        sources=sources,
        document_kind=kind,
    )


def identify_document(content: bytes, filename: str, *, registry: Any = None,
                      mime: Optional[str] = None) -> CompanyIdentity:
    """Who and when a stored document is about, from its own bytes.

    ``registry`` is an ``engine.public_ro.store.PublicRoStore`` (or anything
    with its ``get_company(cui)`` / ``search_companies(q, limit)``), or None.
    ``mime`` is the MIME type the upload declared (``documents.mime_type``):
    with the name it picks the branch the pipeline reads the bytes on, and
    the identifier reads them on the same one (``extract_document_text``).
    It picks a READER only — nothing of the identity is taken from it.
    Empty or unreadable bytes still yield an identity — from the filename
    alone, with ``document_kind="unreadable"``."""
    doc = extract_document_text(content, filename, mime)
    ident = identify_text(doc, filename, registry=registry)
    if doc.read_error:
        srcs = dict(ident.sources)
        srcs["document_kind"] = {"signal": "read_error", "evidence": doc.read_error}
        ident = replace(ident, sources=srcs)
    return ident


# ── operator-verified identities ───────────────────────────────────────

def match_known_identity(rules: Sequence[Mapping[str, Any]], *, user_id: Optional[str],
                         content_sha256: Optional[str], filename: Optional[str]
                         ) -> Optional[Mapping[str, Any]]:
    """The first rule for this document: an exact content hash beats a
    filename glob; a rule naming a ``user_id`` applies only to that user."""
    def _scoped(rule: Mapping[str, Any]) -> bool:
        uid = rule.get("user_id")
        return not uid or (user_id is not None and str(uid) == str(user_id))

    if content_sha256:
        for rule in rules:
            if _scoped(rule) and rule.get("content_sha256") and \
                    str(rule["content_sha256"]).lower() == content_sha256.lower():
                return rule
    if filename:
        low = filename.lower()
        for rule in rules:
            glob = rule.get("filename_glob")
            if _scoped(rule) and glob and fnmatch.fnmatchcase(low, str(glob).lower()):
                return rule
    return None


def apply_known_identity(identity: CompanyIdentity, rule: Mapping[str, Any], *,
                         registry: Any = None) -> Tuple[CompanyIdentity, Optional[str]]:
    """Layer an operator-verified identity over what the document says.

    Returns (identity, conflict). A CUI the document itself prints always
    wins: a rule that disagrees with it is reported as ``conflict`` and not
    applied. A rule with ``"cui": null`` and a ``company_name`` pins a
    company that has no known CUI (it also blocks a registry NAME match,
    which could otherwise hand the book a stranger's CUI)."""
    evidence = str(rule.get("evidence") or "operator-verified identity")
    has_cui_key = "cui" in rule
    rule_cui = normalize_cui(rule.get("cui")) if rule.get("cui") is not None else None
    if rule.get("cui") is not None and rule_cui is None:
        return identity, "rule CUI %r fails the control digit" % (rule.get("cui"),)

    doc_signal = (identity.sources.get("cui") or {}).get("signal")
    if identity.cui and doc_signal == "document_header_cui":
        if rule_cui != identity.cui:
            return identity, ("document prints CUI %s, rule says %s — the document wins"
                              % (identity.cui, rule_cui))
        # Same company: the document keeps its CUI and name; a CAEN the
        # operator verified still layers on (it used to be dropped here).
        rule_caen = _normalize_caen(rule.get("caen_code"))
        if not rule_caen or rule_caen == identity.caen_code:
            return identity, None
        sources = dict(identity.sources)
        sources["caen_code"] = {"signal": "operator_verified", "evidence": evidence}
        industry = industry_key_for_caen(rule_caen)
        if industry:
            sources["industry_key"] = {"signal": "caen_catalogue", "evidence": "CAEN %s" % rule_caen}
        else:
            sources.pop("industry_key", None)
        return replace(identity, caen_code=rule_caen, industry_key=industry, sources=sources), None

    sources = dict(identity.sources)
    if rule_cui:
        reg = _registry_company(registry, rule_cui)
        name = (reg or {}).get("name") or rule.get("company_name") or identity.company_name
        caen = _normalize_caen(rule.get("caen_code")) or _normalize_caen((reg or {}).get("caen")) \
            or identity.caen_code
        sources["cui"] = {"signal": "operator_verified", "evidence": evidence}
        sources["company_name"] = {"signal": "registry" if (reg or {}).get("name") else "operator_verified",
                                   "evidence": ("registry CUI %s" % rule_cui) if (reg or {}).get("name") else evidence}
        cui = rule_cui
    elif has_cui_key and rule.get("company_name"):
        name = str(rule["company_name"])
        caen = _normalize_caen(rule.get("caen_code")) or None
        sources.pop("cui", None)
        sources["company_name"] = {"signal": "operator_verified", "evidence": evidence}
        cui = None
    else:
        return identity, "rule carries neither a CUI nor a company name"

    if caen:
        sources["caen_code"] = {"signal": "operator_verified" if rule.get("caen_code") else "registry",
                                "evidence": evidence if rule.get("caen_code") else "registry CUI %s" % cui}
    else:
        sources.pop("caen_code", None)
    industry = industry_key_for_caen(caen) if caen else None
    if industry:
        sources["industry_key"] = {"signal": "caen_catalogue", "evidence": "CAEN %s" % caen}
    else:
        sources.pop("industry_key", None)
    return replace(identity, cui=cui, company_name=name, caen_code=caen,
                   industry_key=industry, sources=sources), None
